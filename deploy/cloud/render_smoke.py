"""Synthetic Linux render smoke: Chinese chart, Word TOC/PDF and PPT PDF."""
import argparse
import json
import os
from pathlib import Path
import subprocess
import sys


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if sys.platform != 'linux':
        raise SystemExit('LibreOffice smoke must run on Linux')
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=False)
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from pptx import Presentation
    from pptx.util import Inches as SlideInches
    import fitz
    from portal.product_assets.libreoffice_render import render

    runtime = Path('/app/services/business/backend/portal/product_assets/mermaid-runtime')
    (root / 'content.json').write_text(json.dumps({'blocks': [{'type': 'figure', 'id': 'smoke',
        'caption': '中文流程测试', 'path': 'chart.png'}]}, ensure_ascii=False), encoding='utf-8')
    (root / 'private-diagrams.json').write_text(json.dumps({'diagrams': [{'figure_id': 'smoke',
        'source': 'flowchart LR\n A[资料接收] --> B[方案编制] --> C[人工复核]'}]}, ensure_ascii=False), encoding='utf-8')
    environment = dict(os.environ, PORTAL_CHROMIUM_PATH='/usr/bin/chromium')
    subprocess.run(['node', str(runtime / 'render.mjs'), str(root / 'content.json'),
                    str(root / 'private-diagrams.json')], env=environment, check=True, timeout=120)
    document = Document()
    document.add_heading('云部署合成测试', 0)
    document.add_paragraph('仅用于测试中文字体、目录刷新和图表渲染，不是业务交付物。')
    paragraph = document.add_paragraph()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
    paragraph._p.append(field)
    document.add_page_break()
    document.add_heading('中文流程测试', 1)
    document.add_picture(str(root / 'chart.png'), width=Inches(6))
    document.add_heading('渲染验收说明', 1)
    document.add_paragraph('需要人工检查目录页码、中文缺字、图表裁切和分页。')
    source = root / 'synthetic.docx'
    document.save(source)
    result = render(source, root / 'word', executable='/usr/bin/libreoffice',
                    uno_python='/usr/bin/python3', timeout=300)
    if not result.get('rendered') or not result.get('structural_pass'):
        raise RuntimeError('Synthetic DOCX rendering failed; inspect private render.json')
    with fitz.open(root / 'word/document.pdf') as pdf:
        text = '\n'.join(page.get_text() for page in pdf)
        if '中文流程测试' not in text:
            raise RuntimeError('Chinese PDF text verification failed')
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = '中文演示渲染测试'
    slide.shapes.add_picture(str(root / 'chart.png'), SlideInches(.5), SlideInches(1.5), width=SlideInches(9))
    deck.save(root / 'synthetic.pptx')
    ppt_output = root / 'ppt'
    ppt_result = render(root / 'synthetic.pptx', ppt_output, executable='/usr/bin/libreoffice',
                        uno_python='/usr/bin/python3', timeout=300, kind='powerpoint')
    if not ppt_result.get('rendered') or not ppt_result.get('structural_pass'):
        raise RuntimeError('Synthetic PPT rendering failed; inspect private ppt/render.json')
    with fitz.open(ppt_output / 'slides.pdf') as pdf:
        if len(pdf) != 1 or '中文演示' not in pdf[0].get_text():
            raise RuntimeError('Synthetic PPT PDF verification failed')
    if ppt_result.get('page_count') != 1 or ppt_result.get('slide_count') != 1:
        raise RuntimeError('Synthetic PPT renderer page count verification failed')
    if len(ppt_result.get('images', [])) != 1 or not (ppt_output / 'page-001.png').is_file():
        raise RuntimeError('Synthetic PPT page image verification failed')
    report = {'synthetic_structure': 'passed', 'word_pages': result['page_count'],
              'toc_count': result['toc_count'], 'ppt_pages': ppt_result['page_count'],
              'visual_review': 'pending', 'real_business_acceptance': 'pending'}
    (root / 'smoke-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))


if __name__ == '__main__':
    main()

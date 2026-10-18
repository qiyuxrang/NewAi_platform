"""Bounded host adapter for the pinned, unmodified PPT Master native exporter."""
import argparse
import base64
import contextlib
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from xml.etree import ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parent
SVG_NS = "http://www.w3.org/2000/svg"
TAGS = {"svg", "g", "defs", "rect", "circle", "ellipse", "line", "polyline", "polygon",
        "path", "text", "tspan", "title", "desc", "linearGradient", "radialGradient", "stop", "marker", "image"}
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}


def sha256(raw):
    return hashlib.sha256(raw).hexdigest()


def provenance():
    raw = (ROOT / "SOURCES.json").read_bytes()
    source = json.loads(raw)
    for item in source["files"]:
        if sha256((ROOT / item["path"]).read_bytes()) != item["sha256"]:
            raise ValueError("runtime_source_hash_mismatch")
    if sha256((ROOT / "LICENSE").read_bytes()) != source["license_sha256"]:
        raise ValueError("runtime_license_hash_mismatch")
    return {"engine": source["engine"], "source_commit": source["source_commit"],
            "source_sha256": sha256(raw), "license": source["license"],
            "runtime_sha256": sha256(Path(__file__).read_bytes() + raw)}


def read_input(path, root, limit):
    if path.is_symlink() or not path.resolve().is_relative_to(root):
        raise ValueError("input_path_invalid")
    if path.stat().st_size > limit:
        raise ValueError("input_too_large")
    return path.read_bytes()


def validate_svg(raw):
    if re.search(br"<!DOCTYPE|<!ENTITY|<\?xml-stylesheet", raw, re.I):
        raise ValueError("unsafe_svg_declaration")
    root = ET.fromstring(raw)
    if root.tag != f"{{{SVG_NS}}}svg" or root.get("viewBox", "").split() != ["0", "0", "1280", "720"]:
        raise ValueError("svg_canvas_invalid")
    for element in root.iter():
        if not element.tag.startswith(f"{{{SVG_NS}}}") or element.tag.split("}")[-1] not in TAGS:
            raise ValueError("svg_element_invalid")
        for key, value in element.attrib.items():
            name = key.split("}")[-1]
            if name.lower().startswith("on") or name in {"style", "class", "base", "data-icon"}:
                raise ValueError("svg_attribute_invalid")
            if name == "href":
                if element.tag != f"{{{SVG_NS}}}image" or not re.fullmatch(r"data:image/(png|jpeg);base64,[A-Za-z0-9+/=\s]+", value):
                    raise ValueError("svg_external_reference")
                data = base64.b64decode(value.split(",", 1)[1], validate=True)
                from PIL import Image
                with Image.open(io.BytesIO(data)) as image:
                    if image.format not in {"PNG", "JPEG"} or image.width * image.height > 20_000_000:
                        raise ValueError("svg_image_invalid")
                    image.verify()
            if "url(" in value and not re.fullmatch(r"url\(#[A-Za-z_][\w:.-]*\)", value):
                raise ValueError("svg_external_reference")
    return root


def validate_chart(chart):
    if not isinstance(chart, dict) or set(chart) - {"title", "unit", "categories", "values", "box"}:
        raise ValueError("chart_invalid")
    categories, values, box = chart.get("categories"), chart.get("values"), chart.get("box")
    if (not isinstance(categories, list) or not 1 <= len(categories) <= 30
            or any(not isinstance(value, str) or not 1 <= len(value) <= 100 for value in categories)
            or not isinstance(values, list) or len(values) != len(categories)
            or any(type(value) not in {int, float} or not math.isfinite(value) for value in values)
            or not isinstance(box, list) or len(box) != 4
            or any(type(value) not in {int, float} or not math.isfinite(value) for value in box)
            or box[0] < 0 or box[1] < 0 or box[2] < 16 or box[3] < 16
            or box[0] + box[2] > 1280 or box[1] + box[3] > 720
            or any(not isinstance(chart.get(key, ""), str) or len(chart.get(key, "")) > 200 for key in ("title", "unit"))):
        raise ValueError("chart_invalid")


def chart_design(plan):
    supplied = plan.get("design", {})
    if not isinstance(supplied, dict):
        raise ValueError("chart_design_invalid")
    design = {key: supplied.get(key, value) for key, value in {
        "font": "Microsoft YaHei", "background": "#F5F8FC", "ink": "#152C48",
        "muted": "#495C73", "accent": "#174EA6"}.items()}
    if (not isinstance(design["font"], str) or not 1 <= len(design["font"]) <= 100
            or any(not isinstance(value, str) or not re.fullmatch(r"#[0-9a-fA-F]{6}", value)
                   for key, value in design.items() if key != "font")):
        raise ValueError("chart_design_invalid")
    return design


def add_native_objects(output, charts, design):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Emu, Pt
    def font(value, size):
        value.name, value.size = design["font"], Pt(size)
        value.color.rgb = RGBColor.from_string(design["ink"][1:])
        # Font.name sets only a:latin; CJK labels also need an explicit a:ea.
        for name, following in (("ea", ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:extLst")),
                                ("cs", ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:extLst"))):
            element = value._rPr.find("a:" + name, NS)
            if element is None:
                element = OxmlElement("a:" + name)
                value._rPr.insert_element_before(element, *following)
            element.set("typeface", design["font"])
    def text_frame(value, size):
        for paragraph in value.paragraphs:
            font(paragraph.font, size)
            for run in paragraph.runs:
                font(run.font, size)
    presentation = Presentation(str(output))
    for index, slide in enumerate(presentation.slides):
        # Upstream emits SVG lines as shapes. Preserve geometry/IDs and expose
        # them as real editable connectors using the equivalent OOXML element.
        for shape in list(slide._element.iter(f"{{{NS['p']}}}sp")):
            properties = shape.find("p:nvSpPr", NS)
            name = properties.find("p:cNvPr", NS).get("name", "")
            if not re.fullmatch(r"Line \d+", name):
                continue
            shape.tag = f"{{{NS['p']}}}cxnSp"
            properties.tag = f"{{{NS['p']}}}nvCxnSpPr"
            properties.find("p:cNvSpPr", NS).tag = f"{{{NS['p']}}}cNvCxnSpPr"
            for body in list(shape.findall("p:txBody", NS)):
                shape.remove(body)
        chart = charts[index]
        if chart is not None:
            data = CategoryChartData()
            data.categories = chart["categories"]
            data.add_series(chart.get("unit") or chart.get("title") or "Values", chart["values"])
            native = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                *(Emu(round(value * 9525)) for value in chart["box"]), data).chart
            font(native.font, 20)
            native.has_legend = False
            for axis in (native.category_axis, native.value_axis):
                font(axis.tick_labels.font, 20)
                axis.format.line.color.rgb = RGBColor.from_string(design["muted"][1:])
            native.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(design["muted"][1:])
            series = native.series[0]
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(design["accent"][1:])
            series.format.line.color.rgb = RGBColor.from_string(design["accent"][1:])
            native.plots[0].has_data_labels = True
            labels = native.plots[0].data_labels
            labels.show_value = True
            font(labels.font, 20)
            background = OxmlElement("c:spPr")
            fill = OxmlElement("a:solidFill")
            color = OxmlElement("a:srgbClr")
            color.set("val", design["background"][1:])
            fill.append(color)
            background.append(fill)
            native._chartSpace.insert_element_before(background, "c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst")
            if chart.get("title"):
                native.has_title = True
                native.chart_title.text_frame.text = chart["title"]
                text_frame(native.chart_title.text_frame, 24)
            if chart.get("unit"):
                native.value_axis.has_title = True
                native.value_axis.axis_title.text_frame.text = chart["unit"]
                text_frame(native.value_axis.axis_title.text_frame, 20)
    presentation.save(str(output))


def inventory(output):
    pages = []
    with zipfile.ZipFile(output) as archive:
        slides = sorted((name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                        key=lambda name: int(re.search(r"slide(\d+)", name).group(1)))
        for name in slides:
            root = ET.fromstring(archive.read(name))
            pages.append({"native_shapes": len(root.findall(".//p:sp", NS)),
                          "connectors": len(root.findall(".//p:cxnSp", NS)),
                          "charts": len(root.findall(".//c:chart", NS)),
                          "images": len(root.findall(".//p:pic", NS))})
        return {"slides": len(pages), "per_slide": pages,
                **{key: sum(page[key] for page in pages) for key in ("native_shapes", "connectors", "charts", "images")},
                "media_count": sum(name.startswith("ppt/media/") and not name.endswith("/") for name in archive.namelist())}


def run(project, output, validate_only=False):
    info = provenance()
    project, output = Path(project).resolve(), Path(output).absolute()
    files = sorted((project / "svg_output").glob("*.svg"))
    if not 1 <= len(files) <= 80:
        raise ValueError("svg_count_invalid")
    inputs, notes = [], {}
    for path in files:
        raw = read_input(path, project, 8 * 1024 * 1024)
        validate_svg(raw)
        inputs.append((path.name, raw))
        note = project / "notes" / (path.stem + ".md")
        if note.exists():
            notes[path.stem] = read_input(note, project, 256 * 1024).decode("utf-8")
    if sum(len(raw) for _, raw in inputs) > 64 * 1024 * 1024:
        raise ValueError("input_too_large")
    plan_path = project / "plan.json"
    plan = json.loads(read_input(plan_path, project, 1024 * 1024)) if plan_path.exists() else {}
    generation_only = plan.get("generation_policy") == "generation-only-v1"
    design = chart_design(plan)
    pages = plan.get("pages", [{} for _ in files])
    if not isinstance(pages, list) or len(pages) != len(files) or any(not isinstance(page, dict) for page in pages):
        raise ValueError("plan_pages_invalid")
    charts = [page.get("chart") for page in pages]
    for chart in charts:
        if chart is not None:
            validate_chart(chart)
    sys.path.insert(0, str(ROOT / "upstream" / "scripts"))
    # A private snapshot keeps upstream discovery away from arbitrary project
    # sidecars and makes its file reads match the already checked inputs.
    with tempfile.TemporaryDirectory(prefix="ppt-master-native-") as temporary:
        directory = Path(temporary)
        paths = []
        for name, raw in inputs:
            path = directory / name
            path.write_bytes(raw)
            paths.append(path)
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            from svg_to_pptx import create_pptx_with_native_svg
            quality = []
            if not generation_only:
                from svg_quality_checker import SVGQualityChecker
                checker = SVGQualityChecker(template_mode=True)
                quality = [checker.check_file(str(path), expected_format="ppt169") for path in paths]
                if any(not page["passed"] for page in quality):
                    raise ValueError("svg_quality_failed:" + json.dumps([page["errors"] for page in quality], ensure_ascii=True))
            candidate = directory / "native.pptx"
            if not validate_only:
                if not create_pptx_with_native_svg(paths, candidate, verbose=False, use_native_shapes=True,
                        use_compat_mode=False, transition=None, notes=notes, workers=1):
                    raise ValueError("native_export_failed")
                add_native_objects(candidate, charts, design)
                counts = inventory(candidate)
                if counts["slides"] != len(files):
                    raise ValueError("native_slide_count_mismatch")
                output.parent.mkdir(parents=True, exist_ok=True)
                with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".pptx", delete=False) as temporary_output:
                    temporary_output.write(candidate.read_bytes())
                    temporary_path = temporary_output.name
                try:
                    os.replace(temporary_path, output)
                finally:
                    Path(temporary_path).unlink(missing_ok=True)
    result = {**info, "validated": True, "validate_only": validate_only,
              "content_checks": "not_run" if generation_only else "program_only",
              "quality": ([{"file": path.name, "status": "not_run"} for path in paths] if generation_only else
                  [{"file": page["file"], "passed": page["passed"], "warnings": page["warnings"]} for page in quality]),
              "sources": [{"file": name, "sha256": sha256(raw)} for name, raw in inputs]}
    if validate_only:
        result["slides"] = len(files)
    else:
        result.update(counts, sha256=sha256(output.read_bytes()), path=str(output))
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    try:
        print(json.dumps(run(args.project, args.output, args.validate_only), ensure_ascii=True))
    except Exception as error:
        print(json.dumps({"error": str(error)[:4000], "type": type(error).__name__}, ensure_ascii=True))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

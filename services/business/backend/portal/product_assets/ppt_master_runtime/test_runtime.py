"""Run with python -B test_runtime.py; no API calls or platform state."""
import base64
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
import subprocess
import sys
from xml.etree import ElementTree as ET

from runner import run, validate_svg

SVG = '''<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">
<defs><marker id="arrow" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="6" markerHeight="6" orient="auto"><path d="M 0 0 L 10 5 L 0 10 Z" fill="#234567"/></marker></defs>
<rect x="0" y="0" width="1280" height="720" fill="#FFFFFF"/>
<text x="64" y="80" font-size="36" font-family="Arial" fill="#234567">Editable native text</text>
<rect x="64" y="140" width="200" height="100" fill="#DDEEFF"/>
<line x1="264" y1="190" x2="440" y2="190" stroke="#234567" stroke-width="3" marker-end="url(#arrow)"/>
<line x1="100" y1="260" x2="300" y2="280" stroke="#234567" stroke-width="3"/>
<line x1="400" y1="310" x2="320" y2="310" stroke="#345678" stroke-width="4"/>
<line x1="360" y1="390" x2="360" y2="330" stroke="#456789" stroke-width="2"/>
<line x1="500" y1="350" x2="450" y2="300" stroke="#56789A" stroke-width="5"/>
<path d="M 480 150 L 580 220 L 480 240 Z" fill="#234567"/>
</svg>'''

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


class NativeRuntimeTest(unittest.TestCase):
    def project(self, directory, slides=1):
        root = Path(directory) / "project"
        (root / "svg_output").mkdir(parents=True)
        (root / "notes").mkdir()
        for index in range(slides):
            (root / "svg_output" / f"{index + 1:02d}.svg").write_text(SVG, encoding="utf-8")
            (root / "notes" / f"{index + 1:02d}.md").write_text("Speaker note with source [R1].", encoding="utf-8")
        return root

    def test_single_and_multiple_pages_remain_native(self):
        from pptx import Presentation
        for total in (1, 2, 17):
            with self.subTest(total=total), tempfile.TemporaryDirectory() as directory:
                root = self.project(directory, total)
                output = Path(directory) / "native.pptx"
                result = run(root, output)
                self.assertEqual(result["slides"], total)
                self.assertEqual(result["connectors"], total * 5)
                self.assertGreaterEqual(result["native_shapes"], total * 4)
                self.assertEqual(result["images"], 0)
                self.assertEqual(result["media_count"], 0)
                with zipfile.ZipFile(output) as archive:
                    slide = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
                connectors = slide.findall(".//p:cxnSp", NS)
                self.assertEqual(len(connectors), 5)
                geometries = [connector.find("p:spPr/a:prstGeom", NS) for connector in connectors]
                self.assertTrue(all(geometry is not None and geometry.get("prst") == "line" for geometry in geometries))
                self.assertIsNotNone(connectors[0].find("p:spPr/a:ln/a:tailEnd", NS))
                reverse_horizontal = connectors[2].find("p:spPr", NS)
                self.assertEqual(reverse_horizontal.find("a:xfrm", NS).get("flipH"), "1")
                self.assertEqual(reverse_horizontal.find("a:xfrm/a:ext", NS).get("cy"), "0")
                self.assertEqual(reverse_horizontal.find("a:ln", NS).get("w"), "38100")
                self.assertEqual(reverse_horizontal.find("a:ln/a:solidFill/a:srgbClr", NS).get("val"), "345678")
                reverse_vertical = connectors[3].find("p:spPr", NS)
                self.assertEqual(reverse_vertical.find("a:xfrm", NS).get("flipV"), "1")
                self.assertEqual(reverse_vertical.find("a:xfrm/a:ext", NS).get("cx"), "0")
                reverse_diagonal = connectors[4].find("p:spPr/a:xfrm", NS)
                self.assertEqual((reverse_diagonal.get("flipH"), reverse_diagonal.get("flipV")), ("1", "1"))
                deck = Presentation(str(output))
                self.assertEqual(len(deck.slides), total)
                self.assertTrue(any("Editable native text" in shape.text for shape in deck.slides[0].shapes if shape.has_text_frame))
                self.assertIn("source [R1]", deck.slides[0].notes_slide.notes_text_frame.text)
                self.assertEqual(result["source_commit"], "d6f6cdfa97eaafd73b9cc2b0343b94a9f3998fc4")

    def test_explicit_chart_and_embedded_image(self):
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            root = self.project(directory)
            buffer = io.BytesIO()
            Image.new("RGB", (20, 20), "blue").save(buffer, format="PNG")
            encoded = base64.b64encode(buffer.getvalue()).decode()
            page = root / "svg_output/01.svg"
            page.write_text(SVG.replace("</svg>", f'<image x="1000" y="80" width="100" height="100" href="data:image/png;base64,{encoded}"/></svg>'), encoding="utf-8")
            (root / "plan.json").write_text(json.dumps({"design": {"font": "Microsoft YaHei", "accent": "#174EA6", "ink": "#152C48", "background": "#F5F8FC"}, "pages": [{"chart": {"title": "Explicit inputs", "unit": "units", "categories": ["A", "B"], "values": [2, 4], "box": [64, 330, 600, 300]}}]}), encoding="utf-8")
            output = Path(directory) / "native.pptx"
            result = run(root, output)
            self.assertEqual((result["charts"], result["images"], result["media_count"]), (1, 1, 1))
            deck = Presentation(str(output))
            chart = next(shape.chart for shape in deck.slides[0].shapes if shape.has_chart)
            self.assertEqual(list(chart.series[0].values), [2, 4])
            self.assertEqual(chart.category_axis.tick_labels.font.size.pt, 20)
            self.assertEqual(chart.value_axis.tick_labels.font.name, "Microsoft YaHei")
            self.assertEqual(chart.chart_title.text_frame.paragraphs[0].font.size.pt, 24)
            self.assertEqual(chart.plots[0].data_labels.font.size.pt, 20)
            self.assertEqual(str(chart.series[0].format.fill.fore_color.rgb), "174EA6")
            with zipfile.ZipFile(output) as archive:
                self.assertTrue(any(name.startswith("ppt/embeddings/") for name in archive.namelist()))
                chart_xml = archive.read("ppt/charts/chart1.xml")
                self.assertIn(b'val="F5F8FC"', chart_xml)
                self.assertIn(b'<a:ea typeface="Microsoft YaHei"', chart_xml)

    def test_validate_only_never_writes_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = self.project(directory)
            output = Path(directory) / "native.pptx"
            self.assertTrue(run(root, output, True)["validated"])
            self.assertFalse(output.exists())
            result = subprocess.run([sys.executable, "-B", str(Path(__file__).with_name("runner.py")),
                                     str(root), str(output), "--validate-only"], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(len(result.stdout.splitlines()), 1)
            self.assertTrue(json.loads(result.stdout)["validate_only"])
            self.assertEqual(result.stderr, "")

    def test_unsafe_svg_and_invalid_chart_are_rejected_without_output_changes(self):
        for value in ('<script>alert(1)</script>', '<image href="file:///secret.png"/>', '<image href="https://example.com/a.png"/>', '<rect onclick="alert(1)"/>', '<foreignObject/>'):
            with self.subTest(value=value), self.assertRaises(ValueError):
                validate_svg(SVG.replace("</svg>", value + "</svg>").encode())
        with tempfile.TemporaryDirectory() as directory:
            root = self.project(directory)
            (root / "plan.json").write_text('{"pages":[{"chart":{"categories":["A"],"values":[true],"box":[0,0,200,200]}}]}', encoding="utf-8")
            output = Path(directory) / "native.pptx"
            output.write_bytes(b"previous-version")
            with self.assertRaisesRegex(ValueError, "chart_invalid"):
                run(root, output)
            self.assertEqual(output.read_bytes(), b"previous-version")


if __name__ == "__main__":
    unittest.main()

"""Private XLSX exports built exclusively from immutable match snapshots."""
from datetime import timedelta
import hashlib
import importlib.util
import io
import json
import os
import uuid
from urllib.parse import urlsplit

from django.conf import settings
from django.core.exceptions import ValidationError
from django.utils import timezone

from .baichuan_assistant_models import BaichuanEvidence, BaichuanExport, BaichuanMatchSnapshot
from .baichuan_assistant_service import require_assistant, scope, storage_path
from .baichuan_conditions import OUTPUT_FIELDS
from .baichuan_jobs import enqueue_job, job_for, leased_job, save_job
from .baichuan_service import BaichuanError
from .product_parallel import atomic
from .security import audit

MAX_ROWS = 5000
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_SHEET_ROWS = 100000
LABELS = {"name": "企业名称", "province": "省份", "city": "城市", "website": "官网",
          "phone": "电话", "email": "邮箱", "credit_code": "统一社会信用代码", "address": "地址"}
STATUS_LABELS = {"confirmed": "已确认匹配", "candidate": "待核实候选", "excluded": "已排除"}
VERDICTS = {"satisfied": "满足", "failed": "不满足", "unknown": "未知", "conflict": "冲突"}


def export_available():
    return importlib.util.find_spec("openpyxl") is not None


def _writer():
    try:
        import openpyxl
    except ImportError:
        raise BaichuanError("export_unavailable", "Excel 导出组件暂不可用，请联系维护人员。", 503) from None
    return openpyxl


def snapshot_for_export(user, snapshot_id):
    require_assistant(user)
    try:
        snapshot = BaichuanMatchSnapshot.objects.select_related("job", "conversation").get(pk=snapshot_id,
            conversation__owner_id=user.pk, conversation__scope=scope(), job__scope=scope(), job__owner_id=user.pk)
    except (BaichuanMatchSnapshot.DoesNotExist, ValidationError, ValueError, TypeError):
        raise BaichuanError("not_found", "结果不存在或不可访问。", 404) from None
    job_for(user, snapshot.job_id)
    return snapshot


def _selection(snapshot, payload):
    allowed = {"client_request_id", "scope", "include_candidates", "limit", "fields", "result_revision"}
    if not isinstance(payload, dict) or set(payload) - allowed:
        raise BaichuanError("invalid_request", "导出参数无效。")
    mode = payload.get("scope", "confirmed")
    if mode not in {"confirmed", "confirmed_and_candidates"}:
        raise BaichuanError("invalid_request", "请选择确认结果或明确包含待核实候选。")
    include = payload.get("include_candidates", mode == "confirmed_and_candidates")
    if type(include) is not bool or (mode == "confirmed_and_candidates" and not include):
        raise BaichuanError("invalid_request", "候选导出范围无效。")
    output = snapshot.condition_tree.get("output", {})
    limit = payload.get("limit", output.get("limit", 20))
    if type(limit) is not int or not 1 <= limit <= MAX_ROWS:
        raise BaichuanError("invalid_request", f"导出条数须为 1 至 {MAX_ROWS}。")
    fields = payload.get("fields", output.get("fields", ["name", "phone"]))
    if (not isinstance(fields, list) or not fields or len(fields) > len(OUTPUT_FIELDS)
            or any(not isinstance(field, str) or field not in OUTPUT_FIELDS for field in fields)
            or len(set(fields)) != len(fields)):
        raise BaichuanError("invalid_request", "导出字段无效。")
    if "name" not in fields:
        fields = ["name", *fields]
    revision = payload.get("result_revision", snapshot.result_revision)
    if type(revision) is not int or revision != snapshot.result_revision:
        raise BaichuanError("revision_conflict", "结果版本已变更，请选择明确的结果快照。", 409)
    return {"include_candidates": include, "limit": limit}, fields


def enqueue_export(user, snapshot_id, payload, *, message=None):
    _writer()
    with atomic():
        snapshot = snapshot_for_export(user, snapshot_id)
        selection, fields = _selection(snapshot, payload)
        request = {"snapshot_id": str(snapshot.pk), "selection": selection, "fields": fields,
                   "result_revision": snapshot.result_revision, "condition_revision": snapshot.condition_revision}
        job, created = enqueue_job(user, kind="export", payload=request,
            idempotency_key=payload.get("client_request_id"), conversation=snapshot.conversation, message=message)
        if created:
            BaichuanExport.objects.create(owner=user, snapshot=snapshot, job=job,
                selection=selection, fields=fields,
                authorization={"owner_id": str(user.pk), "scope": scope(),
                               "condition_revision": snapshot.condition_revision,
                               "result_revision": snapshot.result_revision})
            audit(user, "baichuan_export_enqueued", job.pk, changes=[str(snapshot.pk)])
        return job


def _cell(sheet, row, column, value):
    # OOXML string type, plus text number format, prevents formula evaluation
    # without changing visible leading zeros or adding apostrophes to the data.
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    text = "" if value is None else str(value)
    if len(text) > 32767 or any(ord(c) < 32 and c not in "\t\n\r" for c in text):
        raise BaichuanError("export_cell_limit", "结果包含超长或非法单元格内容，未截断导出。", 413)
    cell = sheet.cell(row, column, text)
    cell.data_type, cell.number_format = "s", "@"
    return cell


def _row(sheet, values, positions):
    # Every cell in these sheets is written here. Track positions per workbook
    # instead of rescanning all existing cells for every appended row.
    number = positions.get(sheet, 0) + 1
    if number > MAX_SHEET_ROWS:
        raise BaichuanError("export_row_limit", "结果超过工作表安全行数，未截断导出。", 413)
    for column, value in enumerate(values, 1):
        _cell(sheet, number, column, value)
    positions[sheet] = number
    return number


def _selected(snapshot, selection):
    statuses = {"confirmed", "candidate"} if selection["include_candidates"] else {"confirmed"}
    return [row for row in snapshot.results if row.get("status") in statuses][:selection["limit"]]


def _source_label(source):
    try:
        parsed = urlsplit(source.url)
        if parsed.scheme in {"http", "https"} and parsed.hostname and not parsed.username and not parsed.password:
            return source.url
    except ValueError:
        pass
    # Storage keys and import_ref can contain private filesystem locations.
    return source.title or "导入来源（内部引用已隐藏）"


def public_locator(locator):
    """Keep source positions while hiding acquisition jobs and private storage."""
    allowed = {"page", "pages", "paragraph", "table", "row", "column", "cell", "sheet", "block", "block_id",
               "text_start", "text_end", "section", "primary_section", "code", "content_kind", "location", "selector"}
    if not isinstance(locator, dict):
        return {}
    return {key: public_locator(value) if isinstance(value, dict) else value
            for key, value in locator.items() if key in allowed}


CERTIFICATE_COLUMNS = [('certificate_no', '证书编号'), ('holder_type', '证书主体类型'),
    ('holder_name', '证书主体名称'), ('holder_company_id', '证书主体企业ID'),
    ('category', '资质类别'), ('grade', '证书等级'), ('issuer', '颁发方'),
    ('valid_from', '有效起始日'), ('valid_until', '有效截止日'),
    ('verification_status', '证书核实状态'), ('assertion_status', '证书声明状态'),
    ('revocation_status', '撤销状态'), ('certificate_id', '证书ID'), ('standard', '适用标准'),
    ('scope_text', '证书适用范围'), ('verified_at', '证书核实时间'),
    ('evidence_access_status', '来源获取状态')]


def _frozen_certificates(row, condition, source_id, export_scope):
    """Never read current certificates to fill legacy or altered snapshots."""
    return [record for record in row.get('certificate_records', []) if isinstance(record, dict)
        and record.get('company_id') == row.get('company_id')
        and record.get('evidence_scope') == export_scope
        and record.get('evidence_id') == source_id
        and condition.get('condition_id') in record.get('condition_ids', [])]


def build_workbook(export, generated_at):
    openpyxl = _writer()
    snapshot = export.snapshot
    selected = _selected(snapshot, export.selection)
    # Reject costly snapshots explicitly before allocating worksheet cells.
    if sum(max(1, sum(max(1, len(_frozen_certificates(row, condition, identifier,
            export.authorization['scope']))) for identifier in condition.get('evidence_ids', [])))
           for row in selected for condition in row.get("conditions", [])) > MAX_SHEET_ROWS - 7:
        raise BaichuanError("export_row_limit", "逐项依据超过安全导出上限，请缩小企业范围。", 413)
    identifiers = set()
    for row in selected:
        for condition in row.get("conditions", []):
            for identifier in condition.get("evidence_ids", []):
                try:
                    identifiers.add(uuid.UUID(identifier))
                except (ValueError, TypeError, AttributeError):
                    continue
    sources = {}
    ordered = sorted(identifiers)
    for index in range(0, len(ordered), 500):
        for source in BaichuanEvidence.objects.filter(pk__in=ordered[index:index + 500], scope=export.authorization["scope"]):
            sources[str(source.pk)] = source
    counts = {kind: sum(row["status"] == kind for row in selected) for kind in ["confirmed", "candidate"]}
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheets = [workbook.create_sheet(name) for name in ["企业结果", "逐项依据", "待核实信息"]]
    positions = {}
    metadata = [
        ["结果快照", str(snapshot.pk), "条件版本", snapshot.condition_revision, "结果修订", snapshot.result_revision],
        ["数据截止时间", snapshot.fact_cutoff_at.isoformat(), "研究状态", snapshot.research_status, "导出时间", generated_at.isoformat()],
        ["实际导出企业数", len(selected), "已确认匹配", counts["confirmed"], "待核实候选", counts["candidate"]],
        ["导出范围", "确认结果及明确选择的候选" if export.selection["include_candidates"] else "仅已确认匹配",
         "条数上限", export.selection["limit"], "字段", ",".join(export.fields)],
        ["说明", "候选仅列入待核实信息，不能视为合格企业。空白表示该固定快照未包含字段，不读取后来修改的数据。"],
        ["筛选条件", json.dumps(snapshot.condition_tree, ensure_ascii=False, sort_keys=True)],
    ]
    for sheet in sheets:
        for row in metadata:
            _row(sheet, row, positions)
        sheet.freeze_panes = "A8"
        sheet.column_dimensions["A"].width = 34
        for column in ["B", "C", "D", "E", "F", "G", "H", "I"]:
            sheet.column_dimensions[column].width = 28
    company_sheet, evidence_sheet, pending_sheet = sheets
    _row(company_sheet, ["企业 ID", "匹配状态", *[LABELS[f] for f in export.fields], "企业版本"], positions)
    _row(evidence_sheet, ["企业 ID", "企业名称", "匹配状态", "条件 ID", "逐项判断", "依据说明", "证据 ID", "来源", "核验日期", "原文定位",
        *[label for _, label in CERTIFICATE_COLUMNS]], positions)
    _row(pending_sheet, ["企业 ID", "匹配状态", *[LABELS[f] for f in export.fields], "待核实原因"], positions)
    for row in selected:
        company = row.get("company", {})
        values = [row["company_id"], STATUS_LABELS[row["status"]], *[company.get(field, "") for field in export.fields]]
        if row["status"] == "confirmed":
            _row(company_sheet, [*values, row.get("company_version", "")], positions)
        else:
            reasons = list(row.get("unresolved", []))
            reasons.extend(c.get("reason", "") for c in row.get("conditions", []) if c.get("verdict") in {"unknown", "conflict"})
            _row(pending_sheet, [*values, "；".join(dict.fromkeys(reasons)) or "仍有条件待核实"], positions)
        for condition in row.get("conditions", []):
            ids = condition.get("evidence_ids", [])
            # Never follow another deployment/company's evidence identifiers.
            evidence = [sources[identifier] for identifier in ids if identifier in sources
                        and sources[identifier].subject.get("company_id") == row["company_id"]]
            records = evidence or [None]
            for source in records:
                certificates = (_frozen_certificates(row, condition, str(source.pk), export.authorization['scope'])
                    if source else [])
                for certificate in certificates or [{}]:
                    values = [certificate.get(field, '') for field, _ in CERTIFICATE_COLUMNS]
                    if certificate:
                        values[9] = {'verified': '已核实', 'unverified': '未核实', 'conflict': '冲突',
                            'superseded': '已被替代'}.get(certificate.get('verification_status'), '未核实')
                    _row(evidence_sheet, [row["company_id"], company.get("name", ""), STATUS_LABELS[row["status"]],
                    condition.get("condition_id", ""), VERDICTS.get(condition.get("verdict"), "未知"),
                    condition.get("reason", ""), str(source.pk) if source else ",".join(ids),
                    _source_label(source) if source else "来源未包含在授权快照证据范围",
                    source.verified_on.isoformat() if source and source.verified_on else "未核实",
                    public_locator(source.locator) if source else {}, *values], positions)
    for sheet in sheets:
        sheet.auto_filter.ref = f"A7:{openpyxl.utils.get_column_letter(sheet.max_column)}{max(7, sheet.max_row)}"
        for cell in sheet[7]:
            cell.font = openpyxl.styles.Font(bold=True)
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="E7EFF8")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = openpyxl.styles.Alignment(vertical="top", wrap_text=True)
    stream = io.BytesIO()
    workbook.save(stream)
    content = stream.getvalue()
    if len(content) > MAX_FILE_BYTES:
        raise BaichuanError("export_size_limit", "导出文件超过 20 MiB，请缩小结果范围。", 413)
    return content, counts


def _publish_file(content):
    key = str(uuid.uuid4())
    path = storage_path(key)
    try:
        with path.open("xb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return key, path


def process_export(job, worker):
    with leased_job(job.pk, worker, job.fence) as current:
        export = BaichuanExport.objects.select_related("snapshot").get(job=current)
        if export.owner_id != current.owner_id or export.authorization.get("scope") != current.scope:
            raise BaichuanError("not_found", "导出不存在或不可访问。", 404)
        # Artifact metadata and job completion commit together below. A crash
        # before that transaction leaves an unreferenced file; a new fence
        # regenerates this same snapshot and can never publish the old writer.
    save_job(job.pk, worker, job.fence, stage="exporting")
    generated_at = timezone.now()
    content, counts = build_workbook(export, generated_at)
    key, path = _publish_file(content)
    try:
        with leased_job(job.pk, worker, job.fence) as current:
            export = BaichuanExport.objects.select_for_update().get(pk=export.pk, job=current)
            ttl = getattr(settings, "BAICHUAN_EXPORT_TTL_SECONDS", 86400)
            if type(ttl) is not int or not 60 <= ttl <= 7 * 86400:
                raise BaichuanError("export_unavailable", "导出保留期限配置无效。", 503)
            export.storage_ref, export.file_hash = key, hashlib.sha256(content).hexdigest()
            export.size_bytes, export.generated_at = len(content), generated_at
            export.expires_at = generated_at + timedelta(seconds=ttl)
            export.save(update_fields=["storage_ref", "file_hash", "size_bytes", "generated_at", "expires_at"])
            save_job(job.pk, worker, job.fence, status="completed", stage="exporting", counts=counts,
                result_ref={"export_id": str(export.pk)}, checkpoint={"exported_count": sum(counts.values()),
                    "condition_revision": export.snapshot.condition_revision})
            audit(current.owner, "baichuan_export_generated", export.pk, changes=[export.file_hash])
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return export


def export_for(user, export_id):
    require_assistant(user)
    try:
        export = BaichuanExport.objects.select_related("job", "snapshot").get(pk=export_id, owner_id=user.pk,
            job__scope=scope(), snapshot__conversation__scope=scope(), snapshot__conversation__owner_id=user.pk)
    except (BaichuanExport.DoesNotExist, ValidationError, ValueError, TypeError):
        raise BaichuanError("not_found", "导出不存在或不可访问。", 404) from None
    job_for(user, export.job_id)
    if export.authorization.get("scope") != scope():
        raise BaichuanError("not_found", "导出不存在或不可访问。", 404)
    return export


def open_export(user, export_id):
    export = export_for(user, export_id)
    if export.expires_at and export.expires_at <= timezone.now():
        raise BaichuanError("export_expired", "导出文件已过期，请重新导出该结果快照。", 410)
    if export.job.status != "completed" or not export.storage_ref or not export.generated_at:
        raise BaichuanError("export_not_ready", "Excel 尚未生成，请稍后重试。", 409)
    path = storage_path(export.storage_ref)
    try:
        stream = path.open("rb")
    except FileNotFoundError:
        raise BaichuanError("export_missing", "导出文件已清理，请重新导出。", 410) from None
    digest, size = hashlib.sha256(), 0
    try:
        while chunk := stream.read(65536):
            size += len(chunk)
            if size > MAX_FILE_BYTES:
                raise BaichuanError("export_changed", "导出文件校验失败，请重新导出。", 409)
            digest.update(chunk)
        if digest.hexdigest() != export.file_hash or size != export.size_bytes:
            raise BaichuanError("export_changed", "导出文件校验失败，请重新导出。", 409)
        # Permission can change while hashing; check it again before download.
        export_for(user, export.pk)
        stream.seek(0)
        audit(user, "baichuan_export_download", export.pk, changes=[export.file_hash])
        return export, stream
    except Exception:
        stream.close()
        raise


def expire_exports(*, limit=100):
    """Bounded maintenance hook; caller schedules it in the authorized runtime."""
    if type(limit) is not int or not 1 <= limit <= 1000:
        raise ValueError("Invalid cleanup limit")
    deleted = 0
    for export in BaichuanExport.objects.filter(expires_at__lte=timezone.now(), job__scope=scope()).exclude(storage_ref="").order_by("expires_at")[:limit]:
        path = storage_path(export.storage_ref)
        if path.is_file():
            path.unlink()
            deleted += 1
        BaichuanExport.objects.filter(pk=export.pk, storage_ref=export.storage_ref,
            expires_at__lte=timezone.now()).update(storage_ref="")
    return deleted

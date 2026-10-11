"""Offline regression tests for actual XLSX parsing and source transparency."""
import importlib
import io
import os
from pathlib import Path
import re
import sys
import tempfile
from unittest.mock import patch
import zipfile

from django.test import SimpleTestCase, override_settings
from openpyxl import Workbook


ENGINEERING = Path(__file__).resolve().parents[4] / "engineering"


class EngineeringEngineTests(SimpleTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.paths = [str(ENGINEERING), str(ENGINEERING / "tools")]
        sys.path[:0] = cls.paths
        cls.cli = importlib.import_module("cost_skill")
        cls.health = importlib.import_module("engine.health")

    @classmethod
    def tearDownClass(cls):
        for path in cls.paths:
            sys.path.remove(path)
        super().tearDownClass()

    def dimensionless(self, path, last_row=2):
        book = Workbook()
        sheet = book.active
        sheet.append(["序号", "项目名称", "单位", "数量"])
        sheet.append([1, "配电箱", "台", 2])
        if last_row > 2:
            sheet.cell(last_row, 1, "超限")
        raw = io.BytesIO(); book.save(raw); book.close()
        with zipfile.ZipFile(raw) as original, zipfile.ZipFile(path, "w") as target:
            for entry in original.infolist():
                contents = original.read(entry.filename)
                if entry.filename.startswith("xl/worksheets/"):
                    contents = re.sub(rb'<dimension\b[^>]*/>', b'', contents)
                target.writestr(entry, contents)

    def test_valid_workbook_without_dimension_is_inspected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "valid.xlsx"
            self.dimensionless(path)
            result = self.cli.inspect_xlsx(path)
        self.assertTrue(result["passed"], result)
        self.assertEqual(result["data_rows"], 1)

    def test_dimensionless_workbook_still_enforces_row_cap(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "oversized.xlsx"
            self.dimensionless(path, self.cli.MAX_SCAN_ROWS + 20)
            result = self.cli.inspect_xlsx(path)
        self.assertFalse(result["passed"])
        self.assertIn("ROW_LIMIT_EXCEEDED", [item["code"] for item in result["issues"]])

    def test_nonzero_estimates_do_not_claim_normal_source_health(self):
        result = self.health.audit([[1, "设备", "", "台", 1, "保守暂估（非网价）", 800, 800]])
        self.assertEqual(result["结论"], "可疑")
        self.assertEqual(result["暂估项"], 1)
        self.assertIn("询价", result["警告"][0])

    def test_labor_estimates_are_also_disclosed(self):
        result = self.health.audit([], [[1, "劳务", "", "项", 1, "品类经验暂估", 100, 100]])
        self.assertEqual(result["暂估项"], 1)
        self.assertTrue(result["警告"])
        self.assertEqual(result["行数"], 1)

    def test_source_notes_keep_height_for_explicit_lines(self):
        from engine import assemble
        book = Workbook(); sheet = book.active
        assemble._notes_block(sheet, 4, ["来源健康度\n仍须询价"], 12)
        before = sheet.row_dimensions[5].height
        self.assertGreaterEqual(before, 34)
        assemble._style(sheet, assemble.WIDTHS, wrap={2, 3, 8, 9, 10, 11, 12})
        self.assertGreaterEqual(sheet.row_dimensions[5].height, before)
        book.close()

    def test_owner_material_zero_is_not_an_unpriced_item(self):
        result = self.health.audit([[1, "甲供", "", "台", 1, "甲方供料", 0, 0]])
        self.assertEqual(result["零价项"], 0)
        self.assertFalse(result["警告"])

    def test_operator_data_root_is_absolute_and_keeps_rules_in_code(self):
        from engine.paths import data_path
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"ENGINEERING_DATA_ROOT": directory}):
            self.assertEqual(Path(data_path("library", "prices.xlsx")), Path(directory) / "library" / "prices.xlsx")
        with patch.dict(os.environ, {"ENGINEERING_DATA_ROOT": "relative/data"}):
            with self.assertRaises(ValueError):
                data_path("library")

    def test_worker_forwards_only_operator_configured_data_directory(self):
        import subprocess
        from portal.engineering_worker import _invoke, WorkerUnavailable
        completed = subprocess.CompletedProcess([], 0, stdout='{"ok":true}')
        with tempfile.TemporaryDirectory() as directory, override_settings(ENGINEERING_DATA_ROOT=directory):
            with patch("portal.engineering_worker.subprocess.run", return_value=completed) as invoked:
                _invoke({"python": Path(sys.executable), "cli": ENGINEERING / "tools" / "cost_skill.py"}, "inspect", [])
            self.assertEqual(invoked.call_args.kwargs["env"]["ENGINEERING_DATA_ROOT"], str(Path(directory).resolve()))
        with override_settings(ENGINEERING_DATA_ROOT="relative/data"):
            with self.assertRaises(WorkerUnavailable):
                _invoke({"python": Path(sys.executable), "cli": ENGINEERING / "tools" / "cost_skill.py"}, "inspect", [])

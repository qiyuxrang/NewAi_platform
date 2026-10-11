import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import SimpleTestCase
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from portal.product_assets.simple_presentation import add_visual_panel, cover


class PresentationImagePlacementTests(SimpleTestCase):
    def assert_contained(self, picture, image_path, size, box):
        left, top, width, height = map(Inches, box)
        self.assertGreaterEqual(picture.left, left)
        self.assertGreaterEqual(picture.top, top)
        self.assertLessEqual(picture.left + picture.width, left + width)
        self.assertLessEqual(picture.top + picture.height, top + height)
        self.assertLessEqual(abs(2 * picture.left + picture.width - (2 * left + width)), 1)
        self.assertLessEqual(abs(2 * picture.top + picture.height - (2 * top + height)), 1)
        self.assertLessEqual(abs(picture.width * size[1] - picture.height * size[0]), max(size))
        self.assertEqual((picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom), (0, 0, 0, 0))
        self.assertEqual(hashlib.sha256(picture.image.blob).hexdigest(), hashlib.sha256(image_path.read_bytes()).hexdigest())

    def test_cover_keeps_wide_portrait_and_square_images_uncropped(self):
        with TemporaryDirectory() as directory:
            for size in ((2400, 1350), (600, 1200), (800, 800)):
                with self.subTest(size=size):
                    image_path = Path(directory) / "source.png"
                    # Unequal DPI metadata must not stretch pixel geometry.
                    Image.new("RGB", size, "#2470AA").save(image_path, dpi=(96, 144))
                    prs = Presentation()
                    slide = cover(prs, {"title": "Synthetic image geometry"},
                                  {"asset_path": str(image_path), "ref": "technical-solution:F1", "caption": "Source figure"})
                    picture = next(shape for shape in slide.shapes
                                   if hasattr(shape, "image") and shape.image.blob == image_path.read_bytes())
                    self.assert_contained(picture, image_path, size, (8.75, 1.72, 3.85, 3.95))
                    self.assertLess(picture.top + picture.height, Inches(5.74))
                    self.assertGreater(picture.left, Inches(0.72 + 6.9))

    def test_report_panels_preserve_ratio_without_covering_titles(self):
        with TemporaryDirectory() as directory:
            for box in ((0.7, 1.25, 7.7, 5.35), (0.82, 4.12, 5.1, 2.18), (6.1, 1.25, 6.1, 4.55)):
                for size in ((2400, 1350), (600, 1200), (800, 800)):
                    with self.subTest(box=box, size=size):
                        image_path = Path(directory) / "source.png"
                        Image.new("RGB", size, "#2470AA").save(image_path, dpi=(96, 144))
                        prs = Presentation()
                        slide = prs.slides.add_slide(prs.slide_layouts[6])
                        add_visual_panel(slide, {"asset_path": str(image_path), "caption": "Source figure"}, *box)
                        picture = slide.shapes[1]
                        x, y, width, height = box
                        self.assert_contained(picture, image_path, size, (x + 0.18, y + 0.42, width - 0.36, height - 0.72))
                        title = slide.shapes[2]
                        self.assertLess(title.top + title.height, picture.top)

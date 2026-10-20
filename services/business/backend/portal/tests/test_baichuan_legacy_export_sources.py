"""Trusted old-catalog matches retain existing source links, not new authority."""
from copy import deepcopy
import io
from pathlib import Path
import tempfile
from unittest.mock import patch

from django.test import override_settings
from django.utils import timezone

from portal.baichuan_assistant_models import (BaichuanCertificate, BaichuanConversation,
    BaichuanEvidence, BaichuanFactRevision, map_legacy_qualifications)
from portal.baichuan_exports import enqueue_export, open_export, process_export
from portal.baichuan_jobs import claim_job, enqueue_job
from portal.baichuan_legacy_evidence import backfill_legacy_source_evidence
from portal.baichuan_matching import _legacy_source_ids, match_companies, save_match_snapshot
from portal.baichuan_models import Company, CompanyQualification
from portal.baichuan_rules import install_compatibility_rules
from .base import PortalTestCase


class LegacyExportSourceTests(PortalTestCase):
    def setUp(self):
        self.owner = self.create_user('legacy-export-source-reader', 'product')
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        config = override_settings(BAICHUAN_ASSISTANT_SCOPE='legacy-export-source-test',
            BAICHUAN_TRUST_LEGACY_QUALIFICATIONS=True, BAICHUAN_ASSISTANT_ENABLED=True,
            BAICHUAN_STORAGE_ROOT=Path(self.temp.name), PORTAL_BAICHUAN_SCOPE_CONFIRMED=True,
            PORTAL_BAICHUAN_READ_USER_IDS=[self.owner.pk], PORTAL_BAICHUAN_WRITE_USER_IDS=[self.owner.pk],
            PORTAL_BAICHUAN_READ_ROLE_CODES=[], PORTAL_BAICHUAN_WRITE_ROLE_CODES=[])
        config.enable()
        self.addCleanup(config.disable)
        for name in ('socket.create_connection', 'socket.getaddrinfo', 'portal.model_gateway.generate_for_use'):
            guard = patch(name, side_effect=AssertionError('source tracing never calls models/network'))
            guard.start()
            self.addCleanup(guard.stop)
        install_compatibility_rules(self.owner)
        self.company = Company.objects.create(name='合成旧库来源企业', phone='0012345',
            created_by=self.owner, updated_by=self.owner)
        self.legacy = CompanyQualification.objects.create(company=self.company, category='electronic',
            grade='2', status='PRESENT', certificate_no='SYNTHETIC-001', source='合成源表第9行声明')
        map_legacy_qualifications(self.company, self.owner)
        self.conditions = dict(schema_version=1, revision=1, must=dict(field='qualification',
            category='electronic', value='2', op='grade_exact', holder='bidder'),
            output=dict(fields=['name', 'phone'], limit=20))

    def linked(self):
        backfill_legacy_source_evidence(self.company, self.owner)
        return BaichuanCertificate.objects.select_related('evidence', 'source_revision__supersedes').get(
            legacy_qualification_id=self.legacy.pk)

    def result(self):
        return match_companies(self.owner, self.conditions)[2][0]

    def snapshot(self, key):
        conv = BaichuanConversation.objects.create(owner=self.owner, scope='legacy-export-source-test',
            condition_revision=1)
        job, _ = enqueue_job(self.owner, kind='query', payload={'condition_revision': 1},
            conversation=conv, idempotency_key=key)
        claimed = claim_job('source-test-worker', job_id=job.pk)
        snapshot = save_match_snapshot(self.owner, job.pk, 'source-test-worker', claimed.fence, self.conditions)
        job.refresh_from_db()
        job.status = 'completed'
        job.save(update_fields=['status'])
        return snapshot

    def workbook(self, snapshot, key):
        job = enqueue_export(self.owner, snapshot.pk, {'client_request_id': key})
        export = process_export(claim_job('source-export-worker', job_id=job.pk), 'source-export-worker')
        _, stream = open_export(self.owner, export.pk)
        try:
            data = stream.read()
        finally:
            stream.close()
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(data), data_only=False)

    def test_new_snapshot_and_actual_xlsx_retain_chain_without_mutating_old_snapshot_or_facts(self):
        old = self.snapshot('old-empty-source-snapshot')
        old_results = deepcopy(old.results)
        certificate = self.linked()
        protected = (list(Company.objects.values()), list(CompanyQualification.objects.values()),
            list(BaichuanEvidence.objects.values()), list(BaichuanFactRevision.objects.values()),
            list(BaichuanCertificate.objects.values()))
        current = self.snapshot('new-linked-source-snapshot')
        self.assertEqual(current.counts, dict(confirmed=1, candidate=0, excluded=0))
        condition = current.results[0]['conditions'][0]
        self.assertEqual(condition['evidence_ids'], [str(certificate.evidence_id)])
        self.assertIn('用户确认的旧资质库', condition['reason'])
        self.assertEqual(current.results[0]['certificate_records'][0]['verification_status'], 'unverified')
        self.assertEqual(current.results[0]['certificate_records'][0]['evidence_access_status'], 'unverified')
        book = self.workbook(current, 'new-source-export')
        evidence = book['逐项依据']
        values = dict(zip([cell.value for cell in evidence[7]], [cell.value for cell in evidence[8]]))
        self.assertEqual(values['证据 ID'], str(certificate.evidence_id))
        self.assertEqual(values['来源'], certificate.evidence.title)
        self.assertIn(str(self.legacy.pk), values['原文定位'])
        self.assertEqual(values['核验日期'], '未核实')
        self.assertEqual(values['证书核实状态'], '未核实')
        self.assertEqual(values['证书编号'], 'SYNTHETIC-001')
        self.assertEqual(values['证书主体类型'], 'bidder')
        self.assertEqual(book['企业结果'].cell(8, 4).value, '0012345')
        # Exporting the original immutable snapshot cannot acquire today's links.
        old_book = self.workbook(old, 'old-source-export')
        self.assertEqual(old_book['逐项依据'].cell(8, 8).value, '来源未包含在授权快照证据范围')
        old.refresh_from_db()
        self.assertEqual(old.results, old_results)
        self.assertEqual(protected, (list(Company.objects.values()), list(CompanyQualification.objects.values()),
            list(BaichuanEvidence.objects.values()), list(BaichuanFactRevision.objects.values()),
            list(BaichuanCertificate.objects.values())))

    def test_missing_source_or_unlinked_mapping_does_not_invent_an_evidence_id(self):
        self.assertEqual(self.result()['status'], 'confirmed')
        self.assertEqual(self.result()['conditions'][0]['evidence_ids'], [])
        self.legacy.source = ''
        self.legacy.save(update_fields=['source'])
        self.assertEqual(self.result()['conditions'][0]['evidence_ids'], [])

    def test_source_or_grade_changed_after_mapping_is_not_cited_but_old_catalog_stays_trusted(self):
        self.linked()
        for source, grade in [('人工后来修改声明', '2'), ('合成源表第9行声明', '1')]:
            with self.subTest(source=source, grade=grade):
                self.legacy.source, self.legacy.grade = source, grade
                self.legacy.save(update_fields=['source', 'grade'])
                self.conditions['must']['value'] = grade
                row = self.result()
                self.assertEqual(row['status'], 'confirmed')
                self.assertEqual(row['conditions'][0]['evidence_ids'], [])

    def test_company_scope_identity_and_source_tampering_fail_closed(self):
        certificate = self.linked()
        evidence = certificate.evidence
        cases = [('scope', 'other-deployment'), ('source_type', 'web'), ('access_status', 'verified'),
            ('verified_on', timezone.localdate()), ('excerpt', 'altered source'), ('content_hash', '0' * 64),
            ('locator', {}), ('subject', []), ('subject', {**evidence.subject, 'company_id': 'other-company'}),
            ('subject', {**evidence.subject, 'identity_verified': True}),
            ('subject', {**evidence.subject, 'legacy_qualification_id': 'other-old-row'}),
            ('subject', {**evidence.subject, 'source_revision_id': 'other-revision'}),
            ('subject', {**evidence.subject, 'holder_type': 'manufacturer'})]
        for field, changed in cases:
            with self.subTest(field=field, changed=changed):
                original = getattr(evidence, field)
                setattr(evidence, field, changed)
                self.assertEqual(_legacy_source_ids(self.company, self.legacy, [certificate]), [])
                setattr(evidence, field, original)
        self.assertEqual(_legacy_source_ids(self.company, self.legacy, [certificate]), [str(evidence.pk)])

    def test_certificate_and_revision_chain_tampering_fail_closed(self):
        certificate = self.linked()
        link = certificate.source_revision
        original = link.supersedes
        cases = [(certificate, 'holder_type', 'manufacturer'), (certificate, 'holder_company_id', None),
            (certificate, 'company_id', None), (certificate, 'category', 'itss'),
            (certificate, 'grade', '1'), (certificate, 'certificate_no', 'OTHER'),
            (certificate, 'verification_status', 'verified'), (certificate, 'verified_at', timezone.now()),
            (certificate, 'legacy_qualification_id', None), (link, 'field', 'public_observation'),
            (link, 'review_status', 'applied'), (link, 'verification_status', 'verified'),
            (link, 'company_id', None), (link, 'object_id', None), (link, 'new_value', {}),
            (original, 'field', 'company.name'), (original, 'object_id', None),
            (original, 'company_id', None), (original, 'review_status', 'applied'),
            (original, 'new_value', {**original.new_value, 'source': 'wrong source'})]
        for obj, field, changed in cases:
            with self.subTest(model=type(obj).__name__, field=field):
                previous = getattr(obj, field)
                setattr(obj, field, changed)
                self.assertEqual(_legacy_source_ids(self.company, self.legacy, [certificate]), [])
                setattr(obj, field, previous)
        self.assertEqual(_legacy_source_ids(self.company, self.legacy, [certificate]), [str(certificate.evidence_id)])

    def test_new_unverified_public_observation_cannot_inherit_old_source_authority(self):
        self.linked()
        other = Company.objects.create(name='合成公开观察企业', created_by=self.owner, updated_by=self.owner)
        BaichuanCertificate.objects.create(company=other, holder_type='bidder', holder_company_id=other.pk,
            category='electronic', grade='2', assertion_status='PRESENT', verification_status='unverified')
        row = next(row for row in match_companies(self.owner, self.conditions)[2] if row['company_id'] == str(other.pk))
        self.assertEqual(row['status'], 'candidate')
        self.assertEqual(row['conditions'][0]['evidence_ids'], [])

    def test_strict_mode_and_nonmatching_grade_do_not_receive_satisfied_legacy_verdict(self):
        self.linked()
        with override_settings(BAICHUAN_TRUST_LEGACY_QUALIFICATIONS=False):
            self.assertEqual(self.result()['status'], 'candidate')
        self.conditions['must']['value'] = '1'
        row = self.result()
        self.assertEqual(row['status'], 'candidate')
        self.assertNotEqual(row['conditions'][0]['verdict'], 'satisfied')

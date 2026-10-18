"""User-requested generation-only policy, including real isolated Office files."""
import hashlib
import json
import os
from pathlib import Path
from unittest.mock import patch
from xml.sax.saxutils import escape
from zipfile import ZipFile

from django.test import SimpleTestCase, override_settings

from portal.product_delivery import artifact_delivery
from portal.product_diagrams import validate_diagram_batch
from portal.product_documents import DocumentError
from portal.product_generation_policy import POLICY
from portal.product_models import DocumentTask
from portal.product_ppt_master import DESIGN, validate_selection, validate_svg
from portal.product_service import ProductError
from portal.product_worker import run_once
from .base import PortalTestCase
from . import test_product_short_chain as short_fixture


class GenerationOnlyContractTests(SimpleTestCase):
    def test_diagram_semantics_are_skipped_but_unsafe_sources_and_unknown_refs_fail(self):
        figure = {"id": "F1"}
        paragraph = {"id": "B1", "text": "项目原有文字"}
        item = {"figure_id": "F1", "title": "示意", "description": "未作核验",
                "source": 'flowchart LR\nA["编制方案"] --> B["后续实施"]',
                "anchor_text": "不是逐字原文", "evidence_block_ids": ["B1"]}
        self.assertEqual(validate_diagram_batch({"diagrams": [item]}, [figure], [paragraph],
                         verify_content=False), [item])
        for changed in ({**item, "source": 'flowchart LR\nclick A "https://example.com"'},
                        {**item, "evidence_block_ids": ["OTHER"]}):
            with self.assertRaises(ValueError):
                validate_diagram_batch({"diagrams": [changed]}, [figure], [paragraph], verify_content=False)

    def test_ppt_skips_quote_and_font_review_but_rejects_external_resources(self):
        page = {"page": 2, "type": "content", "title": "项目概述",
                "evidence": [{"ref": "technical-solution:B1", "quote": "已保存正文"}]}
        selected = {"pages": [{"page": 2, "items": [{"ref": "technical-solution:B1", "quote": "编写摘要"}], "chart": None}]}
        self.assertEqual(validate_selection(selected, [page], verify_content=False), selected)
        svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><text x="48" y="80" font-size="12">编写摘要</text></svg>'
        self.assertEqual(validate_svg({"svg": svg}, page | {"items": ["编写摘要"]}, DESIGN,
                                    verify_content=False), {"svg": svg})
        for addition in ('<script>alert(1)</script>', '<image href="https://example.com/a.png"/>'):
            with self.assertRaises(ProductError):
                validate_svg({"svg": svg.replace('</svg>', addition + '</svg>')}, page | {"items": []}, DESIGN,
                             verify_content=False)


class GenerationOnlyRecordTests(PortalTestCase):
    setUp = short_fixture.ShortTriadRealChainTests.setUp
    create_task = short_fixture.ShortTriadRealChainTests.create_task
    save_blueprint = short_fixture.ShortTriadRealChainTests.save_blueprint
    approve_blueprint = short_fixture.ShortTriadRealChainTests.approve_blueprint
    input_payload = short_fixture.ShortTriadRealChainTests.input_payload
    blueprint_payload = short_fixture.ShortTriadRealChainTests.blueprint_payload
    paragraph = staticmethod(short_fixture.ShortTriadRealChainTests.paragraph)

    @override_settings(PRODUCT_CONTENT_VERIFICATION_ENABLED=False, PRODUCT_MODEL_CALLS_ALLOWED=True)
    def test_short_unverified_conflicting_reports_feed_ppt_without_rewriting_old_review(self):
        from portal.product_models import DocumentArtifact
        from portal.product_pair import _report_generation_hash, pair_snapshot, save_report_content
        from portal.product_service import append_revision
        from portal.product_worker import _ensure_content_review, claim_task
        created = self.approve_blueprint(self.save_blueprint(self.create_task()))["task"]
        task = DocumentTask.objects.get(pk=created["id"])
        _, fence, attempt = claim_task()
        task.refresh_from_db()
        current = task.revisions.get(kind="input", version=task.input_version)
        blueprint = task.revisions.get(kind="blueprint", version=task.blueprint_version)
        old = append_revision(task, "review", {"passed": False, "issues": [{"code": "old_rejected_fact"}]},
                              input_hash=current.sha256, blueprint_hash=blueprint.sha256, family="technical-solution")
        old_payload, old_sha = dict(old.payload), old.sha256
        artifacts = []
        with patch("portal.product_worker.content_checks", side_effect=AssertionError("no checks")), \
             patch("portal.product_worker.generate_for_use", side_effect=AssertionError("no review model")):
            for index, family in enumerate(("technical-solution", "feasibility"), 1):
                # Deliberately short and inconsistent with the two input devices.
                chapter = append_revision(task, "chapter", {"chapter_id": "overview", "title": "合成机制与验证边界",
                    "paragraphs": ["合成防火墙99台。"], "source_ids": ["1"]}, family=family,
                    input_hash=current.sha256, blueprint_hash=blueprint.sha256)
                review = _ensure_content_review(task, fence, attempt, current, blueprint, [chapter], family)
                content = {"blocks": [{"id": "BODY", "type": "paragraph", "text": "合成防火墙99台。"}], "pending": []}
                report = save_report_content(task, family, content, input_hash=current.sha256,
                                              blueprint_hash=blueprint.sha256, actor=self.owner)
                target = Path(self.storage.name) / f"{family}.docx"
                from docx import Document
                document = Document()
                document.add_paragraph("合成防火墙99台。")
                document.save(target)
                artifact = DocumentArtifact.objects.create(task=task, family=family, version=index,
                    path=target.name, sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
                    review=review, input_hash=current.sha256, blueprint_hash=blueprint.sha256,
                    generation_hash=_report_generation_hash(task, family, current, blueprint), template_hash="a"*64,
                    render_evidence={"generation_policy": POLICY, "report_id": str(report.pk),
                        "diagram_generation": {"pending_count": 1, "quality_issue_count": 1}})
                self.assertTrue(artifact_delivery(artifact)["delivered"])
                self.assertEqual(artifact_delivery(artifact)["quality_scope"], "not_run")
                artifacts.append(artifact)
            pair = pair_snapshot(task)
            self.assertEqual(pair["generation_policy"], POLICY)
            self.assertEqual(len(pair["sources"]), 2)
        old.refresh_from_db()
        self.assertEqual((old.payload, old.sha256), (old_payload, old_sha))
        from portal.product_models import DocumentSource
        source = DocumentSource.objects.create(task=task, original_name="unbound.txt", media_type="text/plain",
            path="unbound.txt", sha256="b"*64, size=1, parsed={}, uploaded_by=self.owner)
        self.assertFalse(artifact_delivery(artifacts[0])["delivered"], "changed source snapshot must invalidate output")
        source.delete()
        self.assertTrue(artifact_delivery(artifacts[0])["delivered"])
        approval = blueprint.approvals.filter(actor=self.owner).latest("created_at")
        approval.decision = "reject"
        approval.save(update_fields=["decision"])
        self.assertFalse(artifact_delivery(artifacts[0])["delivered"])


@override_settings(PRODUCT_CONTENT_VERIFICATION_ENABLED=False, PRODUCT_MODEL_CALLS_ALLOWED=True,
                   PRODUCT_TECHNICAL_TARGET_CHARACTERS=50000, PRODUCT_FEASIBILITY_TARGET_CHARACTERS=70000,
                   PRODUCT_MAX_MODEL_CALLS=120, PRODUCT_STEP_RETRY_BASE_SECONDS=0,
                   PRODUCT_DOCUMENT_RENDER_TIMEOUT_SECONDS=180, PRODUCT_OFFICE_RENDER_TIMEOUT_SECONDS=120)
class GenerationOnlyRealChainTests(PortalTestCase):
    def setUp(self):
        short_fixture.ShortTriadRealChainTests.setUp(self)
        enabled = override_settings(PRODUCT_MODEL_CALLS_ALLOWED=True)
        enabled.enable()
        self.addCleanup(enabled.disable)
    create_task = short_fixture.ShortTriadRealChainTests.create_task
    save_blueprint = short_fixture.ShortTriadRealChainTests.save_blueprint
    approve_blueprint = short_fixture.ShortTriadRealChainTests.approve_blueprint
    input_payload = short_fixture.ShortTriadRealChainTests.input_payload
    blueprint_payload = short_fixture.ShortTriadRealChainTests.blueprint_payload
    paragraph = staticmethod(short_fixture.ShortTriadRealChainTests.paragraph)

    def model_reply(self, *args, **kwargs):
        request = json.loads(args[2][1]["content"])
        action = request["action"]
        if action in {"chapter", "diagram"}:
            return short_fixture.ShortTriadRealChainTests.model_reply(self, *args)
        self.model_actions.append(action)
        if action == "presentation_plan":
            value = {"pages": [{"page": page["page"], "items": page["evidence"][:2], "chart": None}
                                for page in request["pages"]]}
        elif action == "presentation_slide":
            page = request["page"]
            nodes = [f'<text x="48" y="80" font-size="40" font-family="Microsoft YaHei">{escape(page["title"])}</text>']
            for index, item in enumerate(page["items"]):
                # Short synthetic fixture for real native export, not visual acceptance.
                lines = [item[start:start+35] for start in range(0, len(item), 35)]
                spans = ''.join(f'<tspan x="48" dy="32">{escape(line)}</tspan>' for line in lines)
                nodes.append(f'<text x="48" y="{140+index*220}" font-size="28" font-family="Microsoft YaHei">{spans}</text>')
            value = {"svg": '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">'+''.join(nodes)+'</svg>'}
        else:
            raise AssertionError("unexpected review/correction: " + action)
        return {"content": json.dumps(value, ensure_ascii=False), "prompt_tokens": 10, "completion_tokens": 20}

    def test_two_words_editable_ppt_preview_download_and_resume_without_content_review(self):
        if os.environ.get("RUN_PRODUCT_GENERATION_ONLY_SMOKE") != "1":
            self.skipTest("opt-in real Office/Node renderer; run scripts/test-product-generation-only.py")
        self.model_actions = []
        import portal.product_presentation as presentation
        real = presentation.render_presentation_draft
        failed = False
        def once(*args, **kwargs):
            nonlocal failed
            if not failed:
                failed = True
                raise DocumentError("presentation_unavailable", diagnostic={"stage": "presentation_render", "reason": "timeout"})
            return real(*args, **kwargs)
        with patch("portal.product_worker.generate_for_use", side_effect=self.model_reply), \
             patch("portal.product_worker.content_checks", side_effect=AssertionError("content checks must not run")), \
             patch("portal.product_worker._content_quality_issues", side_effect=AssertionError("quality checks must not run")), \
             patch("portal.product_writing_revision.ensure_writing_consistency", side_effect=AssertionError("consistency review must not run")), \
             patch("portal.product_presentation.render_presentation_draft", side_effect=once):
            created = self.approve_blueprint(self.save_blueprint(self.create_task()))["task"]
            task = DocumentTask.objects.get(pk=created["id"])
            self.assertTrue(run_once())
            task.refresh_from_db()
            self.assertEqual((task.state, task.error_code), ("QUEUED", "step_retry_pending"), task.checkpoint.get("last_failure"))
            saved = {a.family: (str(a.pk), a.sha256) for a in task.artifacts.all()}
            self.assertEqual(set(saved), {"technical-solution", "feasibility"})
            chapter_calls = self.model_actions.count("chapter")
            self.assertTrue(run_once())
            task.refresh_from_db()
            self.assertEqual((task.state, task.error_code), ("COMPLETED", ""), task.checkpoint.get("last_failure"))
            self.assertEqual(self.model_actions.count("chapter"), chapter_calls)
            self.assertEqual({a.family: (str(a.pk), a.sha256) for a in task.artifacts.exclude(family="presentation")}, saved)
        self.assertNotIn("independent_review", self.model_actions)
        self.assertEqual(task.artifacts.count(), 3)
        self.assertEqual(task.attempts.count(), 2)
        for review in task.revisions.filter(kind="review"):
            self.assertEqual(review.payload["generation_policy"], POLICY)
            self.assertEqual(review.payload["content_checks"], "not_run")
            self.assertFalse(review.payload["passed"])
        records = []
        from portal.product_output_preview import draft_preview_evidence
        for artifact in task.artifacts.all():
            self.assertTrue(artifact_delivery(artifact)["delivered"])
            path = Path(self.storage.name) / artifact.path
            with ZipFile(path) as package:
                self.assertIsNone(package.testzip())
                if artifact.family == "presentation":
                    from pptx import Presentation
                    deck = Presentation(path)
                    self.assertGreaterEqual(len(deck.slides), 17)
                    self.assertTrue(any(shape.has_text_frame for slide in deck.slides for shape in slide.shapes))
                    self.assertEqual(artifact.render_evidence["quality_gate"]["status"], "not_run")
                else:
                    from docx import Document
                    self.assertTrue(Document(path).paragraphs)
                    self.assertFalse(artifact.review.payload["program_checks_complete"])
            response = self.owner_client.get(f"/api/product/outputs/{artifact.pk}/download/")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(hashlib.sha256(b''.join(response.streaming_content)).hexdigest(), artifact.sha256)
            response.close()
            self.assertEqual(self.other_client.get(f"/api/product/outputs/{artifact.pk}/download/").status_code, 404)
            preview = draft_preview_evidence(artifact)
            self.assertTrue(preview and preview["pages"], artifact.family)
            page = preview["pages"][0]
            response = self.owner_client.get(f"/api/product/outputs/{artifact.pk}/preview/?page={page['page']}")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(hashlib.sha256(b''.join(response.streaming_content)).hexdigest(), page["sha256"])
            response.close()
            records.append({"family": artifact.family, "sha256": artifact.sha256, "pages": preview["page_count"]})
            raw = path.read_bytes()
            path.write_bytes(b"damaged")
            self.assertFalse(artifact_delivery(artifact)["delivered"])
            path.write_bytes(raw)
        destination = os.environ.get("PRODUCT_GENERATION_ONLY_OUTPUT")
        if destination:
            import shutil
            shutil.copytree(self.storage.name, Path(destination) / "artifacts", dirs_exist_ok=True)
            (Path(destination) / "evidence.json").write_text(json.dumps({"policy": POLICY,
                "scope": "synthetic_short_body_real_word_ppt_office_render", "state": task.state,
                "model": "deterministic_substitute", "model_actions": self.model_actions,
                "targets": {"technical": 50000, "feasibility": 70000}, "content_checks": "not_run",
                "artifacts": records, "execution_attempts": task.attempts.count(),
                "report_ids_hashes_reused": True}, ensure_ascii=False, indent=2), encoding="utf-8")

"""The user's existing verified catalog is a source, not a new web claim."""
import io
from pathlib import Path
import tempfile
from datetime import timedelta
from unittest.mock import patch

from django.test import override_settings
from django.utils import timezone

from portal.baichuan_assistant_models import BaichuanCertificate, BaichuanConversation, BaichuanEvidence
from portal.baichuan_exports import enqueue_export, open_export, process_export
from portal.baichuan_jobs import claim_job, enqueue_job
from portal.baichuan_matching import save_match_snapshot
from portal.baichuan_matching import match_companies
from portal.baichuan_models import Company, CompanyQualification
from portal.baichuan_rules import install_compatibility_rules
from .base import PortalTestCase


def q(category='electronic', value='2', **extra):
    return dict(field='qualification', category=category, value=value, op='grade_exact', holder='bidder', **extra)


class TrustedLegacyTests(PortalTestCase):
    def setUp(self):
        self.owner = self.create_user('trusted-legacy-owner', 'product')
        self.settings = override_settings(BAICHUAN_ASSISTANT_SCOPE='trusted-legacy-test',
            BAICHUAN_TRUST_LEGACY_QUALIFICATIONS=True, PORTAL_BAICHUAN_SCOPE_CONFIRMED=True,
            PORTAL_BAICHUAN_READ_USER_IDS=[self.owner.pk], PORTAL_BAICHUAN_WRITE_USER_IDS=[self.owner.pk],
            PORTAL_BAICHUAN_READ_ROLE_CODES=[], PORTAL_BAICHUAN_WRITE_ROLE_CODES=[])
        self.settings.enable(); self.addCleanup(self.settings.disable)
        for target in ('socket.create_connection', 'socket.getaddrinfo', 'portal.model_gateway.generate_for_use'):
            guard = patch(target, side_effect=AssertionError('catalog matching must not call network/models'))
            guard.start(); self.addCleanup(guard.stop)
        install_compatibility_rules(self.owner)
        self.yes = self.company('Z精确二级', province='陕西省', city='西安市')
        self.no = self.company('A其他等级', province='广东省', city='广州市')
        self.empty = self.company('B未录')
        self.old = CompanyQualification.objects.create(company=self.yes, category='electronic', status='PRESENT', grade='2')
        CompanyQualification.objects.create(company=self.no, category='electronic', status='PRESENT', grade='1')

    def company(self, name, **fields):
        return Company.objects.create(name=name, created_by=self.owner, updated_by=self.owner, **fields)

    def match(self, node=None, **extra):
        return match_companies(self.owner, dict(schema_version=1, revision=1, must=node or q(), **extra))

    def rows(self, node=None, **extra):
        return {r['company_id']: r for r in self.match(node, **extra)[2]}

    def public_certificate(self, **extra):
        args = dict(company=self.empty, holder_company_id=self.empty.pk, holder_type='bidder',
            category='electronic', grade='2', assertion_status='PRESENT', verification_status='unverified')
        args.update(extra)
        return BaichuanCertificate.objects.create(**args)

    def test_trusted_exact_returns_original_matching_company_first_without_updating_rows(self):
        before = list(CompanyQualification.objects.values())
        _, _, results, counts = self.match()
        self.assertEqual(counts, dict(confirmed=1, candidate=2, excluded=0))
        self.assertEqual(results[0]['company_id'], str(self.yes.pk))
        self.assertEqual(results[0]['status'], 'confirmed')
        self.assertIn('用户确认的旧资质库', results[0]['conditions'][0]['reason'])
        self.assertIn(str(self.old.pk), results[0]['conditions'][0]['reason'])
        self.assertEqual(before, list(CompanyQualification.objects.values()))
        self.assertEqual(set(r['company_id'] for r in results), {str(c.pk) for c in (self.yes,self.no,self.empty)})

    @override_settings(BAICHUAN_TRUST_LEGACY_QUALIFICATIONS=False)
    def test_flag_off_keeps_prior_strict_unknown_behavior(self):
        self.assertEqual(self.match()[3], dict(confirmed=0, candidate=3, excluded=0))

    def test_new_public_observations_and_missing_or_other_grades_do_not_inherit_trust(self):
        cert = self.public_certificate()
        self.assertEqual(self.rows()[str(self.empty.pk)]['status'], 'candidate')
        self.assertEqual(self.rows()[str(self.no.pk)]['status'], 'candidate')
        cert.refresh_from_db()
        self.assertEqual(cert.verification_status, 'unverified')
        self.assertIsNone(cert.verified_at)

    def test_manufacturer_person_scope_and_unproven_grade_relation_stay_unknown(self):
        for holder in ('manufacturer', 'person'):
            node=q();node['holder']=holder
            self.assertEqual(self.rows(node)[str(self.yes.pk)]['status'], 'candidate')
        self.assertEqual(self.rows(q(scope='指定范围'))[str(self.yes.pk)]['status'], 'candidate')
        node=q();node['op']='grade_at_least'
        self.assertEqual(self.rows(node)[str(self.yes.pk)]['status'], 'candidate')

    def test_given_dates_do_not_invent_certificate_validity(self):
        for offset in (-1, 1):
            as_of=dict(date=(timezone.localdate()+timedelta(days=offset)).isoformat(),basis='tender_date')
            self.assertEqual(self.rows(as_of=as_of)[str(self.yes.pk)]['status'], 'candidate')

    def test_current_company_name_and_region_filter_and_missing_field_is_unknown(self):
        node=dict(all=[q(),dict(field='company.city',op='eq',value='西安市')])
        rows=self.rows(node)
        self.assertEqual(rows[str(self.yes.pk)]['status'],'confirmed')
        self.assertEqual(rows[str(self.no.pk)]['status'],'excluded')
        self.assertEqual(rows[str(self.empty.pk)]['status'],'candidate')
        node=dict(field='company.name',op='eq',value=self.yes.name)
        self.assertEqual(self.match(node)[3],dict(confirmed=1,candidate=0,excluded=2))

    def test_unknown_not_does_not_become_present_and_explicit_absent_is_trusted(self):
        self.assertEqual(self.rows({'not':q()})[str(self.empty.pk)]['status'],'candidate')
        self.assertEqual(self.rows({'not':q()})[str(self.yes.pk)]['status'],'excluded')
        CompanyQualification.objects.create(company=self.empty,category='electronic',status='ABSENT')
        self.assertEqual(self.rows()[str(self.empty.pk)]['status'],'excluded')
        self.assertEqual(self.rows({'not':q()})[str(self.empty.pk)]['status'],'confirmed')

    def test_and_or_and_preference_do_not_weaken_hard_conditions(self):
        CompanyQualification.objects.create(company=self.no,category='itss',status='PRESENT',grade='2')
        self.assertEqual(self.rows({'all':[q(),q('itss')]})[str(self.no.pk)]['status'],'candidate')
        self.assertEqual(self.match({'any':[q(),q('itss')]})[3]['confirmed'],2)
        self.assertEqual(self.match(should=[q('itss')])[2][0]['company_id'],str(self.yes.pk))

    def test_known_expired_legacy_entry_is_not_confirmed(self):
        self.old.valid_until=timezone.localdate()-timedelta(days=1);self.old.save()
        self.assertEqual(self.rows()[str(self.yes.pk)]['status'],'candidate')

    def test_known_conflict_or_revocation_does_not_get_hidden_by_trusted_present(self):
        cert=self.public_certificate(company=self.yes,holder_company_id=self.yes.pk,verification_status='conflict')
        self.assertEqual(self.rows()[str(self.yes.pk)]['verdict'],'conflict')
        cert.verification_status='unverified';cert.revocation_status='revoked';cert.save()
        self.assertEqual(self.rows()[str(self.yes.pk)]['status'],'candidate')

    def test_public_unknown_different_grade_does_not_override_trusted_old_grade(self):
        self.public_certificate(company=self.yes,holder_company_id=self.yes.pk,grade='1')
        self.assertEqual(self.rows()[str(self.yes.pk)]['status'],'confirmed')

    def test_internal_legacy_absent_present_conflict_cannot_confirm(self):
        CompanyQualification.objects.create(company=self.yes,category='electronic',status='ABSENT')
        self.assertEqual(self.rows()[str(self.yes.pk)]['verdict'],'conflict')

    def test_legacy_conflict_cannot_be_overridden_by_independently_verified_positive(self):
        today=timezone.localdate()
        evidence=BaichuanEvidence.objects.create(scope='trusted-legacy-test',source_type='fixture',import_ref='fixture:trusted-positive',
            retrieved_at=timezone.now(),verified_on=today,content_hash='a'*64,excerpt='synthetic fixture',
            locator={'page':1},access_status='verified',subject=dict(company_id=str(self.yes.pk),
                identity_verified=True,holder_type='bidder',authority='issuer',scope_verified=True,
                revocation_checked=True,revocation_checked_on=today.isoformat()))
        self.public_certificate(company=self.yes,holder_company_id=self.yes.pk,holder_name=self.yes.name,
            verification_status='verified',evidence=evidence,valid_from=today-timedelta(days=1),
            valid_until=today+timedelta(days=30),revocation_status='clear')
        # Prove the positive is actually accepted before introducing the conflict.
        with override_settings(BAICHUAN_TRUST_LEGACY_QUALIFICATIONS=False):
            self.assertEqual(self.rows()[str(self.yes.pk)]['status'],'confirmed')
        CompanyQualification.objects.create(company=self.yes,category='electronic',status='ABSENT')
        self.assertEqual(self.rows()[str(self.yes.pk)]['verdict'],'conflict')
        self.old.delete()
        self.assertEqual(self.rows()[str(self.yes.pk)]['verdict'],'conflict')

    def test_other_holder_observation_does_not_conflict_with_bidder_absence(self):
        CompanyQualification.objects.create(company=self.empty,category='electronic',status='ABSENT')
        cert=self.public_certificate(holder_type='manufacturer')
        self.assertEqual(self.rows()[str(self.empty.pk)]['status'],'excluded')
        cert.holder_type='bidder';cert.holder_company_id=self.yes.pk;cert.save()
        self.assertEqual(self.rows()[str(self.empty.pk)]['status'],'excluded')

    def test_requested_standard_without_legacy_standard_evidence_stays_unknown(self):
        from portal.baichuan_legacy_matching import trusted_legacy_verdict
        from portal.baichuan_matching import company_facts, load_rules
        compiled,rules,_,_=self.match()
        rule={**rules['electronic'],'standard':'specific standard'}
        self.assertIsNone(trusted_legacy_verdict(q(),company_facts(self.yes),compiled,rule,[]))

    def test_real_snapshot_default_xlsx_and_old_snapshot_survive_policy_and_catalog_changes(self):
        temp=tempfile.TemporaryDirectory();self.addCleanup(temp.cleanup)
        with override_settings(BAICHUAN_ASSISTANT_ENABLED=True,BAICHUAN_STORAGE_ROOT=Path(temp.name)):
            conv=BaichuanConversation.objects.create(owner=self.owner,scope='trusted-legacy-test',condition_revision=1)
            payload=dict(schema_version=1,revision=1,must=q(),output=dict(fields=['name','phone'],limit=20))
            self.yes.phone='001234';self.yes.save()
            def snapshot(key):
                job,_=enqueue_job(self.owner,kind='query',payload={'condition_revision':1},conversation=conv,idempotency_key=key)
                claimed=claim_job('trusted-worker',job_id=job.pk)
                saved=save_match_snapshot(self.owner,job.pk,'trusted-worker',claimed.fence,payload)
                job.refresh_from_db();job.status='completed';job.save()
                return saved
            with override_settings(BAICHUAN_TRUST_LEGACY_QUALIFICATIONS=False):old=snapshot('strict')
            old_results=old.results
            current=snapshot('trusted')
            self.assertEqual(current.counts['confirmed'],1)
            self.old.grade='1';self.old.save()
            self.yes.phone='changed';self.yes.save()
            def exported(snap,key):
                job=enqueue_export(self.owner,snap.pk,{'client_request_id':key})
                export=process_export(claim_job('xlsx-worker',job_id=job.pk),'xlsx-worker')
                _,stream=open_export(self.owner,export.pk)
                try:data=stream.read()
                finally:stream.close()
                import openpyxl
                return openpyxl.load_workbook(io.BytesIO(data),data_only=False)
            old_book=exported(old,'old-xlsx'); new_book=exported(current,'new-xlsx')
            self.assertEqual(old_book.worksheets[0].max_row,7)
            text=[[c.value for c in row] for row in new_book.worksheets[0].iter_rows()]
            self.assertTrue(any(self.yes.name in row and '001234' in row for row in text))
            self.assertFalse(any('changed' in row for row in text))
            old.refresh_from_db();current.refresh_from_db()
            self.assertEqual(old.results,old_results)
            self.assertEqual(current.results[0]['status'],'confirmed')

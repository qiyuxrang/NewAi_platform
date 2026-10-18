from copy import deepcopy
from datetime import timedelta
import hashlib
import io
from pathlib import Path
import tempfile
import uuid
from unittest.mock import patch

from django.test import override_settings
from django.utils import timezone

from portal.baichuan_assistant_models import (BaichuanConversation, BaichuanEvidence,
    BaichuanExport, BaichuanJob, BaichuanMatchSnapshot, BaichuanCertificate, BaichuanQualificationRule)
from portal.baichuan_assistant_service import storage_path
from portal.baichuan_exports import (build_workbook, enqueue_export, expire_exports,
    open_export, process_export)
from portal.baichuan_jobs import claim_job, enqueue_job
from portal.baichuan_matching import save_match_snapshot
from portal.baichuan_models import Company
from portal.baichuan_service import BaichuanError
from .base import PortalTestCase


class BaichuanExportTests(PortalTestCase):
    def setUp(self):
        self.owner = self.create_user("export-reader", "product")
        self.other = self.create_user("export-other", "product")
        self.temp = tempfile.TemporaryDirectory(prefix="baichuan-export-test-")
        self.addCleanup(self.temp.cleanup)
        config = override_settings(PORTAL_BAICHUAN_SCOPE_CONFIRMED=True,
            PORTAL_BAICHUAN_READ_USER_IDS=[self.owner.pk, self.other.pk], PORTAL_BAICHUAN_WRITE_USER_IDS=[],
            PORTAL_BAICHUAN_READ_ROLE_CODES=[], PORTAL_BAICHUAN_WRITE_ROLE_CODES=[],
            BAICHUAN_ASSISTANT_SCOPE="isolated", BAICHUAN_ASSISTANT_ENABLED=True,
            BAICHUAN_STORAGE_ROOT=Path(self.temp.name))
        config.enable()
        self.addCleanup(config.disable)
        self.conversation = BaichuanConversation.objects.create(owner=self.owner, scope="isolated", condition_revision=1)
        self.query, _ = enqueue_job(self.owner, kind="query", payload={"synthetic": True},
            conversation=self.conversation, idempotency_key="query-snapshot")
        self.query.status = "completed"
        self.query.save()
        self.company_id = str(uuid.uuid4())
        evidence = BaichuanEvidence.objects.create(scope="isolated", url="https://example.test/source",
            title="虚构证据", retrieved_at=timezone.now(), verified_on=timezone.localdate(),
            locator={"page": 2}, subject={"company_id": self.company_id}, excerpt="虚构测试证据")
        self.rows = [
            {"company_id": self.company_id, "status": "confirmed", "company_version": 4,
             "company": {"name": '=HYPERLINK("https://example.test")', "phone": "00123456789",
                         "credit_code": "001234567890123456"},
             "conditions": [{"condition_id": "must.0", "verdict": "satisfied", "reason": "虚构逐项证据",
                             "evidence_ids": [str(evidence.pk)]}], "unresolved": []},
            {"company_id": str(uuid.uuid4()), "status": "confirmed", "company_version": 2,
             "company": {"name": "+虚构企业B", "phone": "0000123", "credit_code": "001111111111111111"},
             "conditions": [], "unresolved": []},
            {"company_id": str(uuid.uuid4()), "status": "candidate", "company_version": 1,
             "company": {"name": "@虚构候选C", "phone": "0000999"},
             "conditions": [{"condition_id": "must.0", "verdict": "unknown", "reason": "缺少主体依据", "evidence_ids": []}],
             "unresolved": ["主体待核实"]},
            {"company_id": str(uuid.uuid4()), "status": "excluded", "company": {"name": "虚构已排除D"}, "conditions": []},
        ]
        self.snapshot = self.make_snapshot(self.rows)

    def make_snapshot(self, rows, revision=1):
        return BaichuanMatchSnapshot.objects.create(job=self.query, conversation=self.conversation,
            condition_revision=1, result_revision=revision, condition_tree={"schema_version": 1,
                "revision": 1, "must": {"field": "company.name", "op": "eq", "value": "虚构"},
                "output": {"fields": ["name", "phone", "credit_code"], "limit": 2}},
            results=deepcopy(rows), counts={"confirmed": 2, "candidate": 1, "excluded": 1},
            fact_cutoff_at=timezone.now(), research_status="partial")

    def enqueue(self, snapshot=None, **changes):
        return enqueue_export(self.owner, (snapshot or self.snapshot).pk,
            {"client_request_id": uuid.uuid4().hex, **changes})

    def generate(self, **changes):
        job = self.enqueue(**changes)
        claimed = claim_job("export-worker", job_id=job.pk)
        export = process_export(claimed, "export-worker")
        _, stream = open_export(self.owner, export.pk)
        try:
            content = stream.read()
        finally:
            stream.close()
        import openpyxl
        return export, content, openpyxl.load_workbook(io.BytesIO(content), data_only=False)

    def assert_code(self, code, function, *args, **kwargs):
        with self.assertRaises(BaichuanError) as caught:
            function(*args, **kwargs)
        self.assertEqual(caught.exception.code, code)
        return caught.exception

    def test_default_confirmed_selection_three_sheets_text_types_and_source_locator(self):
        export, content, workbook = self.generate(scope="confirmed")
        self.assertTrue(content.startswith(b"PK"))
        self.assertEqual(workbook.sheetnames, ["企业结果", "逐项依据", "待核实信息"])
        company = workbook["企业结果"]
        self.assertEqual(company.max_row, 9)
        self.assertEqual(company["C8"].value, self.rows[0]["company"]["name"])
        self.assertEqual(company["D8"].value, "00123456789")
        self.assertEqual(company["E8"].value, "001234567890123456")
        self.assertEqual(company["D9"].value, "0000123")
        self.assertEqual(workbook["待核实信息"].max_row, 7)
        self.assertEqual(workbook["逐项依据"]["H8"].value, "https://example.test/source")
        self.assertEqual(workbook["逐项依据"]["J8"].value, '{"page": 2}')
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        self.assertEqual(cell.data_type, "s")
                        self.assertIsNone(cell.hyperlink)
        self.assertEqual(hashlib.sha256(content).hexdigest(), export.file_hash)
        export.job.refresh_from_db()
        self.assertEqual(export.job.status, "completed")
        self.assertEqual(export.job.counts, {"confirmed": 2, "candidate": 0})

    def test_actual_snapshot_certificate_numbers_survive_xlsx_and_later_certificate_changes(self):
        company = Company.objects.create(name='FICTIONAL NUMBER COMPANY', created_by=self.owner, updated_by=self.owner)
        source = BaichuanEvidence.objects.create(scope='isolated', source_type='fixture',
            url='https://example.test/certificates', retrieved_at=timezone.now(), verified_on=timezone.localdate(),
            content_hash='a' * 64, excerpt='Synthetic original certificate declarations', locator={'page': 3},
            access_status='verified', subject={'company_id': str(company.pk)})
        BaichuanQualificationRule.objects.create(category='itss', version=1, legal_grades=['2'], evidence=source)
        certificates = [BaichuanCertificate.objects.create(company=company, category='itss', grade='2',
            holder_type='manufacturer', holder_name='FICTIONAL MANUFACTURER', certificate_no=number,
            verification_status='unverified', evidence=source) for number in ['001239', '=HYPERLINK("https://example.test")']]
        job, _ = enqueue_job(self.owner, kind='query', payload={'condition_revision': 1},
            conversation=self.conversation, idempotency_key='certificate-metadata-query')
        claimed = claim_job('certificate-query-worker', job_id=job.pk)
        tree = {'schema_version': 1, 'revision': 1, 'must': {'field': 'qualification', 'category': 'itss',
            'op': 'grade_exact', 'value': '2', 'holder': 'manufacturer'}}
        snapshot = save_match_snapshot(self.owner, job.pk, 'certificate-query-worker', claimed.fence, tree)
        self.assertEqual(snapshot.counts, {'confirmed': 0, 'candidate': 1, 'excluded': 0})
        original = deepcopy(snapshot.results)
        export, content, workbook = self.generate(snapshot=snapshot, include_candidates=True)
        sheet = workbook['逐项依据']
        self.assertEqual([sheet.cell(7, col).value for col in range(1, 11)],
            ['企业 ID', '企业名称', '匹配状态', '条件 ID', '逐项判断', '依据说明', '证据 ID', '来源', '核验日期', '原文定位'])
        self.assertEqual(sheet['K7'].value, '证书编号')
        self.assertEqual({sheet.cell(row, 11).value for row in (8, 9)},
            {'001239', '=HYPERLINK("https://example.test")'})
        for row in (8, 9):
            self.assertEqual(sheet.cell(row, 11).data_type, 's')
            self.assertIsNone(sheet.cell(row, 11).hyperlink)
            self.assertEqual(sheet.cell(row, 12).value, 'manufacturer')
            self.assertEqual(sheet.cell(row, 13).value, 'FICTIONAL MANUFACTURER')
            self.assertEqual(sheet.cell(row, 20).value, '未核实')
            self.assertEqual(sheet.cell(row, 5).value, '未知')
        for certificate in certificates:
            certificate.certificate_no = 'CHANGED-LATER'
            certificate.holder_name = 'CHANGED-LATER-HOLDER'
            certificate.save()
        # A new export from the same snapshot must not ask the mutable table.
        with patch.object(BaichuanCertificate.objects, 'get_queryset', side_effect=AssertionError('No mutable certificate lookup')):
            _, _, repeated = self.generate(snapshot=snapshot, include_candidates=True)
        self.assertEqual({repeated['逐项依据'].cell(row, 11).value for row in (8, 9)},
            {'001239', '=HYPERLINK("https://example.test")'})
        snapshot.refresh_from_db()
        self.assertEqual(snapshot.results, original)
        _, stream = open_export(self.owner, export.pk)
        try:
            self.assertEqual(stream.read(), content)
        finally:
            stream.close()

    def test_legacy_snapshot_does_not_backfill_certificate_numbers_at_export_time(self):
        with patch.object(BaichuanCertificate.objects, 'get_queryset', side_effect=AssertionError('No certificate fallback')):
            _, _, workbook = self.generate()
        self.assertIsNone(workbook['逐项依据']['K8'].value)
        self.assertIsNone(workbook['逐项依据']['L8'].value)

    def test_frozen_metadata_cannot_cross_company_scope_evidence_or_condition(self):
        source = BaichuanEvidence.objects.get(subject__company_id=self.company_id)
        record = {'company_id': self.company_id, 'evidence_scope': 'isolated', 'evidence_id': str(source.pk),
            'condition_ids': ['must.0'], 'certificate_no': '000123', 'verification_status': 'unverified'}
        malicious = [{**record, 'company_id': 'another-company'}, {**record, 'evidence_scope': 'another-scope'},
            {**record, 'evidence_id': str(uuid.uuid4())}, {**record, 'condition_ids': ['should.9']}]
        rows = deepcopy(self.rows)
        rows[0]['certificate_records'] = malicious + [record]
        snapshot = self.make_snapshot(rows, revision=2)
        _, _, workbook = self.generate(snapshot=snapshot)
        sheet = workbook['逐项依据']
        self.assertEqual(sheet.max_row, 8)
        self.assertEqual(sheet['K8'].value, '000123')
        self.assertEqual(sheet['T8'].value, '未核实')

    def test_certificate_expansion_counts_toward_evidence_sheet_limit(self):
        source = BaichuanEvidence.objects.get(subject__company_id=self.company_id)
        record = {'company_id': self.company_id, 'evidence_scope': 'isolated', 'evidence_id': str(source.pk),
            'condition_ids': ['must.0'], 'certificate_no': '000123'}
        rows = deepcopy(self.rows[:1])
        rows[0]['certificate_records'] = [record, {**record, 'certificate_no': '000124'}]
        snapshot = self.make_snapshot(rows, revision=2)
        job = self.enqueue(snapshot=snapshot)
        with patch('portal.baichuan_exports.MAX_SHEET_ROWS', 8):
            self.assert_code('export_row_limit', build_workbook, BaichuanExport.objects.get(job=job), timezone.now())

    def test_explicit_candidates_separate_limit_and_fields_apply(self):
        _, _, workbook = self.generate(include_candidates=True, limit=3, fields=["phone"])
        self.assertEqual(workbook["企业结果"].max_row, 9)
        pending = workbook["待核实信息"]
        self.assertEqual(pending.max_row, 8)
        self.assertEqual(pending["B8"].value, "待核实候选")
        self.assertEqual(pending["C8"].value, "@虚构候选C")
        self.assertIn("主体待核实", pending["E8"].value)
        _, _, limited = self.generate(limit=1)
        self.assertEqual(limited["企业结果"].max_row, 8)
        self.assertEqual(limited["企业结果"]["B3"].value, "1")

    def test_missing_snapshot_fields_remain_blank_and_new_revision_cannot_change_old_file(self):
        export, original, workbook = self.generate(fields=["name", "address"])
        self.assertIsNone(workbook["企业结果"]["D8"].value)
        changed = deepcopy(self.rows)
        changed[0]["company"]["name"] = "后续研究新名称"
        changed[0]["company"]["address"] = "后续地址"
        self.make_snapshot(changed, revision=2)
        _, stream = open_export(self.owner, export.pk)
        try:
            self.assertEqual(stream.read(), original)
        finally:
            stream.close()

    def test_empty_snapshot_generates_headers_without_fake_companies(self):
        empty = self.make_snapshot([], revision=2)
        job = self.enqueue(snapshot=empty)
        claimed = claim_job("export-worker", job_id=job.pk)
        export = process_export(claimed, "export-worker")
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(storage_path(export.storage_ref).read_bytes()))
        self.assertEqual(workbook["企业结果"].max_row, 7)
        self.assertEqual(workbook["企业结果"]["B3"].value, "0")

    def test_sheet_limit_includes_metadata_and_each_sheet_keeps_its_own_rows(self):
        # Six metadata rows and one header leave one actual result per sheet.
        # Two confirmed rows must fail, while a confirmed and a candidate fit.
        with patch('portal.baichuan_exports.MAX_SHEET_ROWS', 8):
            job = self.enqueue()
            export = BaichuanExport.objects.get(job=job)
            self.assert_code('export_row_limit', build_workbook, export, timezone.now())
            mixed = deepcopy([self.rows[1], self.rows[2]])
            mixed[1]['conditions'] = []
            snapshot = self.make_snapshot(mixed, revision=2)
            job = self.enqueue(snapshot=snapshot, include_candidates=True)
            export = BaichuanExport.objects.get(job=job)
            content, counts = build_workbook(export, timezone.now())
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
            self.assertEqual(counts, {'confirmed': 1, 'candidate': 1})
            self.assertEqual([sheet.max_row for sheet in workbook], [8, 7, 8])
            self.assertEqual(workbook['企业结果']['C8'].value, mixed[0]['company']['name'])
            self.assertEqual(workbook['待核实信息']['C8'].value, mixed[1]['company']['name'])

    def test_idempotent_export_rejects_changed_selection_and_invalid_parameters(self):
        first = self.enqueue(client_request_id="export-same")
        second = self.enqueue(client_request_id="export-same")
        self.assertEqual(first.pk, second.pk)
        self.assertEqual(BaichuanExport.objects.count(), 1)
        self.assert_code("idempotency_conflict", self.enqueue, client_request_id="export-same", include_candidates=True)
        for values in [{"limit": 5001}, {"limit": True}, {"fields": ["arbitrary_sql"]},
                       {"include_candidates": "yes"}, {"result_revision": 2}, {"scope": "all"}]:
            with self.subTest(values=values), self.assertRaises(BaichuanError):
                self.enqueue(**values)

    def test_owner_scope_revocation_expiry_and_cleanup_are_enforced(self):
        export, _, _ = self.generate()
        self.assertEqual(self.assert_code("not_found", open_export, self.other, export.pk).status, 404)
        with override_settings(BAICHUAN_ASSISTANT_SCOPE="other"):
            self.assert_code("not_found", open_export, self.owner, export.pk)
        with override_settings(PORTAL_BAICHUAN_READ_USER_IDS=[]):
            with self.assertRaises(BaichuanError):
                open_export(self.owner, export.pk)
        path = storage_path(export.storage_ref)
        BaichuanExport.objects.filter(pk=export.pk).update(expires_at=timezone.now() - timedelta(seconds=1))
        self.assertEqual(self.assert_code("export_expired", open_export, self.owner, export.pk).status, 410)
        self.assertEqual(expire_exports(), 1)
        self.assertFalse(path.exists())
        self.assertEqual(expire_exports(), 0)
        export.refresh_from_db()
        self.assertEqual(export.storage_ref, "")
        self.assertTrue(export.file_hash)

    def test_unready_tampered_and_missing_artifacts_do_not_download(self):
        job = self.enqueue()
        export = BaichuanExport.objects.get(job=job)
        self.assert_code("export_not_ready", open_export, self.owner, export.pk)
        claimed = claim_job("export-worker", job_id=job.pk)
        process_export(claimed, "export-worker")
        export.refresh_from_db()
        path = storage_path(export.storage_ref)
        path.write_bytes(b"synthetic corruption")
        self.assert_code("export_changed", open_export, self.owner, export.pk)
        path.unlink()
        self.assert_code("export_missing", open_export, self.owner, export.pk)

    def test_failed_publication_cleans_file_and_can_resume_with_new_fence(self):
        job = self.enqueue()
        first = claim_job("export-worker", job_id=job.pk)
        with patch.object(BaichuanExport, "save", side_effect=RuntimeError("synthetic commit crash")):
            with self.assertRaises(RuntimeError):
                process_export(first, "export-worker")
        export = BaichuanExport.objects.get(job=job)
        self.assertEqual(export.storage_ref, "")
        self.assertEqual(list(Path(self.temp.name).iterdir()), [])
        BaichuanJob.objects.filter(pk=job.pk).update(lease_expires_at=timezone.now() - timedelta(seconds=1))
        second = claim_job("recovery-worker", job_id=job.pk)
        self.assertGreater(second.fence, first.fence)
        process_export(second, "recovery-worker")
        export.refresh_from_db()
        self.assertTrue(storage_path(export.storage_ref).is_file())

    def test_stale_writer_cannot_publish_file_after_new_worker_claims(self):
        from portal import baichuan_exports
        job = self.enqueue()
        first = claim_job("old-worker", job_id=job.pk)
        real_publish = baichuan_exports._publish_file
        replacement = []
        def publish_and_reclaim(content):
            result = real_publish(content)
            BaichuanJob.objects.filter(pk=job.pk).update(lease_expires_at=timezone.now() - timedelta(seconds=1))
            replacement.append(claim_job("new-worker", job_id=job.pk))
            return result
        with patch.object(baichuan_exports, "_publish_file", side_effect=publish_and_reclaim):
            self.assert_code("stale_lease", process_export, first, "old-worker")
        self.assertEqual(list(Path(self.temp.name).iterdir()), [])
        process_export(replacement[0], "new-worker")
        self.assertEqual(BaichuanExport.objects.count(), 1)

    def test_missing_writer_and_oversized_cells_report_explicit_capability_or_limit(self):
        with patch("portal.baichuan_exports._writer", side_effect=BaichuanError("export_unavailable", "缺组件", 503)):
            self.assertEqual(self.assert_code("export_unavailable", self.enqueue).status, 503)
        changed = deepcopy(self.rows)
        changed[0]["company"]["name"] = "x" * 32768
        snapshot = self.make_snapshot(changed, revision=2)
        job = self.enqueue(snapshot=snapshot)
        export = BaichuanExport.objects.get(job=job)
        self.assert_code("export_cell_limit", build_workbook, export, timezone.now())
        self.assertEqual(list(Path(self.temp.name).iterdir()), [])

    def test_http_export_download_headers_and_other_owner_denial(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        from portal.baichuan_export_api import download, exports
        factory = APIRequestFactory()
        request = factory.post("/exports/", {"scope": "confirmed", "client_request_id": "http-export"}, format="json")
        force_authenticate(request, user=self.owner)
        response = exports(request, snapshot_id=self.snapshot.pk)
        self.assertEqual(response.status_code, 202)
        claimed = claim_job("export-worker", job_id=response.data["job"]["id"])
        export = process_export(claimed, "export-worker")
        request = factory.get("/download/")
        force_authenticate(request, user=self.other)
        self.assertEqual(download(request, export_id=export.pk).status_code, 404)
        request = factory.get("/download/")
        force_authenticate(request, user=self.owner)
        response = download(request, export_id=export.pk)
        try:
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Cache-Control"], "private, no-store")
            self.assertEqual(response["X-Content-Type-Options"], "nosniff")
            self.assertIn("attachment", response["Content-Disposition"])
            self.assertEqual(hashlib.sha256(b"".join(response.streaming_content)).hexdigest(), export.file_hash)
        finally:
            # APIRequestFactory did not open a request lifecycle; closing a
            # FileResponse would emit request_finished inside TestCase's outer
            # atomic transaction and close its file SQLite connection.
            response.file_to_stream.close()

    def test_excessive_evidence_rows_rejected_before_workbook_allocation(self):
        job = self.enqueue()
        export = BaichuanExport.objects.get(job=job)
        with patch("portal.baichuan_exports.MAX_SHEET_ROWS", 7):
            self.assert_code("export_row_limit", build_workbook, export, timezone.now())
        self.assertEqual(list(Path(self.temp.name).iterdir()), [])

    def test_retry_http_creates_export_attempt_and_real_worker_finishes_it(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        from portal.baichuan_assistant_api import retry
        from portal.baichuan_assistant_worker import run_one
        original = self.enqueue(client_request_id="retry-original")
        with patch("portal.baichuan_exports.build_workbook", side_effect=BaichuanError("export_unavailable", "虚构故障", 503)):
            self.assertTrue(run_one("failed-export-worker", job_id=original.pk))
        original.refresh_from_db()
        self.assertEqual(original.status, "failed")
        old_export = BaichuanExport.objects.get(job=original)
        request = APIRequestFactory().post("/retry/", {"client_request_id": "retry-child"}, format="json")
        force_authenticate(request, user=self.owner)
        response = retry(request, job_id=original.pk)
        self.assertEqual(response.status_code, 202)
        child = BaichuanJob.objects.get(pk=response.data["job"]["id"])
        new_export = BaichuanExport.objects.get(job=child)
        self.assertNotEqual(new_export.pk, old_export.pk)
        self.assertEqual(new_export.snapshot_id, old_export.snapshot_id)
        self.assertEqual(new_export.storage_ref, "")
        self.assertEqual(child.parent_job_id, original.pk)
        self.assertTrue(run_one("retry-export-worker", job_id=child.pk))
        child.refresh_from_db()
        self.assertEqual(child.status, "completed")
        new_export.refresh_from_db()
        _, stream = open_export(self.owner, new_export.pk)
        stream.close()
        old_export.refresh_from_db()
        self.assertEqual(old_export.storage_ref, "")

    def test_retry_rejects_tampered_export_scope_or_selection_without_new_job(self):
        from portal.baichuan_jobs import retry_job
        job = self.enqueue()
        job.status = "failed"
        job.payload = {**job.payload, "selection": {"include_candidates": True, "limit": 1}}
        job.save()
        before = BaichuanJob.objects.count()
        self.assert_code("invalid_export_retry", retry_job, self.owner, job.pk, idempotency_key="tampered-retry")
        self.assertEqual(BaichuanJob.objects.count(), before)

    def test_private_import_reference_is_hidden_from_reader_workbook(self):
        evidence = BaichuanEvidence.objects.create(scope="isolated", import_ref="C:/private/internal/source.xlsx",
            title="原始整理工作表（脱敏）", retrieved_at=timezone.now(), verified_on=timezone.localdate(),
            locator={"sheet": "Sheet1", "row": 2}, subject={"company_id": self.company_id}, excerpt="虚构")
        rows = deepcopy(self.rows)
        rows[0]["conditions"][0]["evidence_ids"] = [str(evidence.pk)]
        snapshot = self.make_snapshot(rows, revision=2)
        job = self.enqueue(snapshot=snapshot)
        export = process_export(claim_job("export-worker", job_id=job.pk), "export-worker")
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(storage_path(export.storage_ref).read_bytes()))
        self.assertEqual(workbook["逐项依据"]["H8"].value, "原始整理工作表（脱敏）")
        for sheet in workbook:
            for row in sheet:
                self.assertNotIn("C:/private", " ".join(str(cell.value or "") for cell in row))

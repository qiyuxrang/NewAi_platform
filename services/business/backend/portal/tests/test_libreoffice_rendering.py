"""Linux renderer contract tests; real UNO execution requires a Linux worker."""

from contextlib import ExitStack
import hashlib
import json
import os
from pathlib import Path
import signal
import subprocess
import sys
import tempfile
from types import SimpleNamespace
from unittest import skipUnless
from unittest.mock import Mock, patch
from zipfile import ZipFile, ZIP_STORED

from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz
from pptx import Presentation

from portal.product_assets import libreoffice_render as renderer
from portal.product_documents import DocumentError, _office_command
from portal import product_office_runtime as runtime


class LibreOfficeRenderingTests(SimpleTestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "source.docx"
        document = Document()
        document.add_heading("A test heading", level=1)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), 'TOC \\o "1-3"')
        document.add_paragraph()._p.append(field)
        document.save(self.source)
        # python-docx's default template contains customXml; rendered product
        # packages intentionally reject it. Keep this synthetic package audited.
        with ZipFile(self.source) as archive:
            main_xml = archive.read("word/document.xml")
        with ZipFile(self.source, "w", ZIP_STORED) as archive:
            archive.writestr("word/document.xml", main_xml)
            archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            archive.writestr("word/_rels/document.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>')
        self.output = self.root / "output"
        self.before = self.source.read_bytes()
        self.details = {"toc_count": 1, "toc_cached": True}
        self.worker_error = None
        self.commands = []
        self.cleanup = []

    def process(self, command, **options):
        self.commands.append((command, options))
        process = Mock(pid=1200 + len(self.commands))
        process.wait.return_value = 0
        if "--uno-worker" in command:
            marker = command.index("--uno-worker")
            source_copy, output, kind = Path(command[marker + 2]), Path(command[marker + 3]), command[marker + 4]
            self.assertNotEqual(source_copy, self.source)
            self.assertEqual(source_copy.read_bytes(), self.before)
            if kind == "word":
                (output / "reviewed.docx").write_bytes(source_copy.read_bytes())
            pdf = fitz.open()
            pdf.new_page().insert_text((50, 50), "Synthetic renderer output")
            pdf.save(output / ("document.pdf" if kind == "word" else "slides.pdf"))
            pdf.close()
            (output / "uno.json").write_text(json.dumps(self.details))
            if self.worker_error:
                process.wait.side_effect = self.worker_error
        return process

    def render(self, *, kind="word", process=None):
        with ExitStack() as stack:
            stack.enter_context(patch.object(sys, "platform", "linux"))
            stack.enter_context(patch.object(os, "pidfd_open", Mock(), create=True))
            stack.enter_context(patch.object(signal, "pidfd_send_signal", Mock(), create=True))
            stack.enter_context(patch.object(signal, "SIGHUP", 1, create=True))
            stack.enter_context(patch.object(signal, "getsignal", return_value=signal.SIG_DFL))
            stack.enter_context(patch.object(signal, "signal"))
            stack.enter_context(patch.object(renderer.subprocess, "Popen", side_effect=process or self.process))
            stack.enter_context(patch.object(renderer, "terminate_owned", side_effect=lambda child: self.cleanup.append(child.pid)))
            return renderer.render(self.source, self.output, executable="/usr/bin/libreoffice",
                                   uno_python="/usr/bin/python3", timeout=10, kind=kind)

    def test_audited_snapshot_private_profile_pipe_and_hashed_outputs(self):
        result = self.render()
        self.assertTrue(result["rendered"])
        self.assertTrue(result["structural_pass"])
        self.assertFalse(result["visually_reviewed"])
        self.assertEqual(result["input_sha256"], hashlib.sha256(self.before).hexdigest())
        self.assertEqual(result["pdf_sha256"], renderer.sha(self.output / "document.pdf"))
        self.assertEqual(result["image_sha256"], [renderer.sha(self.output / "page-001.png")])
        self.assertEqual(self.source.read_bytes(), self.before)
        self.assertEqual(self.cleanup, [1202, 1201])
        listener, worker = self.commands
        self.assertIn("--headless", listener[0])
        self.assertTrue(any(item.startswith("-env:UserInstallation=file:") for item in listener[0]))
        self.assertTrue(any(item.startswith("--accept=pipe,name=portal_") for item in listener[0]))
        self.assertEqual(listener[1]["process_group"], 0)
        self.assertEqual(worker[1]["process_group"], 0)
        self.assertIn("-I", worker[0])
        self.assertFalse(list(self.root.glob("portal-lo-*")))

    def test_external_relationships_and_macros_never_start_office(self):
        from zipfile import ZipFile, ZIP_DEFLATED
        for unsafe in ("external", "macro", "field"):
            with self.subTest(unsafe=unsafe):
                with ZipFile(self.source) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                if unsafe == "external":
                    parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(
                        b"</Relationships>", b'<Relationship Id="evil" Type="hyperlink" Target="https://invalid/" TargetMode="External"/></Relationships>')
                elif unsafe == "macro":
                    parts["word/vbaProject.bin"] = b"unsafe macro"
                else:
                    parts["word/document.xml"] = parts["word/document.xml"].replace(b"TOC ", b"INCLUDETEXT ")
                with ZipFile(self.source, "w", ZIP_DEFLATED) as archive:
                    for name, data in parts.items():
                        archive.writestr(name, data)
                self.before = self.source.read_bytes()
                self.output = self.root / unsafe
                result = self.render()
                self.assertFalse(result["rendered"])
                self.assertEqual(self.commands, [])
                self.assertEqual(self.source.read_bytes(), self.before)

    def test_timeout_and_cancellation_cleanup_both_owned_children(self):
        for error in (subprocess.TimeoutExpired("UNO", 10), renderer.RenderCancelled("cancelled")):
            with self.subTest(error=error):
                self.output = self.root / type(error).__name__
                self.commands, self.cleanup = [], []
                self.worker_error = error
                result = self.render()
                self.assertFalse(result["rendered"])
                self.assertFalse(result["structural_pass"])
                self.assertEqual(self.cleanup, [1202, 1201])
                self.assertTrue((self.output / "render.json").is_file())
                self.assertFalse(list(self.root.glob("portal-lo-*")))

    def test_failed_render_also_checks_original_hash(self):
        def mutate(command, **options):
            child = self.process(command, **options)
            if "--uno-worker" in command:
                self.source.write_bytes(b"source changed")
                child.wait.side_effect = RuntimeError("conversion failed")
            return child
        result = self.render(process=mutate)
        self.assertFalse(result["rendered"])
        self.assertIn("Source document changed", result["unverified_items"][0])
        self.assertEqual(result["input_sha256"], hashlib.sha256(self.before).hexdigest())

    def test_toc_count_and_cache_must_match_audited_intent(self):
        for number, details in enumerate(({"toc_count": 2, "toc_cached": True}, {"toc_count": 1},
                                         {"toc_count": True, "toc_cached": True})):
            with self.subTest(details=details):
                self.details = details
                self.output = self.root / ("toc" + str(number))
                self.assertFalse(self.render()["rendered"])

    def test_presentation_pdf_covers_every_slide_without_rewriting_source(self):
        self.source = self.root / "source.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(self.source)
        self.before = self.source.read_bytes()
        self.details = {"slide_count": 1}
        result = self.render(kind="powerpoint")
        self.assertTrue(result["rendered"])
        self.assertEqual(result["page_count"], 1)
        self.assertFalse((self.output / "reviewed.docx").exists())
        self.assertEqual(self.source.read_bytes(), self.before)
        self.output = self.root / "missing_slide"
        self.details = {"slide_count": 2}
        self.assertFalse(self.render(kind="powerpoint")["rendered"])

    def test_invalid_pdf_is_not_reported_as_a_render(self):
        def invalid_pdf(command, **options):
            child = self.process(command, **options)
            if "--uno-worker" in command:
                (self.output / "document.pdf").write_bytes(b"%PDF-invalid")
            return child
        result = self.render(process=invalid_pdf)
        self.assertFalse(result["rendered"])
        self.assertFalse(result["structural_pass"])

    def test_external_svg_resources_are_rejected_before_office(self):
        with ZipFile(self.source) as archive:
            package = {name: archive.read(name) for name in archive.namelist()}
        for number, svg in enumerate((
                '<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///private"/></svg>',
                '<svg xmlns="http://www.w3.org/2000/svg"><script>unsafe()</script></svg>',
                '<svg xmlns="http://www.w3.org/2000/svg"><style>@import "https://invalid";</style></svg>',
                '<svg xmlns="http://www.w3.org/2000/svg"><rect style="fill:url(https://invalid)"/></svg>')):
            package["word/media/unsafe.svg"] = svg.encode()
            with ZipFile(self.source, "w", ZIP_STORED) as archive:
                for name, data in package.items():
                    archive.writestr(name, data)
            self.before = self.source.read_bytes()
            self.output = self.root / ("svg" + str(number))
            self.assertFalse(self.render()["rendered"])
            self.assertEqual(self.commands, [])


class UnoWorkerContractTests(SimpleTestCase):
    def invoke(self, *, cached="Section One\t1", writer=True):
        entry = Mock()
        entry.supportsService.return_value = True
        entry.getAnchor.return_value.getString.return_value = cached
        other_index = Mock()
        other_index.supportsService.return_value = False
        document = Mock()
        document.supportsService.return_value = writer
        document.DocumentIndexes = SimpleNamespace(Count=2, getByIndex=lambda index: (entry, other_index)[index])
        desktop = SimpleNamespace(loadComponentFromURL=Mock(return_value=document))
        remote = SimpleNamespace(ServiceManager=SimpleNamespace(createInstanceWithContext=Mock(return_value=desktop)))
        resolver = SimpleNamespace(resolve=Mock(return_value=remote))
        context = SimpleNamespace(ServiceManager=SimpleNamespace(createInstanceWithContext=Mock(return_value=resolver)))
        uno = SimpleNamespace(createUnoStruct=lambda name: SimpleNamespace(),
                              getComponentContext=lambda: context, getConstantByName=Mock(return_value=0))
        with tempfile.TemporaryDirectory() as directory, patch.dict(sys.modules, {"uno": uno}):
            output = Path(directory)
            try:
                renderer.uno_worker("unique_private_pipe", output / "private-input.docx", output)
                details = json.loads((output / "uno.json").read_text())
            except RuntimeError:
                details = None
        return desktop, document, uno, details

    def test_never_executes_macros_or_updates_links_and_counts_only_toc(self):
        desktop, document, uno, details = self.invoke()
        properties = {item.Name: item.Value for item in desktop.loadComponentFromURL.call_args.args[3]}
        self.assertTrue(properties["Hidden"])
        self.assertTrue(properties["ReadOnly"])
        self.assertFalse(properties["PickListEntry"])
        self.assertFalse(properties["StartPresentation"])
        self.assertEqual(properties["MacroExecutionMode"], 0)
        self.assertEqual(properties["UpdateDocMode"], 0)
        self.assertEqual([call.args[0] for call in uno.getConstantByName.call_args_list], [
            "com.sun.star.document.MacroExecMode.NEVER_EXECUTE", "com.sun.star.document.UpdateDocMode.NO_UPDATE"])
        self.assertEqual(details, {"toc_count": 1, "toc_cached": True})
        document.close.assert_called_once_with(True)
        self.assertIn("reviewed.docx", document.storeAsURL.call_args.args[0])

    def test_unresolved_cache_or_wrong_document_service_is_closed_without_saving(self):
        for options in ({"cached": ""}, {"cached": "Error! Bookmark not defined"}, {"writer": False}):
            with self.subTest(options=options):
                desktop, document, uno, details = self.invoke(**options)
                self.assertIsNone(details)
                document.storeAsURL.assert_not_called()
                document.close.assert_called_once_with(True)


class LinuxProcessIsolationTests(SimpleTestCase):
    def test_exited_launcher_does_not_skip_descendant_cleanup(self):
        process = Mock(pid=500)
        process.poll.return_value = 0
        with patch.object(os, "getpgrp", return_value=1, create=True), patch.object(renderer, "_signal_group") as send, \
                patch.object(signal, "SIGKILL", 9, create=True):
            renderer.terminate_owned(process)
        self.assertEqual([call.args[0] for call in send.call_args_list], [500, 500])

    def test_owned_session_excludes_unrelated_processes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for pid, session in ((101, 101), (102, 101), (103, 999)):
                target = root / str(pid)
                target.mkdir()
                fields = ["S", "1", str(pid), str(session)] + ["0"] * 15 + [str(pid + 1000)]
                (target / "stat").write_text(f"{pid} (name with ) spaces) " + " ".join(fields))
            self.assertEqual(runtime._session_members(101, root), {101: "1101", 102: "1102"})

    def test_reused_pid_is_not_signalled(self):
        with patch.object(runtime, "_session_members", side_effect=[{101: "original"}, {101: "reused"}]), \
                patch.object(os, "getsid", return_value=1, create=True), \
                patch.object(os, "pidfd_open", return_value=9, create=True), \
                patch.object(os, "close"), patch.object(signal, "pidfd_send_signal", create=True) as send:
            runtime._signal_session_members(101, signal.SIGTERM)
        send.assert_not_called()

    def test_outer_timeout_cleans_only_created_session(self):
        process = Mock(pid=600, returncode=-15)
        process.communicate.side_effect = [subprocess.TimeoutExpired("render", 10), ("", "")]
        process.poll.return_value = -15
        with patch.object(runtime, "sys_platform_linux", return_value=True), \
                patch.object(os, "pidfd_open", create=True), patch.object(signal, "pidfd_send_signal", create=True), \
                patch.object(runtime.subprocess, "Popen", return_value=process) as start, \
                patch.object(runtime, "_cleanup_session") as cleanup:
            with self.assertRaises(subprocess.TimeoutExpired):
                runtime.run_office(["render"], linux=True, capture_output=True, timeout=10)
        self.assertTrue(start.call_args.kwargs["start_new_session"])
        process.terminate.assert_called_once()
        cleanup.assert_called_once_with(600)

    def test_renderer_selection_is_explicit_and_has_no_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            executable, uno_python = Path(directory) / "soffice", Path(directory) / "python"
            executable.write_bytes(b"trusted runtime")
            uno_python.write_bytes(b"trusted runtime")
            with override_settings(PRODUCT_OFFICE_RENDERER="libreoffice", PRODUCT_LIBREOFFICE_EXECUTABLE=executable,
                                   PRODUCT_UNO_PYTHON=uno_python):
                command = _office_command(Path("python"), Path("input.pptx"), Path("out"), 30, kind="powerpoint")
                self.assertEqual(command[3], "powerpoint")
                self.assertTrue(command[2].endswith("libreoffice_render.py"))
            with override_settings(PRODUCT_OFFICE_RENDERER="unknown"):
                with self.assertRaises(DocumentError):
                    _office_command(Path("python"), Path("input.docx"), Path("out"), 30)

    @skipUnless(sys.platform.startswith("linux"), "Native pidfd/session cleanup requires Linux")
    def test_native_timeout_kills_owned_descendant_and_preserves_unrelated_session(self):
        unrelated = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(30)"], start_new_session=True)
        try:
            with tempfile.TemporaryDirectory() as directory:
                marker = Path(directory) / "child.pid"
                code = ("import subprocess,sys,time; from pathlib import Path; "
                        "child=subprocess.Popen([sys.executable,'-c','import time; time.sleep(30)'], "
                        "stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL); "
                        "Path(sys.argv[1]).write_text(str(child.pid)); time.sleep(30)")
                with self.assertRaises(subprocess.TimeoutExpired):
                    runtime.run_office([sys.executable, "-c", code, str(marker)], linux=True,
                                       capture_output=True, timeout=1.5)
                child_pid = int(marker.read_text())
                child_stat = Path("/proc") / str(child_pid) / "stat"
                if child_stat.exists():
                    self.assertEqual(child_stat.read_text().rsplit(")", 1)[1].split()[0], "Z")
                self.assertIsNone(unrelated.poll())
        finally:
            unrelated.terminate()
            unrelated.wait(timeout=5)


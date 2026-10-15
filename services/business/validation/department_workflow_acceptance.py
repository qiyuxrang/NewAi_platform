"""Isolated browser -> real platform proxy -> real Django acceptance.

Run with a Python environment containing backend/platform dependencies and
Playwright, and an installed Edge browser. Every database/file is under a new
.runtime directory. Model calls are stubbed; business HTTP responses are real.
"""
import argparse
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import time
import io
import zipfile

ROOT = Path(__file__).resolve().parents[3]
PASSWORD = 'Sdd!IsolatedOnly2026-Qx'


def django_setup():
    sys.path.insert(0, str(ROOT / 'services/business/backend'))
    os.environ['DJANGO_SETTINGS_MODULE'] = 'config.settings'
    import django
    django.setup()


def seed():
    django_setup()
    from django.core.management import call_command
    from portal.models import User, Role
    from portal.business_models import BusinessLedgerGrant
    call_command('migrate', interactive=False, verbosity=0)
    call_command('seed_portal', verbosity=0)
    for name, role, department in [('manager', 'general_manager', None), ('presales', 'product', 'product'),
                                    ('engineering', 'engineering', 'engineering'), ('finance', 'finance', 'finance'), ('hr', 'hr', 'hr')]:
        user = User.objects.create_user(username='sdd-' + name, password=PASSWORD, department_code=department or '', must_change_password=False)
        user.roles.set(Role.objects.filter(code=role))
        if name in {'presales', 'engineering', 'finance'}:
            BusinessLedgerGrant.objects.create(user=user, department=name, can_edit=True,
                                               can_submit=name != 'finance', can_publish=name != 'finance')


def gateway():
    sys.path.insert(0, str(ROOT / 'platform'))
    from fastapi import FastAPI, Response
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from app.config import settings
    from app.db.engine import get_db
    from app.db.models import Base, User
    from app.auth.service import hash_password
    from app.routes import auth, enterprise
    import uvicorn
    settings.jwt_secret = 'sdd-isolated-proxy-only-' + 'x' * 48
    settings.portal_base_url = os.environ['SDD_PORTAL_ORIGIN']
    engine = create_engine('sqlite:///' + str(Path(os.environ['SDD_OUTPUT']) / 'platform.sqlite3'),
                           connect_args={'check_same_thread': False})
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        for name in ['manager', 'presales', 'engineering', 'finance', 'hr']:
            db.add(User(id='sdd-' + name, username='sdd-' + name, email=name+'@isolated.invalid', password_hash=hash_password(PASSWORD)))
        db.commit()
    class AsyncDB:
        def __init__(self, db): self.db = db
        async def execute(self, statement): return self.db.execute(statement)
        async def commit(self): self.db.commit()
        async def rollback(self): self.db.rollback()
        async def refresh(self, user): self.db.refresh(user)
        def add(self, value): self.db.add(value)
    async def database():
        with Session(engine, expire_on_commit=False) as db:
            yield AsyncDB(db)
    app = FastAPI()
    app.include_router(auth.router)
    app.include_router(enterprise.router)
    app.dependency_overrides[get_db] = database
    # The container is not part of this acceptance; only its shell indicators are fixtures.
    @app.get('/api/openclaw/ping')
    def ping(): return {'message': 'isolated shell fixture'}
    @app.get('/api/openclaw/container/info')
    def container(): return {'status': 'running'}
    @app.get('/api/openclaw/events/stream')
    def events(): return Response(status_code=204)
    uvicorn.run(app, host='127.0.0.1', port=int(os.environ['SDD_GATEWAY_PORT']), log_level='warning')


def tick(kind):
    django_setup()
    if kind == 'engineering':
        from portal.engineering_worker import run_once
        assert run_once(), 'No engineering work claimed'
    else:
        from unittest.mock import patch
        from portal.hr_screening_worker import claim_one, process_one
        with patch('portal.hr_screening_worker.generate_for_use', side_effect=[
            {'content': json.dumps({'skills': {'value': ['SQL'], 'status': 'extracted', 'source_ref': {'quote': 'SQL'}}})},
            {'content': json.dumps({'requirements': []})},
        ]):
            process_one(*claim_one())


def free_port():
    with socket.socket() as sock:
        sock.bind(('127.0.0.1', 0))
        return sock.getsockname()[1]


def wait_url(url):
    import httpx
    for _ in range(100):
        try:
            if httpx.get(url, timeout=1).status_code < 500: return
        except httpx.HTTPError: pass
        time.sleep(.2)
    raise RuntimeError('Local fixture did not start: ' + url)


def browser(base, out, env):
    from playwright.sync_api import sync_playwright, expect
    results, errors = [], []
    def worker(kind): subprocess.run([sys.executable, __file__, '--mode', 'tick-' + kind], env=env, check=True, stdout=subprocess.DEVNULL)
    with sync_playwright() as p:
        browser = p.chromium.launch(channel='msedge', headless=True)
        contexts = []
        def session(name):
            context = browser.new_context(viewport={'width': 1440, 'height': 1000})
            contexts.append(context)
            response = context.request.post(base+'/api/auth/login', data={'username': 'sdd-'+name, 'password': PASSWORD})
            assert response.ok, response.text()
            token = response.json()['access_token']
            response = context.request.post(base+'/api/enterprise/link', headers={'Authorization': 'Bearer '+token},
                data={'username': 'sdd-'+name, 'password': PASSWORD})
            assert response.ok, response.text()
            context.add_init_script('localStorage.setItem("openclaw_access_token", '+json.dumps(token)+')')
            page = context.new_page()
            page.on('pageerror', lambda error: errors.append(str(error)))
            return page
        def api(page, path, body=None, method=None):
            value = page.evaluate('''async ({path,body,method}) => { const response = await fetch('/api/enterprise/business/'+path, {
                method: method || (body ? 'POST' : 'GET'), headers: {Authorization:'Bearer '+localStorage.getItem('openclaw_access_token'),
                ...(body ? {'Content-Type':'application/json'} : {})}, ...(body ? {body:JSON.stringify(body)} : {})});
                return {status:response.status, body: await response.json()}; }''', {'path': path, 'body': body, 'method': method})
            assert 200 <= value['status'] < 300, value
            return value['body']
        try:
            manager = session('manager')
            for code in ['presales', 'engineering', 'finance']:
                employee = session(code)
                employee.goto(base+'/departments/'+('product' if code == 'presales' else code)+'/entry')
                expect(employee.get_by_role('button', name='新增记录' if code == 'finance' else '新增项目', exact=True)).to_be_visible()
                manager.goto(base+'/departments/business/'+('finance' if code == 'finance' else 'projects?source='+code))
                expect(manager.get_by_role('heading', name='财务看板' if code == 'finance' else '部门项目汇总', exact=True).last).to_be_visible()
                name = '隔离同名项目'
                row = {'project_name': name, 'owner': '隔离负责人'}
                if code == 'presales': row.update(status='方案编制', amount='500.00', follow_up_date='2026-10-15')
                elif code == 'engineering': row.update(status='实施中', planned_end='2026-12-01', progress='25')
                else: row.update(contract_amount='100.00', received_amount='20.00', due_date='')
                path = 'business/ledgers/'+code+'/'
                employee.get_by_role('button', name='新增记录' if code == 'finance' else '新增项目', exact=True).click()
                for key, value in row.items():
                    field = employee.locator('#'+('finance-' if code == 'finance' else 'entry-')+key)
                    if code == 'finance' and key == 'owner': continue
                    if key == 'status': field.select_option(label=value)
                    else: field.fill(value)
                employee.get_by_role('button', name='保存草稿', exact=True).click()
                expect(employee.locator('#'+('finance-' if code == 'finance' else 'entry-')+'project_name')).not_to_be_attached()
                ledger = api(employee, path)
                row = ledger['records'][0]
                assert not manager.get_by_text(name, exact=True).count(), 'Draft leaked into manager page'
                ledger = api(employee, path, {'expected_revision': ledger['revision'], 'as_of': '2026-10-15'}, 'PATCH')
                started = time.monotonic()
                if code == 'finance':
                    meta = ledger['record_meta']
                    record_id, entry = next(iter(meta.items()))
                    payload = {'expected_revision': ledger['revision'], 'expected_published_revision': 0,
                               'records': [{'record_id': record_id, 'source_revision': entry['draft_revision'], 'source_checksum': entry['draft_checksum']}]}
                    employee.reload()
                    employee.get_by_role('checkbox', name='选择发布 '+name, exact=True).check()
                    employee.once('dialog', lambda dialog: dialog.accept())
                    employee.get_by_role('button', name='确认发布本人记录', exact=True).click()
                    expect(employee.get_by_text('选定本人记录已发布。', exact=True)).to_be_visible()
                    published = api(employee, path)
                    replay = api(employee, path+'publish/', payload)
                    assert replay['revision'] == published['revision']
                else:
                    employee.reload()
                    employee.get_by_role('button', name='提交台账', exact=True).click()
                    expect(employee.get_by_role('button', name='发布台账', exact=True)).to_be_enabled()
                    employee.get_by_role('button', name='发布台账', exact=True).click()
                    expect(employee.get_by_text('台账已发布。', exact=True)).to_be_visible()
                    published = api(employee, path)
                # The manager page remains open; refresh is driven by the actual hook.
                manager.bring_to_front()
                expect(manager.get_by_text(name, exact=True).first).to_be_visible(timeout=30000)
                elapsed = time.monotonic()-started
                assert elapsed <= 30, elapsed
                board = api(manager, 'business/boards/'+code+'/')
                assert board['source']['revision'] == published['revision']
                assert len(board['records']) == 1
                if code == 'finance':
                    assert next(m for m in board['metrics'] if m['key']=='receivable')['value'] == '80.00'
                    changed = {**row, 'contract_amount':'200.00'}
                    api(employee, path+'records/'+row['project_id']+'/', {'expected_revision':published['revision'], 'record':changed}, 'PUT')
                    assert api(manager, 'business/boards/finance/')['records'][0]['contract_amount']=='100.00'
                else:
                    manager.get_by_role('button', name='查看明细', exact=True).click()
                    expect(manager.get_by_label('项目明细').get_by_text(name, exact=True).first).to_be_visible()
                manager.screenshot(path=str(out/(code+'-published.png')), full_page=True)
                results.append({'department':code, 'published_revision':published['revision'], 'visible_seconds':round(elapsed,2), 'records':1})
            # Engineering UI, real CLI and private download.
            engineering = session('engineering')
            engineering.goto(base+'/departments/engineering/estimate')
            from openpyxl import Workbook
            book = Workbook(); book.active.append(['序号','项目名称','单位','数量']); book.active.append([1,'配电箱','台',2])
            boq = out/'synthetic.xlsx'; book.save(boq); book.close()
            engineering.get_by_label('清单文件', exact=True).set_input_files(str(boq))
            engineering.get_by_role('button', name='上传并预检', exact=True).click()
            expect(engineering.get_by_text('等待预检 · 陕西', exact=True)).to_be_visible()
            worker('engineering')
            engineering.get_by_role('button', name='刷新准备记录', exact=True).click()
            confirm = engineering.get_by_role('button', name='确认 v1 清单与依据，开始测算', exact=True)
            expect(confirm).to_be_visible(); confirm.click()
            worker('engineering')
            engineering.get_by_role('button', name='刷新任务状态', exact=True).click()
            expect(engineering.get_by_role('link', name='下载内部成本草稿', exact=True)).to_be_visible(timeout=15000)
            def signature(content):
                from openpyxl import load_workbook
                workbooks = [content]
                # XLSX is itself a ZIP, distinguished by its content types entry.
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    if '[Content_Types].xml' not in archive.namelist():
                        workbooks = [archive.read(name) for name in sorted(archive.namelist())]
                values = []
                for raw in workbooks:
                    workbook = load_workbook(io.BytesIO(raw), data_only=False)
                    cells = []
                    for sheet in workbook:
                        for row in sheet:
                            for cell in row:
                                value = cell.value
                                if (isinstance(value, (int, float)) or cell.data_type == 'f'
                                        or isinstance(value, str) and (value in {'台','元','元/台'} or '企业价格库' in value or '隔离合成来源' in value)):
                                    cells.append([sheet.title, cell.coordinate, value])
                    workbook.close(); values.append(cells)
                return values
            for count in [1, 2]:
                paths = [boq]
                if count == 2:
                    second = out/'synthetic-second.xlsx'; second.write_bytes(boq.read_bytes()); paths.append(second)
                    engineering.get_by_label('清单文件', exact=True).set_input_files([str(path) for path in paths])
                    engineering.get_by_role('button', name='上传并预检', exact=True).click()
                    expect(engineering.get_by_text('等待预检 · 陕西', exact=True)).to_be_visible()
                    worker('engineering'); engineering.get_by_role('button', name='刷新准备记录', exact=True).click()
                    engineering.get_by_role('button', name='确认 v1 清单与依据，开始测算', exact=True).click()
                    worker('engineering'); engineering.get_by_role('button', name='刷新任务状态', exact=True).click()
                with engineering.expect_download() as saved:
                    engineering.get_by_role('link', name='下载内部成本草稿', exact=True).click()
                downloaded = saved.value; destination = out/f'confirmed-{count}.artifact'
                downloaded.save_as(str(destination))
                token = engineering.evaluate("localStorage.getItem('openclaw_access_token')")
                import httpx
                response = httpx.post(base+'/api/enterprise/business/engineering/jobs/',
                    headers={'Authorization':'Bearer '+token}, data={'region':'陕西'},
                    files=[('files',(path.name,path.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')) for path in paths], timeout=30, trust_env=False)
                assert response.status_code == 201, response.text
                old_id = response.json()['job']['id']; worker('engineering')
                old = api(engineering, 'engineering/jobs/'+old_id+'/')['job']
                assert old['status'] == 'completed', old
                response = httpx.get(base+f'/api/enterprise/business/engineering/jobs/{old_id}/download/', headers={'Authorization':'Bearer '+token}, timeout=30, trust_env=False)
                assert response.status_code == 200
                assert signature(destination.read_bytes()) == signature(response.content), 'Legacy/new numerical, formula, unit or source mismatch'
                results.append({'engineering_equivalence_files':count,'numeric_formula_unit_source_cells':'identical'})
            engineering.screenshot(path=str(out/'engineering-confirmed-result.png'), full_page=True)
            results.append({'engineering':'preflight -> confirm -> real CLI -> downloadable draft'})
            human = session('hr')
            human.goto(base+'/departments/hr/resumes')
            human.get_by_label('上传 JD 文件（TXT、DOCX、PDF）').set_input_files({'name':'jd.txt','mimeType':'text/plain','buffer':'岗位：数据工程师\n任职要求：熟悉SQL\n工作地点：西安'.encode()})
            confirm = human.get_by_role('button', name='确认该版本 JD 及筛选标准', exact=True)
            expect(confirm).to_be_visible()
            assert not api(human, 'hr/recruitment/confirmed-jds/'), 'Upload automatically confirmed a JD'
            confirm.click()
            human.get_by_label('批量上传简历（TXT、DOCX、PDF）').set_input_files({'name':'resume.txt','mimeType':'text/plain','buffer':'姓名：隔离候选人\n技能：SQL'.encode()})
            human.get_by_role('button', name='进行筛选', exact=True).click()
            expect(human.get_by_text('筛选已提交，后台正在处理。', exact=True)).to_be_visible()
            worker('hr')
            human.get_by_role('button', name='刷新进度', exact=True).click()
            human.get_by_role('link', name='查看筛选结果', exact=True).click()
            expect(human.get_by_text('resume.txt', exact=True).first).to_be_visible()
            human.get_by_role('checkbox', name='resume.txt', exact=True).check()
            expect(human.get_by_role('table', name='候选人对比')).to_be_visible()
            human.screenshot(path=str(out/'hr-explicit-confirm-and-results.png'), full_page=True)
            results.append({'hr':'upload draft -> explicit confirmation -> model stub worker -> evidence comparison'})
            assert not errors, errors
        finally:
            for context in contexts: context.close()
            browser.close()
    (out/'results.json').write_text(json.dumps({'scope':'Isolated real HTTP business/proxy; model responses stubbed; no production data', 'results':results, 'page_errors':errors}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'passed':True, 'evidence':str(out), 'results':results}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--mode', default='run', choices=['run','seed','gateway','tick-engineering','tick-hr'])
    args = parser.parse_args()
    if args.mode != 'run':
        output = os.environ.get('SDD_OUTPUT')
        if not output or not Path(output).resolve().is_relative_to((ROOT/'.runtime').resolve()):
            raise RuntimeError('Internal fixture modes require an isolated .runtime output directory')
        database = os.environ.get('PORTAL_SQLITE_PATH')
        if not database or not Path(database).resolve().is_relative_to(Path(output).resolve()):
            raise RuntimeError('Fixture database must be inside its isolated output directory')
    if args.mode == 'seed': return seed()
    if args.mode == 'gateway': return gateway()
    if args.mode.startswith('tick-'): return tick(args.mode[5:])
    import tempfile
    (ROOT/'.runtime').mkdir(exist_ok=True)
    out = Path(tempfile.mkdtemp(prefix='department-acceptance-', dir=ROOT/'.runtime'))
    env = {k:v for k,v in os.environ.items() if not k.startswith(('PORTAL_', 'SDD_')) and k not in {'DJANGO_SETTINGS_MODULE','DATABASE_URL'}}
    portal, proxy, vite = free_port(), free_port(), free_port()
    env.update(PORTAL_SECRET_KEY='sdd-isolated-only-'+'x'*48, PORTAL_DEBUG='1', PORTAL_HTTPS='0',
               PORTAL_ALLOWED_HOSTS='localhost,127.0.0.1,testserver', PORTAL_SQLITE_PATH=str(out/'business.sqlite3'),
               PORTAL_PRODUCT_STORAGE_ROOT=str(out/'product-private'), PORTAL_HR_STORAGE_ROOT=str(out/'hr-private'),
               PORTAL_ENGINEERING_STORAGE_ROOT=str(out/'engineering-private'), PORTAL_ENGINEERING_PYTHON=sys.executable,
               PORTAL_ENGINEERING_COST_CLI=str(ROOT/'services/engineering/tools/cost_skill.py'),
               PORTAL_ENGINEERING_DATA_ROOT=str(out/'engineering-data'), PORTAL_ENGINEERING_ALLOW_ONLINE='0',
               SDD_OUTPUT=str(out), SDD_PORTAL_ORIGIN=f'http://127.0.0.1:{portal}', SDD_GATEWAY_PORT=str(proxy))
    (out/'engineering-data').mkdir()
    # Explicitly synthetic local prices; never borrow the operator's real library.
    from openpyxl import Workbook
    library = out/'engineering-data/library'; library.mkdir()
    book = Workbook(); materials = book.active; materials.title = '物料主数据'
    materials.append(['物料编码','设备名称','单位']); materials.append(['SDD-BOX','配电箱','台'])
    quotes = book.create_sheet('报价记录'); quotes.append(['物料编码','单价（元）','供应商','价格类型','价格年份'])
    quotes.append(['SDD-BOX',100,'隔离合成来源','内部合成价格','2026'])
    book.save(library/'设备价格库_入库版.xlsx'); book.close()
    subprocess.run([sys.executable, __file__, '--mode','seed'], env=env, check=True, stdout=subprocess.DEVNULL)
    config = out/'vite.config.mjs'
    config.write_text(f"import base from {json.dumps((ROOT/'frontend/vite.config.ts').as_posix())}; export default {{...base, server: {{host:'127.0.0.1',port:{vite},strictPort:true,proxy:{{'/api':'http://127.0.0.1:{proxy}'}}}}}}", encoding='utf-8')
    processes, logs = [], []
    try:
        commands = [([sys.executable, str(ROOT/'services/business/backend/manage.py'),'runserver',f'127.0.0.1:{portal}','--noreload'], ROOT, 'portal'),
                    ([sys.executable,__file__,'--mode','gateway'], ROOT, 'gateway'),
                    (['node',str(ROOT/'frontend/node_modules/vite/bin/vite.js'),'--config',str(config)], ROOT/'frontend', 'vite')]
        for command, cwd, name in commands:
            log = (out/(name+'.log')).open('w', encoding='utf-8'); logs.append(log)
            processes.append(subprocess.Popen(command,cwd=cwd,env=env,stdout=log,stderr=subprocess.STDOUT,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0))
        wait_url(f'http://127.0.0.1:{portal}/api/csrf/')
        wait_url(f'http://127.0.0.1:{proxy}/docs')
        wait_url(f'http://127.0.0.1:{vite}/')
        browser(f'http://127.0.0.1:{vite}', out, env)
    finally:
        for process in reversed(processes):
            process.terminate()
            try: process.wait(timeout=5)
            except subprocess.TimeoutExpired: process.kill(); process.wait()
        for log in logs: log.close()
        print('Isolated evidence: '+str(out))


if __name__ == '__main__': main()

# -*- coding: utf-8 -*-
"""把『陕西省消耗量定额』Word 版解析成人工含量 / 材料含量 / 机械含量。

源表结构（纵向：左 4 列为名称区，右侧各列为定额子目）：
    行：编号        | 8-3-11 | 8-3-12 | ...
    行：项目        | 水平敷设(截面mm²以内) ...
    行：项目        | 35 | 50 | 70 | ...
    行：名称 | 单位 | 消耗量 ...
    行：人/工 | 000204000 | 综合二类工 | 工日 | 3.593 | 3.902 | ...
    行：材/料 | 280008013 | 铜芯电缆 | m | (101.000) | ...   ← 括号 = 未计价主材
    行：机/械 | 990304004 | 汽车式起重机 8t | 台班 | 0.009 | ...

输出：
    data/library/定额含量库.json   每条子目一条记录（嵌套 人工/材料/机械）
    data/library/定额含量明细.xlsx 长表：编号|类别|编码|名称|单位|数量|是否未计价

用法：
    python tools/parse_quota_consumption.py 文件1.docx ...
    python tools/parse_quota_consumption.py --scan 目录
"""
import json
import os
import re
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT_DIR = os.path.join(ROOT, "data", "library")

CODE_RE = re.compile(r"^\d+(?:\s*[-–—]\s*\d+){1,3}[a-zA-Z]?$")
UNIT_RE = re.compile(r"计量单位[:：]\s*([^\s]+)")
QUOTED_RE = re.compile(r"^[（(]\s*([0-9.]+)\s*[)）]$")

CAT_MAP = [("人工", "人"), ("材料", "材"), ("机械", "机")]


def norm(v):
    return re.sub(r"[\s\u3000]+", "", (v or "").replace("\n", ""))


def norm_keep(v):
    return re.sub(r"\s+", " ", (v or "").replace("\n", " ").replace("\u3000", " ")).strip()


def row_cells(r):
    return [norm_keep(c.text) for c in r.cells]


def parse_file(path, book=""):
    doc = docx.Document(path)
    items = []
    for table in doc.tables:
        rows = [[norm(c.text) for c in r.cells] for r in table.rows]
        if not rows:
            continue

        # 找编号行
        code_row = None
        name_col = None
        for ri, r in enumerate(rows[:8]):
            for ci in range(len(r) - 1):
                if "编号" in r[ci] and r[ci + 1] and not "编号" in r[ci + 1]:
                    code_row, name_col = ri, ci + 1
                    break
            if code_row is not None:
                break
        if code_row is None:
            continue
        codes = rows[code_row][name_col:]
        if not any(CODE_RE.match(c) for c in codes):
            continue

        chapter = rows[0][0] if rows[0] else ""
        head_text = " ".join(rows[code_row - 1]) if code_row >= 1 else ""
        unit_m = UNIT_RE.search(head_text)
        unit = unit_m.group(1) if unit_m else ""
        work = head_text.split("计量单位")[0].strip()

        # 名称区：编号行之后到「名称/单位/消耗量」表头行
        idx_name = idx_unit = idx_qty = None
        for ri in range(code_row + 1, min(code_row + 6, len(rows))):
            r = rows[ri]
            if "名称" in r and "单位" in r:
                idx_unit = r.index("单位")
                idx_name = idx_unit - 1          # 名称列紧挨单位列左侧
                for ci in range(idx_unit + 1, len(r)):
                    if r[ci] and "消耗" in r[ci]:   # 第一列消耗量 = 第一个子目
                        idx_qty = ci
                        break
                if idx_qty is None:
                    idx_qty = idx_unit + 1
                head_row = ri
                break
        if idx_name is None:
            continue

        # 规格行 = 表头前一行（如 "35 / 50 / 70"），项目行 = 再上一行
        specs = rows[head_row - 1][name_col:] if head_row >= 1 else []
        proj = rows[head_row - 2][name_col:] if head_row >= 2 else specs

        def col_get(arr, j):
            return arr[j] if j < len(arr) else ""

        per_code = {}
        for j, code in enumerate(codes):
            if not CODE_RE.match(code):
                continue
            proj_name = norm_keep(col_get(proj, j))
            spec = norm_keep(col_get(specs, j))
            full = " ".join(x for x in (proj_name, spec) if x)
            per_code[j] = {
                "册别": book,
                "章节": chapter,
                "编号": re.sub(r"\s+", "", code),
                "项目名称": full,
                "单位": unit,
                "工作内容": work,
                "人工": [],
                "材料": [],
                "机械": [],
                "来源": os.path.basename(path),
            }

        # 数据行
        for ri in range(head_row + 1, len(rows)):
            r = rows[ri]
            if len(r) <= idx_unit:
                continue
            cat_cell = r[0] if r else ""
            cat = None
            for name, mark in CAT_MAP:
                if cat_cell.startswith(mark) or cat_cell.startswith(name):
                    cat = name
                    break
            if cat is None:
                continue
            mat_code = r[1] if len(r) > 1 else ""
            mat_name = r[idx_name] if idx_name < len(r) else ""
            mat_unit = r[idx_unit] if idx_unit < len(r) else ""
            if not mat_name:
                continue
            for j, rec in per_code.items():
                qty_cell = col_get(r, idx_qty + j) if idx_qty is not None else ""
                if not qty_cell or qty_cell in ("—", "-", "--", "－"):
                    continue
                q = QUOTED_RE.match(qty_cell)
                unpriced = bool(q)
                try:
                    qty = float(q.group(1) if q else qty_cell.replace(",", ""))
                except ValueError:
                    continue
                rec[cat].append({
                    "编码": mat_code,
                    "名称": mat_name,
                    "单位": mat_unit,
                    "数量": qty,
                    "未计价": unpriced,
                })

        for rec in per_code.values():
            rec["人工工日"] = sum(x["数量"] for x in rec["人工"])
            rec["机械台班"] = sum(x["数量"] for x in rec["机械"])
            items.append(rec)
    return dedupe(items)


def dedupe(items):
    """同一编号重复出现时（跨页重复表），保留内容最全的一条。"""
    best = {}
    for it in items:
        key = (it["册别"], it["编号"])
        cur = best.get(key)
        if cur is None:
            best[key] = it
            continue
        score_new = len(it["人工"]) * 2 + len(it["材料"]) + len(it["机械"])
        score_old = len(cur["人工"]) * 2 + len(cur["材料"]) + len(cur["机械"])
        if score_new > score_old:
            best[key] = it
    return list(best.values())


def find_files(root):
    out = []
    for dirpath, _, files in os.walk(root):
        for f in files:
            if f.endswith(".docx") and "消耗量定额" in dirpath and not f.startswith("~$"):
                out.append(os.path.join(dirpath, f))
    return sorted(out)


def write_outputs(items):
    os.makedirs(OUT_DIR, exist_ok=True)
    js = os.path.join(OUT_DIR, "定额含量库.json")
    with open(js, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False)

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "定额含量明细"
    ws.append(["册别", "章节", "编号", "项目名称", "单位", "类别",
               "材料编码", "名称", "单位2", "数量", "未计价"])
    for it in items:
        for cat in ("人工", "材料", "机械"):
            for d in it[cat]:
                ws.append([it["册别"], it["章节"], it["编号"], it["项目名称"], it["单位"],
                           cat, d["编码"], d["名称"], d["单位"], d["数量"],
                           "是" if d["未计价"] else ""])
    ws.freeze_panes = "A2"
    for i, w in enumerate([18, 26, 12, 40, 9, 8, 14, 30, 9, 11, 9], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    xlsx = os.path.join(OUT_DIR, "定额含量明细.xlsx")
    wb.save(xlsx)
    return js, xlsx


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        return
    files = find_files(args[1]) if args[0] == "--scan" else args

    all_items = []
    for f in files:
        book = os.path.basename(os.path.dirname(f)) or ""
        try:
            items = parse_file(f, book=book)
        except Exception as e:  # noqa: BLE001
            print("解析失败 %s: %s" % (f, e))
            continue
        print("%-56s 子目 %5d  人工行 %5d  材料行 %5d" % (
            os.path.basename(f)[:54], len(items),
            sum(len(i["人工"]) for i in items), sum(len(i["材料"]) for i in items)))
        all_items.extend(items)

    print("-" * 70)
    print("子目合计 %d" % len(all_items))
    if all_items:
        with_labor = [i for i in all_items if i["人工"]]
        print("带人工含量的子目 %d（%.0f%%）" % (
            len(with_labor), len(with_labor) / len(all_items) * 100))
        js, xlsx = write_outputs(all_items)
        print("→", js)
        print("→", xlsx)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()

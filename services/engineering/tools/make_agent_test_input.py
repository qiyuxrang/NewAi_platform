# -*- coding: utf-8 -*-
"""从企业共享资料库生成一份测试用「设备清单 + 项目背景」，放到桌面 agent测试 目录。

来源：
    工程量清单：02-工程技术部/成本测算/工程量清单/2025-337神盘项目劳务机电1.23.xlsx
    项目背景：  02-工程技术部/施工方案/2025-G337...机电施工施工组织设计.1.14.docx

用法：  .venv\\Scripts\\python.exe tools\\make_agent_test_input.py
"""
import os
import re
import sys

import docx
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SRC = r"C:\Users\123\Desktop\项目\企业共享资料库（2018-2026整理）"
XLSX = os.path.join(SRC, r"02-工程技术部\成本测算\工程量清单\2025-337神盘项目劳务机电1.23.xlsx")
DOCX = os.path.join(SRC, r"02-工程技术部\施工方案\2025-G337神木至盘塘公路改扩建工程机电施工施工组织设计.1.14.docx")
OUT_DIR = r"C:\Users\123\Desktop\agent测试"

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "微软雅黑"


def norm(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def read_sheet(ws, name_col, spec_col, unit_col, qty_col, skip_header_at=3, limit=None):
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for r in rows:
        if len(r) <= max(name_col, spec_col, unit_col, qty_col):
            continue
        nm = norm(r[name_col])
        if not nm or not re.match(r"^\d", norm(r[0] or "0")):
            continue
        qty = r[qty_col]
        if not isinstance(qty, (int, float)):
            continue
        out.append({"名称": nm, "规格": norm(r[spec_col])[:180],
                    "单位": norm(r[unit_col]) or "套", "数量": qty})
        if limit and len(out) >= limit:
            break
    return out


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    # 安装类：序号|子目名称|工作内容|单位|数量
    install = read_sheet(wb["1.收费站、道路、服务区（安装）"], 1, 2, 3, 4, limit=12)
    # 土建基础类：序号|子目名称|工作内容|单位|数量
    base = read_sheet(wb["3.道路及服务区-基础部分（土建）"], 1, 2, 3, 4, limit=8)
    wb.close()

    items = []
    for x in install:
        items.append(("设备安装", x))
    for x in base:
        items.append(("土建基础", x))

    out = Workbook()
    ws = out.active
    ws.title = "工程量清单"
    ws["A1"] = "G337神木至盘塘公路改扩建工程（机电）　工程量清单"
    ws.merge_cells("A1:F1")
    ws["A1"].font = Font(name=FONT, size=13, bold=True, color="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    ws["A2"] = "（测试用清单：由企业共享资料库真实项目整理，含设备安装与土建基础两类）"
    ws.merge_cells("A2:F2")
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")
    for i, h in enumerate(["序号", "项目名称", "备注（规格/工作内容）", "单位", "数量", "类别"], 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = Alignment(horizontal="center", vertical="center")
    r = 4
    for k, (cat, x) in enumerate(items, 1):
        vals = [k, x["名称"], x["规格"], x["单位"], x["数量"], cat]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 15 * max(1, len(x["规格"]) // 40 + 1) + 8
        r += 1
    for col, w in zip("ABCDEF", [6, 30, 50, 8, 10, 12]):
        ws.column_dimensions[col].width = w
    path = os.path.join(OUT_DIR, "设备清单_G337神盘机电.xlsx")
    out.save(path)
    print("设备清单 →", path, "共 %d 行" % len(items))

    # 项目背景
    d = docx.Document(DOCX)
    texts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    picked = [t for t in texts if any(k in t for k in (
        "工程概况", "沙峁", "工期", "开工", "完工", "编制依据", "CCTV", "情报板",
        "隧道", "收费", "保通", "交通", "路灯"))][:22]
    bg = """项目背景（G337神木至盘塘公路改扩建工程 · 机电工程）

一、项目基本情况
1. 项目名称：G337神木至盘塘公路改扩建工程（机电工程）
2. 建设地点：陕西省榆林市神木市沙峁镇一带，沿 G337 国道
3. 工程内容：收费站、道路、服务区、隧道的机电系统施工，含收费系统、监控系统（CCTV）、
   F 型可变情报板、道路照明、隧道机电、光缆敷设与熔纤等
4. 工期要求：现场计划 2025 年 3 月开工，2025 年 12 月完工（以开工令为准）
5. 项目经理部驻地：榆林市神木市沙峁镇

二、施工条件与不利因素
1. 边通车边施工：作业面在通车道路及收费广场，需交通保畅，成立交通保畅小组（保通员 5 名），
   作业人员须经交通与路政部门培训，配备专职引导员
2. 隧道内作业：照明与通风条件差，需专项方案
3. 高空与吊装作业：灯杆、情报板立柱、广场摄像机立柱等需吊装与高空作业
4. 与土建、房建交叉施工：机电安装受土建进度制约，需协调界面交接
5. 冬季施工：11-12 月气温低，混凝土浇筑需保温养护

三、甲供与采购口径（本项目惯例）
1. 主要设备（摄像机、情报板、交换机、服务器、收费设备、灯具、灯杆等）由甲方供货，我方只承担安装
2. 基础与土建部分材料（C30/C40 混凝土、钢筋、模板）按清单口径执行
3. 辅材（穿线管、预埋件、螺栓、线缆附件、接地材料等）由我方采购
4. 劳务按总包形式分包给作业队，考虑地理与施工条件因素，报价不压价

四、质量与安全要求
1. 杜绝人身、行车、设备责任事故，施工期负伤率为零
2. 隐蔽工程与接地装置须经监理验收后方可覆盖
3. 联合调试与设备安装按预定计划实施，配合业主与监理完成系统联调

五、资料原文摘录（来自施组设计）
""" + "\n".join("　- " + t[:150] for t in picked)
    bg_path = os.path.join(OUT_DIR, "项目背景_G337神盘机电.txt")
    with open(bg_path, "w", encoding="utf-8") as f:
        f.write(bg)
    print("项目背景 →", bg_path)

    guide = """【测试说明】

1. 打开 http://localhost → 工作室 → 成本测算 Agent
2. 「工程量清单」上传本目录下的：设备清单_G337神盘机电.xlsx
3. 「检索关键词」填（可复制）：
   收费站 摄像机基础 C30混凝土 单价
   （想测别的就问别的，比如：F型可变情报板 基础 单价）
4. 点运行，约 20~40 秒出结果：
   - 生成测算报告（中文报告 + 四版明细）
   - 原始数据（四版 JSON）
   - 同时会在引擎侧生成「太阳能路灯项目六件套」到 data\\output\\太阳能路灯项目
     （注意：六件套目前是按 2431 座太阳能路灯项目的参数出的，用来验证链路；
       换成这个公路机电项目要改参数，跟我说一声我加参数入口）
5. 项目背景_G337神盘机电.txt 的内容可以直接贴进对话，用来测「结合项目背景输出结果」
"""
    guide_path = os.path.join(OUT_DIR, "测试说明.txt")
    with open(guide_path, "w", encoding="utf-8") as f:
        f.write(guide)
    print("测试说明 →", guide_path)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()

# -*- coding: utf-8 -*-
"""把企业资料库里的历年报价单 / 供货一览表 / 投资概算抽成『历史价格库』。

只读源文件，不改原件。输出：
    data/library/历史价格库.xlsx
    data/library/历史价格库.json

用法：
    .venv\\Scripts\\python.exe tools\\parse_history_prices.py "目录"
"""
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT_DIR = os.path.join(ROOT, "data", "library")

NAME_KEYS = ["设备名称", "材料名称", "产品名称", "货物名称", "名称", "子目名称",
             "项目名称", "费用名称", "分项工程名称", "建设内容"]
PRICE_KEYS = ["综合单价", "单价", "价格", "报价", "合价", "金额", "小计"]
SKIP_NAME_RE = re.compile(r"^(小计|合计|总计|累计|其中|备注|说明|以下空白|分部分项合计)")
UNIT_WAN = "万元"


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()


def to_num(v):
    s = cell_text(v).replace(",", "").replace("￥", "").replace("¥", "").replace("元", "")
    s = s.replace("（", "(").replace("）", ")")
    m = re.match(r"^-?\d+(?:\.\d+)?$", s)
    if not m:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def find_header(rows, max_scan=15):
    for i, r in enumerate(rows[:max_scan]):
        cells = [cell_text(c) for c in r]
        joined = " ".join(cells)
        has_name = any(k in joined for k in NAME_KEYS)
        has_price = any(k in joined for k in PRICE_KEYS)
        if has_name and has_price:
            return i, cells
    return None, None


def map_cols(header):
    def find(keys, exclude=()):
        for i, h in enumerate(header):
            if not h:
                continue
            if any(k in h for k in keys) and not any(x in h for x in exclude):
                return i
        return None

    return {
        "序号": find(["序号", "编号", "项次"]),
        "名称": find(NAME_KEYS),
        "品牌": find(["品牌", "厂家", "制造厂"]),
        "型号": find(["型号"]),
        "规格": find(["技术规格", "规格", "配置", "参数"], exclude=("单价",)),
        "单位": find(["单位"], exclude=("单价",)),
        "数量": find(["数量", "工程量"]),
        "单价": find(["综合单价", "单价", "价格", "报价"], exclude=("合价", "金额", "小计", "总价")),
        "金额": find(["合计", "合价", "金额", "小计", "总价"]),
        "备注": find(["备注", "说明"]),
    }


def scale_of(header, idx):
    """表头写『万元』时，单价与金额要乘 10000。"""
    if idx is None or idx >= len(header):
        return 1.0
    return 10000.0 if UNIT_WAN in header[idx] else 1.0


def guess_category(path, sheet):
    text = "%s %s" % (path, sheet)
    if re.search(r"劳务|工队|机电|安装费|分包|人工", text):
        return "劳务"
    if re.search(r"软件|平台|系统|功能|开发|服务", text):
        return "软件/服务"
    if re.search(r"土建|基础|道路|桥涵|管网|土方|混凝土", text):
        return "土建"
    return "设备/材料"


def parse_workbook(path, stats):
    rows_out = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        stats["failed"].append("%s: %s" % (os.path.basename(path), str(e)[:60]))
        return rows_out

    rel = os.path.relpath(path, stats["root"])
    for ws in wb.worksheets:
        try:
            data = list(ws.iter_rows(values_only=True))
        except Exception:
            continue
        hr, header = find_header(data)
        if hr is None:
            continue
        cols = map_cols(header)
        if cols["名称"] is None or (cols["单价"] is None and cols["金额"] is None):
            continue
        cat = guess_category(rel, ws.title)
        got = 0
        for n, r in enumerate(data[hr + 1:], start=hr + 2):
            def get(key):
                i = cols.get(key)
                return r[i] if i is not None and i < len(r) else None

            name = cell_text(get("名称"))
            if not name or SKIP_NAME_RE.match(name) or len(name) > 120:
                continue
            price = to_num(get("单价"))
            amount = to_num(get("金额"))
            if price is None and amount is None:
                continue
            if price is not None:
                price *= scale_of(header, cols["单价"])
            if amount is not None:
                amount *= scale_of(header, cols["金额"])
            qty = to_num(get("数量"))
            rows_out.append({
                "来源文件": rel,
                "表名": ws.title,
                "行号": n,
                "类别": cat,
                "名称": name,
                "品牌": cell_text(get("品牌")),
                "型号": cell_text(get("型号")),
                "规格": cell_text(get("规格"))[:200],
                "单位": cell_text(get("单位")),
                "数量": qty,
                "单价元": price,
                "金额元": amount,
                "备注": cell_text(get("备注"))[:120],
            })
            got += 1
        if got:
            stats["sheets"] += 1
            stats["by_file"][rel] = stats["by_file"].get(rel, 0) + got
    wb.close()
    return rows_out


def main():
    root = sys.argv[1] if len(sys.argv) > 1 else \
        r"C:\Users\123\Desktop\项目\企业共享资料库（2018-2026整理）\02-工程技术部\成本测算"
    stats = {"root": root, "sheets": 0, "failed": [], "by_file": {}}

    files = []
    for dirpath, _, names in os.walk(root):
        for f in names:
            if f.endswith((".xlsx", ".xlsm")) and not f.startswith("~$"):
                files.append(os.path.join(dirpath, f))
    files.sort()
    print("待解析 Excel：%d 个" % len(files))

    all_rows = []
    for p in files:
        rows = parse_workbook(p, stats)
        if rows:
            print("  %-58s %4d 行" % (os.path.basename(p)[:56], len(rows)))
        all_rows.extend(rows)

    print("-" * 72)
    print("抽取价格行：%d 条，来自 %d 张表" % (len(all_rows), stats["sheets"]))

    # 去重：同一个『名称+品牌+型号+单位+单价』在多份/多表里重复出现时只留一条，并记出现次数
    keyed = {}
    for r in all_rows:
        key = (r["名称"], r["品牌"], r["型号"], r["单位"], r["单价元"])
        cur = keyed.get(key)
        if cur is None:
            r["出现次数"] = 1
            keyed[key] = r
        else:
            cur["出现次数"] += 1
    uniq = list(keyed.values())
    print("去重后：%d 条（同一物料在多项目重复出现的已合并）" % len(uniq))
    all_rows = uniq

    if stats["failed"]:
        print("解析失败 %d 个文件" % len(stats["failed"]))
        for x in stats["failed"][:6]:
            print("   ", x)

    os.makedirs(OUT_DIR, exist_ok=True)
    with open(os.path.join(OUT_DIR, "历史价格库.json"), "w", encoding="utf-8") as f:
        json.dump(all_rows, f, ensure_ascii=False)

    from openpyxl import Workbook
    heads = ["来源文件", "表名", "行号", "类别", "名称", "品牌", "型号", "规格",
             "单位", "数量", "单价元", "金额元", "出现次数", "备注"]
    wb = Workbook()
    ws = wb.active
    ws.title = "历史价格库"
    ws.append(heads)
    for r in all_rows:
        ws.append([r.get(h) for h in heads])
    ws.freeze_panes = "A2"
    for i, w in enumerate([34, 22, 6, 10, 36, 14, 18, 30, 8, 10, 12, 14, 20], 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w
    xlsx = os.path.join(OUT_DIR, "历史价格库.xlsx")
    wb.save(xlsx)
    print("→", xlsx)
    print("→", os.path.join(OUT_DIR, "历史价格库.json"))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()

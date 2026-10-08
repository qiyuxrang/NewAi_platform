# -*- coding: utf-8 -*-
"""从企业资料库的『采购清单（无价）』里拼一份测试用工程量清单，用来验证联网取价。

用法：  .venv\\Scripts\\python.exe tools\\make_test_boq.py
输出：  data/input/测试清单_设备.xlsx
"""
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BASE = r"C:\Users\123\Desktop\项目\企业共享资料库（2018-2026整理）\02-工程技术部\成本测算"

PICKS = [
    (r"采购清单及产品说明\智能化系统\网络交换设备\网络交换设备.xlsx", 2),
    (r"采购清单及产品说明\智能化系统\光纤网线类\光纤网线类.xlsx", 2),
    (r"采购清单及产品说明\智能化系统\机柜箱体配电\机柜箱体配电.xlsx", 1),
    (r"采购清单及产品说明\生产集控及调度系统\安防车辆管控\安防车辆管控.xlsx", 2),
    (r"采购清单及产品说明\智能化系统\安装辅材杂项\安装辅材杂项.xlsx", 1),
]


def rows_of(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        data = list(ws.iter_rows(values_only=True))
        head = None
        for i, r in enumerate(data[:6]):
            cells = [str(c).strip() if c is not None else "" for c in r]
            if any("设备名称" in c or "名称" in c for c in cells) and any(
                    "数量" in c for c in cells):
                head = i
                idx_name = next(i2 for i2, c in enumerate(cells) if "名称" in c)
                idx_spec = next((i2 for i2, c in enumerate(cells)
                                 if "参数" in c or "规格" in c or "型号" in c), None)
                idx_unit = next((i2 for i2, c in enumerate(cells) if c == "单位"), None)
                idx_qty = next((i2 for i2, c in enumerate(cells) if "数量" in c), None)
                break
        if head is None:
            continue
        for r in data[head + 1:]:
            def g(i):
                return r[i] if i is not None and i < len(r) and r[i] is not None else ""

            name = str(g(idx_name)).strip()
            qty = g(idx_qty)
            if not name or not isinstance(qty, (int, float)) or qty <= 0:
                continue
            out.append({"名称": name, "规格": str(g(idx_spec))[:120],
                        "单位": str(g(idx_unit)).strip() or "台", "数量": qty})
    wb.close()
    return out


def main():
    rows = []
    for rel, cap in PICKS:
        p = os.path.join(BASE, rel)
        if not os.path.exists(p):
            print("缺文件:", rel)
            continue
        got = rows_of(p)[:cap]
        print("  %-46s 取 %d 行" % (os.path.basename(rel), len(got)))
        rows.extend(got)

    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "工程量清单"
    ws["A1"] = "测试清单（设备类，价格库与历史价均未收录，用于验证联网取价）"
    ws["A1"].font = Font(bold=True, size=12)
    for i, h in enumerate(["序号", "项目名称", "备注", "单位", "数量"], 1):
        ws.cell(row=3, column=i, value=h).font = Font(bold=True)
    for i, r in enumerate(rows, 1):
        ws.cell(row=3 + i, column=1, value=i)
        ws.cell(row=3 + i, column=2, value=r["名称"])
        ws.cell(row=3 + i, column=3, value=r["规格"])
        ws.cell(row=3 + i, column=4, value=r["单位"])
        ws.cell(row=3 + i, column=5, value=r["数量"])
    for col, w in zip("ABCDE", [6, 34, 60, 8, 10]):
        ws.column_dimensions[col].width = w
    out = os.path.join(ROOT, "data", "input", "测试清单_设备.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print("→", out, "共 %d 行" % len(rows))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()

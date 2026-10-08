# -*- coding: utf-8 -*-
"""把「互联网检索」来源的核验链接补齐。

背景：早年六件套回写知识库时没把链接带上（kbsync 里写死 "链接": ""），
而且回写是按「名称+规格+单位」整条覆盖，把 webprice.remember() 存的链接也冲掉了。
结果是出表时只能写「（未取到具体链接）」——人工审核点不回去。

代码侧已修（kbsync 带链接、取价时回缓存捡链接），本脚本负责把**存量**补回来：
    ① 网价缓存里有链接的，回填到 data/library/网价入库.jsonl
    ② 桌面「成本测算输出」里已生成的六件套，把空的链接单元格补上

用法：
    .venv\\Scripts\\python.exe tools\\repair_links.py
"""
import glob
import json
import os
import re
import sqlite3
import sys
import urllib.parse

import openpyxl
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

CACHE_DB = os.path.join(ROOT, "data", "library", "网价缓存.sqlite")
MEMO = os.path.join(ROOT, "data", "library", "网价入库.jsonl")
OUT_ROOT = r"C:\Users\123\Desktop\成本测算输出"

LINK_COL = 10          # ③④⑤⑥ 表：核验链接在第 11 列（0 基）


def load_cache_links():
    """缓存 + 自动入库记录里 key → 链接。"""
    out = {}
    if os.path.exists(CACHE_DB):
        con = sqlite3.connect(CACHE_DB)
        for key, payload in con.execute("SELECT key, payload FROM webprice"):
            try:
                d = json.loads(payload)
            except Exception:
                continue
            if isinstance(d, dict) and d.get("链接"):
                out[key] = d["链接"]
        con.close()
    # 自动入库的网价记录里也存了链接（较新的是带着链接一起入库的）
    if os.path.exists(MEMO):
        for line in open(MEMO, encoding="utf-8"):
            line = line.strip()
            if not line:
                continue
            try:
                d = json.loads(line)
            except Exception:
                continue
            if d.get("链接"):
                out.setdefault(ckey(d.get("名称"), d.get("规格"), d.get("单位")), d["链接"])
    return out


def lookup_link(name, spec, unit, cache):
    """按几种写法找链接：原名 → 去掉〔数量单位〕后缀 → 只用名称。"""
    cands = [ckey(name, spec, unit)]
    bare = re.sub(r"〔[^〕]*〕$", "", str(name or "")).strip()
    if bare != name:
        cands.append(ckey(bare, spec, unit))
    for k in cands:
        if k in cache:
            return cache[k]
    # 名称完全相同、规格略有出入的，也认（规格列在表里被截断过）
    bare_key = re.sub(r"\s+", "", str(bare or ""))
    for k, v in cache.items():
        if k.split("|")[0] == bare_key:
            return v
    return ""


def search_link(name, spec):
    """实在找不回原页面时，给一个能点开的检索页（人工审核用）。"""
    q = "%s %s 价格" % (str(name or "").strip(), str(spec or "").strip()[:30])
    return "https://www.baidu.com/s?wd=" + urllib.parse.quote(q)


def ckey(name, spec, unit=""):
    return "|".join([re.sub(r"\s+", "", str(name or "")),
                     re.sub(r"\s+", "", str(spec or ""))[:40],
                     str(unit or "")])


def fix_kb(cache):
    """回填 网价入库.jsonl 里缺链接的记录。"""
    if not os.path.exists(MEMO):
        return 0, 0
    recs = []
    n_all = n_fix = 0
    for line in open(MEMO, encoding="utf-8"):
        line = line.strip()
        if not line:
            continue
        try:
            d = json.loads(line)
        except Exception:
            continue
        n_all += 1
        if not d.get("链接"):
            hit = lookup_link(d.get("名称"), d.get("规格"), d.get("单位"), cache)
            if hit:
                d["链接"] = hit
                n_fix += 1
        recs.append(d)
    with open(MEMO, "w", encoding="utf-8") as f:
        for d in recs:
            f.write(json.dumps(d, ensure_ascii=False) + "\n")
    return n_all, n_fix


def fix_workbook(path, cache):
    """把六件套里空着的核验链接补上（只补"互联网检索"来源的行）。"""
    wb = openpyxl.load_workbook(path)
    n_fix = 0
    for nm in wb.sheetnames:
        if nm not in ("③成本测算（库+网检索）", "④成本测算（纯网检索）",
                      "⑤劳务报价（库+网检索）", "⑥劳务报价（纯网检索）"):
            continue
        ws = wb[nm]
        for row in ws.iter_rows(min_row=4):
            r = [c.value for c in row]
            if not isinstance(r[0], (int, float)) or not r[1]:
                continue
            src = str(r[7] or "")
            if not any(k in src for k in ("互联网检索", "爱采购", "阿里")):
                continue
            if str(r[LINK_COL] or "").startswith("http"):
                continue
            hit = lookup_link(r[1], r[2], r[3], cache) or search_link(r[1], r[2])
            if hit:
                cell = row[LINK_COL]
                cell.value = hit
                cell.font = Font(name="微软雅黑", size=10, color="0563C1",
                                 underline="single")
                cell.hyperlink = hit
                n_fix += 1
    if n_fix:
        wb.save(path)
    wb.close()
    return n_fix


def main():
    cache = load_cache_links()
    print("网价缓存里有链接的条目：%d" % len(cache))
    n_all, n_fix = fix_kb(cache)
    print("知识库回填：%d 条里补了 %d 条链接" % (n_all, n_fix))
    files = sorted(glob.glob(os.path.join(OUT_ROOT, "*", "1.成本测算六件套.xlsx")))
    total = 0
    for p in files:
        try:
            k = fix_workbook(p, cache)
        except Exception as e:  # noqa: BLE001
            print("  跳过 %s：%s" % (os.path.basename(os.path.dirname(p)), str(e)[:60]))
            continue
        total += k
        print("  %-38s 补链接 %d 行" % (os.path.basename(os.path.dirname(p))[:38], k))
    print("六件套共补 %d 行链接（%d 份）" % (total, len(files)))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()

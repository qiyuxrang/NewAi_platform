# -*- coding: utf-8 -*-
"""探查定额 Word 文件结构：表格数量、形状、样例行。

用法：  python tools/inspect_quota_docx.py "路径.docx" [样例表数]
"""
import collections
import sys

import docx


def main():
    path = sys.argv[1]
    show = int(sys.argv[2]) if len(sys.argv) > 2 else 5
    d = docx.Document(path)
    print("文件：", path)
    print("段落数：", len(d.paragraphs), "表格数：", len(d.tables))
    shapes = collections.Counter((len(t.rows), len(t.columns)) for t in d.tables)
    print("表格形状 (行, 列) -> 个数：")
    for shape, cnt in shapes.most_common(12):
        print("   ", shape, "->", cnt)
    for ti, t in enumerate(d.tables[:show]):
        print("=" * 30, "表", ti, "行", len(t.rows), "列", len(t.columns))
        for r in t.rows[:10]:
            cells = [c.text.replace("\n", "|").strip() for c in r.cells]
            print("    " + " ┃ ".join(cells)[:320])


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()

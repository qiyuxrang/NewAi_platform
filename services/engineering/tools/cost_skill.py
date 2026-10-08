# -*- coding: utf-8 -*-
"""Standalone cost-estimation CLI with blocking XLSX preflight checks."""
import argparse
import contextlib
import hashlib
import json
import os
import re
import sys
import tempfile
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook


CODE_ROOT = Path(__file__).resolve().parents[1]
MAX_XLSX_BYTES = 20 * 1024 * 1024
MAX_SCAN_ROWS = 2000
MAX_QUOTA_NAME_CHARS = 100
MAX_QUOTA_UNIT_CHARS = 20
MAX_QUOTA_CANDIDATES = 5
DEFAULT_OUTPUT_DIR = CODE_ROOT / "data" / "jobs"
INTERNAL_DRAFT_NAME = "1.成本测算六件套（内部草稿）.xlsx"
GENERATED_SHEET_PREFIXES = (
    "①设备清单", "②劳务清单", "③成本测算",
    "④成本测算", "⑤劳务报价", "⑥劳务报价",
)


class CliError(Exception):
    pass


class JsonArgumentParser(argparse.ArgumentParser):
    def error(self, message):
        raise CliError(message)


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _combined_hash(values):
    values = list(values)
    if not values or any(not value for value in values):
        return None
    return hashlib.sha256("".join(values).encode("ascii")).hexdigest()


def _text(value):
    return "" if value is None else str(value).strip()


def _issue(code, message, sheet=None, row=None, **details):
    issue = {"code": code, "message": message, "blocking": True}
    if sheet is not None:
        issue["sheet"] = sheet
    if row is not None:
        issue["row"] = row
    issue.update(details)
    return issue


def inspect_xlsx(raw_path):
    path = Path(raw_path).expanduser().resolve()
    result = {
        "input": str(path), "sha256": None, "size_bytes": None,
        "passed": False, "data_rows": 0, "issues": [],
    }
    if not path.is_file():
        result["issues"].append(_issue("FILE_NOT_FOUND", "文件不存在"))
        return result
    result["size_bytes"] = path.stat().st_size
    result["sha256"] = _sha256(path)
    if path.suffix.lower() != ".xlsx":
        result["issues"].append(_issue("INVALID_EXTENSION", "仅支持 .xlsx 文件"))
    if result["size_bytes"] > MAX_XLSX_BYTES:
        result["issues"].append(_issue(
            "FILE_TOO_LARGE", "文件超过 20MB", limit_bytes=MAX_XLSX_BYTES))
    if result["issues"]:
        return result

    if str(CODE_ROOT) not in sys.path:
        sys.path.insert(0, str(CODE_ROOT))
    from engine import boq

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        result["issues"].append(_issue(
            "INVALID_XLSX", "无法读取 xlsx：%s" % str(exc)[:160]))
        return result

    try:
        if all(any(name.startswith(prefix) for name in workbook.sheetnames)
               for prefix in GENERATED_SHEET_PREFIXES):
            result["issues"].append(_issue(
                "GENERATED_OUTPUT_INPUT",
                "检测到完整六件套输出文件，不能作为原始清单再次测算"))
            return result
        for sheet in workbook.worksheets:
            if sheet.max_row > MAX_SCAN_ROWS:
                result["issues"].append(_issue(
                    "ROW_LIMIT_EXCEEDED", "Sheet 行数超过预检上限",
                    sheet=sheet.title, actual_rows=sheet.max_row,
                    limit_rows=MAX_SCAN_ROWS))
            rows = [list(row) for row in sheet.iter_rows(
                max_row=MAX_SCAN_ROWS, values_only=True)]
            header_row, columns = boq.find_header(rows)
            if header_row is None:
                continue
            missing_columns = [name for name in ("单位", "数量") if name not in columns]
            for name in missing_columns:
                result["issues"].append(_issue(
                    "MISSING_COLUMN", "缺少%s列" % name, sheet=sheet.title,
                    row=header_row + 1, column=name))

            last_name = ""
            last_unit = ""
            seen = {}
            for offset, row_values in enumerate(rows[header_row + 1:], header_row + 2):
                def cell(name):
                    index = columns.get(name)
                    return row_values[index] if index is not None and index < len(row_values) else None

                raw_name = _text(cell("项目名称"))
                spec = _text(cell("技术参数"))
                note = _text(cell("备注"))
                if raw_name in ("小计", "合计", "总计", "合 计") or raw_name.startswith("报价包含"):
                    continue
                quantity_raw = cell("数量")
                quantity = boq.to_number(quantity_raw)
                if not raw_name:
                    if last_name and (quantity is not None or spec):
                        name = last_name
                    else:
                        continue
                else:
                    name = raw_name

                raw_unit = _text(cell("单位"))
                unit = raw_unit or (last_unit if not raw_name else "")
                last_name = name
                if raw_unit:
                    last_unit = raw_unit
                result["data_rows"] += 1

                if quantity is None:
                    result["issues"].append(_issue(
                        "MISSING_QUANTITY", "数量缺失或无法识别", sheet.title, offset,
                        item=name))
                elif quantity < 0:
                    result["issues"].append(_issue(
                        "NEGATIVE_QUANTITY", "数量不能为负数", sheet.title, offset,
                        item=name, quantity=quantity))
                if not unit:
                    result["issues"].append(_issue(
                        "MISSING_UNIT", "单位缺失", sheet.title, offset, item=name))

                duplicate_key = (
                    re.sub(r"\s+", "", name), re.sub(r"\s+", "", spec),
                    re.sub(r"\s+", "", unit), quantity,
                    re.sub(r"\s+", "", note),
                )
                if quantity is not None and unit:
                    if duplicate_key in seen:
                        result["issues"].append(_issue(
                            "DUPLICATE_ROW", "同一 Sheet 存在重复行", sheet.title, offset,
                            item=name, first_row=seen[duplicate_key], blocking=False))
                    else:
                        seen[duplicate_key] = offset
    finally:
        workbook.close()

    if not result["data_rows"]:
        result["issues"].append(_issue(
            "NO_RECOGNIZABLE_DATA", "未识别到可测算的清单明细"))
    result["passed"] = not any(issue["blocking"] for issue in result["issues"])
    return result


def _read_background(raw_path):
    if not raw_path:
        return ""
    path = Path(raw_path).expanduser().resolve()
    if not path.is_file():
        raise CliError("背景文件不存在：%s" % path)
    suffix = path.suffix.lower()
    if suffix == ".txt":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return raw.decode(encoding)[:100000]
            except UnicodeDecodeError:
                pass
        return raw.decode("utf-8", "ignore")[:100000]
    if suffix == ".docx":
        try:
            with zipfile.ZipFile(path) as archive:
                root = ElementTree.fromstring(archive.read("word/document.xml"))
        except Exception as exc:
            raise CliError("无法读取 docx 背景：%s" % str(exc)[:160]) from exc
        paragraphs = []
        for paragraph in root.iter():
            if paragraph.tag.endswith("}p"):
                text = "".join(node.text or "" for node in paragraph.iter()
                               if node.tag.endswith("}t")).strip()
                if text:
                    paragraphs.append(text)
        return "\n".join(paragraphs)[:100000]
    raise CliError("背景文件仅支持 .txt 或 .docx")


def _safe_stem(path):
    value = re.sub(r'[\\/:*?"<>|\s]+', "_", Path(path).stem).strip("._")
    return value[:40] or "job"


def _bounded_text(value, label, limit):
    value = _text(value)
    if not value:
        raise CliError("%s不能为空" % label)
    if len(value) > limit:
        raise CliError("%s不能超过 %d 个字符" % (label, limit))
    if any(ord(char) < 32 or ord(char) == 127 for char in value):
        raise CliError("%s不能包含控制字符" % label)
    return value


def _quota_candidates(raw_name, raw_unit):
    name = _bounded_text(raw_name, "名称", MAX_QUOTA_NAME_CHARS)
    unit = _bounded_text(raw_unit, "单位", MAX_QUOTA_UNIT_CHARS)
    if str(CODE_ROOT) not in sys.path:
        sys.path.insert(0, str(CODE_ROOT))
    from engine import pipeline

    quota = pipeline.load_quota()
    candidates = quota.search(
        pipeline.keywords_of(name), unit=unit, top=MAX_QUOTA_CANDIDATES)
    if not candidates:
        words, majors = pipeline.object_rule(name)
        item_text = pipeline.strip_paren(name)
        keywords = [word for word in (words or []) if word in item_text] or (words or [])
        for major in majors or []:
            candidates.extend(quota.search(
                keywords, unit=unit, top=MAX_QUOTA_CANDIDATES,
                source_kw=major))
        candidates.sort(key=lambda item: (-item.get("score", 0), item.get("编号", "")))
        candidates = candidates[:MAX_QUOTA_CANDIDATES]

    result = []
    for candidate in candidates:
        source = quota.price(candidate.get("编号"), candidate.get("专业")) or {}
        result.append({
            "code": candidate.get("编号"),
            "major": candidate.get("专业"),
            "name": candidate.get("项目名称"),
            "unit": candidate.get("单位"),
            "score": candidate.get("score"),
            "original_source": source.get("来源"),
            "unit_compatible": pipeline.unit_same_strict(candidate.get("单位"), unit),
        })
    return {
        "ok": True, "command": "quota-candidates",
        "query": {"name": name, "unit": unit}, "candidates": result,
    }


def _run_files(args, inspections):
    if not args.allow_online:
        os.environ["DISABLE_WEB_LOOKUP"] = "1"
    else:
        os.environ.pop("DISABLE_WEB_LOOKUP", None)
    background = _read_background(args.background)
    output_base = Path(args.output_dir).expanduser().resolve()
    output_base.mkdir(parents=True, exist_ok=True)
    if str(CODE_ROOT) not in sys.path:
        sys.path.insert(0, str(CODE_ROOT))
    from engine import foursheet

    outputs = []
    for inspection in inspections:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        job_dir = output_base / ("%s_%s_%s" % (
            _safe_stem(inspection["input"]), stamp, uuid.uuid4().hex[:8]))
        job_dir.mkdir(parents=True, exist_ok=False)
        try:
            result = foursheet.run(
                inspection["input"], over={"地区": args.region},
                bg_text=background, out_name=INTERNAL_DRAFT_NAME,
                export=False, out_dir=str(job_dir))
            output_file = Path(result["四件套文件"]).resolve()
            verification = result.get("校验") or {}
            validation_passed = bool(verification.get("通过"))
            outputs.append({
                "input": inspection["input"],
                "input_sha256": inspection["sha256"],
                "preflight_issues": inspection["issues"],
                "status": "completed" if validation_passed else "validation_failed",
                "output_dir": str(job_dir),
                "output_file": str(output_file),
                "output_sha256": _sha256(output_file),
                "validation_passed": validation_passed,
                "validation_issues": verification.get("问题") or [],
                "source_health": result.get("健康度") or {},
                "pending_confirmations": result.get("待确认项") or [],
                "internal_draft": True,
            })
        except SystemExit as exc:
            outputs.append({
                "input": inspection["input"], "input_sha256": inspection["sha256"],
                "preflight_issues": inspection["issues"], "status": "failed",
                "output_dir": str(job_dir), "internal_draft": True,
                "error": str(exc),
            })
        except Exception as exc:
            outputs.append({
                "input": inspection["input"], "input_sha256": inspection["sha256"],
                "preflight_issues": inspection["issues"], "status": "failed",
                "output_dir": str(job_dir), "internal_draft": True,
                "error": "%s: %s" % (type(exc).__name__, str(exc)[:500]),
            })
    return outputs


def _build_parser():
    parser = JsonArgumentParser(prog="cost_skill.py")
    subparsers = parser.add_subparsers(dest="command", required=True)
    inspect_parser = subparsers.add_parser("inspect")
    inspect_parser.add_argument("xlsx", nargs="+", help="1~2 个 .xlsx 清单")
    run_parser = subparsers.add_parser("run")
    run_parser.add_argument("xlsx", nargs="+", help="1~2 个 .xlsx 清单")
    run_parser.add_argument("--background", help="可选 .txt/.docx 项目背景")
    run_parser.add_argument("--region", default="陕西")
    run_parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    run_parser.add_argument("--allow-online", action="store_true")
    quota_parser = subparsers.add_parser("quota-candidates")
    quota_parser.add_argument("--name", required=True)
    quota_parser.add_argument("--unit", required=True)
    return parser


def _self_check():
    with tempfile.TemporaryDirectory(prefix="cost_skill_check_") as temp_dir:
        valid_path = Path(temp_dir) / "valid.xlsx"
        duplicate_path = Path(temp_dir) / "duplicate.xlsx"
        invalid_path = Path(temp_dir) / "invalid.xlsx"
        generated_path = Path(temp_dir) / "generated.xlsx"
        oversized_path = Path(temp_dir) / "oversized.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "项目名称", "单位", "数量", "备注"])
        sheet.append([1, "配电箱", "台", 2, "室内"])
        workbook.save(valid_path)
        sheet.append([2, "配电箱", "台", 2, "室内"])
        workbook.save(duplicate_path)
        workbook.close()

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "项目名称", "单位", "数量", "备注"])
        sheet.append([1, "配电箱", "台", 2, "室内"])
        sheet.append([2, "配电箱", "台", 2, "室内"])
        sheet.append([3, "电缆", "", -1, ""])
        sheet.append([4, "交换机", "台", None, ""])
        workbook.save(invalid_path)
        workbook.close()

        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in (
                "①设备清单", "②劳务清单", "③成本测算（库+网检索）",
                "④成本测算（纯网检索）", "⑤劳务报价（库+网检索）",
                "⑥劳务报价（纯网检索）"):
            workbook.create_sheet(name)
        workbook.save(generated_path)
        workbook.close()

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "项目名称", "单位", "数量", "备注"])
        sheet.append([1, "配电箱", "台", 2, "室内"])
        sheet.cell(row=MAX_SCAN_ROWS + 1, column=1, value="超出扫描范围")
        workbook.save(oversized_path)
        workbook.close()

        before = _sha256(valid_path)
        valid = inspect_xlsx(valid_path)
        duplicate = inspect_xlsx(duplicate_path)
        invalid = inspect_xlsx(invalid_path)
        generated = inspect_xlsx(generated_path)
        oversized = inspect_xlsx(oversized_path)
        codes = {item["code"] for item in invalid["issues"]}
        generated_codes = {item["code"] for item in generated["issues"]}
        oversized_codes = {item["code"] for item in oversized["issues"]}
        checks = {
            "valid_passes": valid["passed"],
            "source_unchanged": before == _sha256(valid_path),
            "duplicate_warning_accepted": duplicate["passed"] and duplicate["data_rows"] == 2
                and any(issue["code"] == "DUPLICATE_ROW" and not issue["blocking"]
                        and issue["sheet"] == "Sheet" and issue["row"] == 3
                        and issue["first_row"] == 2 for issue in duplicate["issues"]),
            "missing_quantity_blocked": "MISSING_QUANTITY" in codes,
            "missing_unit_blocked": "MISSING_UNIT" in codes,
            "negative_quantity_blocked": "NEGATIVE_QUANTITY" in codes,
            "generated_output_blocked": "GENERATED_OUTPUT_INPUT" in generated_codes,
            "row_limit_blocked": "ROW_LIMIT_EXCEEDED" in oversized_codes,
        }
        return {"ok": all(checks.values()), "command": "self-check", "checks": checks}


def dispatch(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv == ["--self-check"]:
        result = _self_check()
        return result, 0 if result["ok"] else 1
    args = _build_parser().parse_args(argv)
    if args.command == "quota-candidates":
        return _quota_candidates(args.name, args.unit), 0
    if not 1 <= len(args.xlsx) <= 2:
        raise CliError("每次必须提供 1~2 个 xlsx 文件")
    inspections = [inspect_xlsx(path) for path in args.xlsx]
    if args.command == "inspect":
        ok = all(item["passed"] for item in inspections)
        return {"ok": ok, "command": "inspect", "files": inspections}, 0 if ok else 2
    if not all(item["passed"] for item in inspections):
        validation_issues = [
            {"input": item["input"], "issue": issue}
            for item in inspections for issue in item["issues"]
        ]
        return {
            "ok": False, "command": "run", "stage": "preflight",
            "message": "预检未通过，未执行任何测算", "files": inspections,
            "input_hash": _combined_hash(item["sha256"] for item in inspections),
            "output_hash": None, "validation_issues": validation_issues,
            "source_health": {}, "result_type": "internal_draft",
        }, 2
    outputs = _run_files(args, inspections)
    ok = all(item["status"] == "completed" for item in outputs)
    validation_issues = [
        {"input": item["input"], "issue": issue}
        for item in outputs for issue in item.get("validation_issues", [])
    ]
    source_health = {
        "files": [{"input": item["input"], "health": item["source_health"]}
                  for item in outputs if "source_health" in item]
    }
    pending_confirmations = [
        {"input": item["input"], "item": pending}
        for item in outputs for pending in item.get("pending_confirmations", [])
    ]
    return {
        "ok": ok, "command": "run", "online_allowed": bool(args.allow_online),
        "input_hash": _combined_hash(item.get("input_sha256") for item in outputs),
        "output_hash": _combined_hash(item.get("output_sha256") for item in outputs),
        "validation_issues": validation_issues, "source_health": source_health,
        "pending_confirmations": pending_confirmations,
        "result_type": "internal_draft", "files": outputs,
    }, 0 if ok else 1


def main(argv=None):
    try:
        with contextlib.redirect_stdout(sys.stderr):
            result, exit_code = dispatch(argv)
    except CliError as exc:
        result, exit_code = {"ok": False, "error": str(exc)}, 2
    except Exception as exc:
        result, exit_code = {
            "ok": False, "error": "%s: %s" % (type(exc).__name__, str(exc)[:500])}, 1
    payload = json.dumps(result, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    sys.stdout.buffer.write(payload + b"\n")
    sys.stdout.buffer.flush()
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())

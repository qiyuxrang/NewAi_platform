# -*- coding: utf-8 -*-
"""入库通道：识别上传文件是「定额数据」还是「价格数据」，写进知识库并返回入库报告。

落库位置（与测算用的价格库/定额库同源，后续取价会自动读到）：
    data/library/入库数据.sqlite      结构化落库
    data/library/入库明细_<日期>.xlsx  人工可查的明细
    data/library/入库日志.jsonl        每次入库的日志
"""
import datetime as dt
import json
import os
import re
import sqlite3

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LIB = os.path.join(ROOT, "data", "library")
DB = os.path.join(LIB, "入库数据.sqlite")

NAME_KEYS = ["设备名称", "材料名称", "产品名称", "品名", "名称", "项目名称", "子目名称"]
PRICE_KEYS = ["单价", "价格", "报价", "含税价", "除税价"]
SPEC_KEYS = ["规格", "型号", "技术参数", "参数", "规格型号"]
QTY_KEYS = ["数量", "工程量"]
UNIT_KEYS = ["单位", "计量单位"]
CODE_KEYS = ["编号", "定额编号", "物料编码", "编码", "材料编码"]
LABOR_KEYS = ["人工费", "人工", "工日", "人工工日"]
MAT_KEYS = ["材料费", "材料"]


def _txt(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def _num(v):
    try:
        return float(str(v).replace(",", "").replace("元", "").strip())
    except Exception:
        return None


def find_header(rows, max_scan=12):
    for i, r in enumerate(rows[:max_scan]):
        cells = [_txt(c) for c in r]
        joined = " ".join(cells)
        if any(k in joined for k in NAME_KEYS) and any(
                k in joined for k in PRICE_KEYS + LABOR_KEYS):
            return i, cells
    return None, None


def classify(cells):
    j = " ".join(cells)
    has_code = any(k in j for k in CODE_KEYS)
    has_labor = any(k in j for k in LABOR_KEYS)
    has_unit_col = any(k in j for k in UNIT_KEYS)
    has_price = any(k in j for k in PRICE_KEYS)
    if has_code and has_labor and has_unit_col:
        return "定额数据"
    if has_price:
        return "价格数据"
    return "未识别"


def col(cells, keys, exclude=()):
    for i, h in enumerate(cells):
        if any(k in h for k in keys) and not any(x in h for x in exclude):
            return i
    return None


def ensure_db():
    os.makedirs(LIB, exist_ok=True)
    con = sqlite3.connect(DB)
    con.executescript(
        "CREATE TABLE IF NOT EXISTS price_rows("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, source TEXT, name TEXT, brand TEXT,"
        " model TEXT, spec TEXT, unit TEXT, price REAL, tax_included TEXT,"
        " note TEXT, ingested_at TEXT);"
        "CREATE TABLE IF NOT EXISTS quota_rows("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, source TEXT, code TEXT, name TEXT,"
        " unit TEXT, labor REAL, material REAL, machine REAL, note TEXT, ingested_at TEXT);"
        "CREATE INDEX IF NOT EXISTS idx_pr_name ON price_rows(name);"
        "CREATE INDEX IF NOT EXISTS idx_qr_code ON quota_rows(code);"
    )
    return con


def ingest(path, kind_hint=""):
    """解析文件并入库，返回报告 dict。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    con = ensure_db()
    now = dt.datetime.now().isoformat(timespec="seconds")
    src = os.path.basename(path)
    report = {"文件": src, "指定类型": kind_hint, "明细": [], "价格条数": 0,
              "定额条数": 0, "跳过条数": 0, "重复条数": 0, "入库时间": now}
    detail_rows = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hr, cells = find_header(rows)
        if hr is None:
            continue
        kind = kind_hint or classify(cells)
        if kind == "未识别":
            report["跳过条数"] += max(0, len(rows) - hr - 1)
            continue
        i_name = col(cells, NAME_KEYS)
        got = 0
        for r in rows[hr + 1:]:
            def g(i):
                return r[i] if i is not None and i < len(r) else None
            nm = _txt(g(i_name))
            if not nm or re.match(r"^(小计|合计|总计)", nm):
                continue
            if kind == "价格数据":
                price = None
                for k in PRICE_KEYS:
                    price = _num(g(col(cells, [k])))
                    if price:
                        break
                if not price:
                    report["跳过条数"] += 1
                    continue
                row = (src, nm, _txt(g(col(cells, ["品牌", "厂家"]))),
                       _txt(g(col(cells, ["型号"]))), _txt(g(col(cells, SPEC_KEYS))),
                       _txt(g(col(cells, UNIT_KEYS))), price,
                       "含税" if "含税" in " ".join(cells) else "未标注",
                       _txt(g(col(cells, ["备注", "说明"]))), now)
                dup = con.execute(
                    "SELECT 1 FROM price_rows WHERE name=? AND IFNULL(spec,'')=? "
                    "AND price=? LIMIT 1", (nm, row[4] or "", price)).fetchone()
                if dup:
                    report["重复条数"] += 1
                    continue
                con.execute("INSERT INTO price_rows(source,name,brand,model,spec,unit,"
                            "price,tax_included,note,ingested_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                            row)
                detail_rows.append(["价格", nm, row[4], row[5], price, src, ws.title])
                report["价格条数"] += 1
            else:
                code = _txt(g(col(cells, CODE_KEYS)))
                if not code:
                    report["跳过条数"] += 1
                    continue
                row = (src, code, nm, _txt(g(col(cells, UNIT_KEYS))),
                       _num(g(col(cells, ["人工费", "人工"]))),
                       _num(g(col(cells, ["材料费", "材料"]))),
                       _num(g(col(cells, ["机械费", "施工机具使用费", "机械"]))),
                       _txt(g(col(cells, ["备注", "说明"]))), now)
                dup = con.execute("SELECT 1 FROM quota_rows WHERE code=? LIMIT 1",
                                  (code,)).fetchone()
                if dup:
                    report["重复条数"] += 1
                    continue
                con.execute("INSERT INTO quota_rows(source,code,name,unit,labor,material,"
                            "machine,note,ingested_at) VALUES(?,?,?,?,?,?,?,?,?)", row)
                detail_rows.append(["定额", nm, code, row[3], "", src, ws.title])
                report["定额条数"] += 1
            got += 1
        report["明细"].append({"sheet": ws.title, "类型": kind, "表头行": hr + 1,
                               "入库条数": got})
    con.commit()
    con.close()
    wb.close()

    con = ensure_db()
    total_price = con.execute("SELECT COUNT(*) FROM price_rows").fetchone()[0]
    total_quota = con.execute("SELECT COUNT(*) FROM quota_rows").fetchone()[0]
    con.close()
    report["库内累计"] = {"价格条数": total_price, "定额条数": total_quota}

    if detail_rows:
        out = openpyxl.Workbook()
        ws = out.active
        ws.title = "入库明细"
        ws.append(["类型", "名称", "规格/编号", "单位", "单价", "来源文件", "Sheet"])
        for r in detail_rows:
            ws.append(r)
        for i, w in enumerate([8, 30, 34, 8, 12, 28, 16], 1):
            ws.column_dimensions[chr(64 + i)].width = w
        xlsx = os.path.join(LIB, "入库明细_%s.xlsx" % dt.date.today().isoformat())
        out.save(xlsx)
        report["入库明细文件"] = xlsx
    with open(os.path.join(LIB, "入库日志.jsonl"), "a", encoding="utf-8") as f:
        f.write(json.dumps(report, ensure_ascii=False) + "\n")
    return report

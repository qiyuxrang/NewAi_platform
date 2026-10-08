# -*- coding: utf-8 -*-
"""读取入库版价格库（物料主数据 + 报价记录）。"""
import os

from openpyxl import load_workbook


class Material(object):
    __slots__ = ("code", "name", "spec", "key_params", "model", "unit",
                 "category", "spec_status")

    def __init__(self, **kw):
        for k in self.__slots__:
            setattr(self, k, kw.get(k, ""))


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(it)]
    for r in it:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        yield dict(zip(header, r))


class PriceLibrary(object):
    """把 xlsx 读进内存；375 个物料 × 511 条报价，量级很小。"""

    def __init__(self, path, mats_sheet="物料主数据", quotes_sheet="报价记录"):
        if not os.path.exists(path):
            raise IOError("找不到价格库文件：%s" % path)
        wb = load_workbook(path, read_only=True, data_only=True)
        self.materials = []
        for r in _rows(wb[mats_sheet]):
            self.materials.append(Material(
                code=str(r.get("物料编码") or "").strip(),
                name=str(r.get("设备名称") or "").strip(),
                spec=str(r.get("规格参数（原表）") or "").strip(),
                key_params=str(r.get("关键参数（整理）") or "").strip(),
                model=str(r.get("型号") or "").strip(),
                unit=str(r.get("单位") or "").strip(),
                category=str(r.get("设备类别") or "").strip(),
                spec_status=str(r.get("规格状态") or "").strip(),
            ))
        self.quotes = {}
        for r in _rows(wb[quotes_sheet]):
            code = str(r.get("物料编码") or "").strip()
            self.quotes.setdefault(code, []).append({
                "project": str(r.get("项目") or "").strip(),
                "supplier": str(r.get("供应商") or "").strip(),
                "price": float(r.get("单价（元）") or 0),
                "tax": str(r.get("是否含税") or "").strip(),
                "price_type": str(r.get("价格类型") or "").strip(),
                "year": str(r.get("价格年份") or "").strip(),
                "note": str(r.get("备注") or "").strip(),
            })
        wb.close()

    def by_code(self, code):
        for m in self.materials:
            if m.code == code:
                return m
        return None

    def index_by_name(self):
        """名称 → 物料列表，用于一级粗筛。"""
        idx = {}
        for m in self.materials:
            idx.setdefault(m.name, []).append(m)
        return idx

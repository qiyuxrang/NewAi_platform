"""Disposable deterministic task -> worker -> snapshot -> XLSX capacity check.

No live service, company database, model or public network is used. Conditions
enter through the worker's real durable-checkpoint resume path, so these timings
exclude natural-language understanding. P95 is nearest-rank on a limited sample.
"""
import argparse
from contextlib import ExitStack
from datetime import datetime, timezone
import hashlib
from io import StringIO
import json
import math
import os
from pathlib import Path
import platform
import statistics
import subprocess
import sys
from tempfile import TemporaryDirectory
from time import perf_counter
import uuid
from zipfile import ZipFile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
PARENT = ROOT / '.runtime/baichuan-checks'


def measure(function, timings, key):
    def wrapped(*args, **kwargs):
        start = perf_counter()
        try:
            return function(*args, **kwargs)
        finally:
            timings[key] = timings.get(key, 0) + (perf_counter() - start) * 1000
    return wrapped


def summarize(samples):
    return {key: {'median_ms': round(statistics.median(row[key] for row in samples), 3),
                  'p95_ms': round(sorted(row[key] for row in samples)[math.ceil(len(samples) * .95) - 1], 3)}
            for key in samples[0] if key.endswith('_ms')}


def execute(args):
    directory = Path(args._directory).resolve(strict=True)
    assert directory.parent == PARENT.resolve() and directory.name.startswith('assistant-e2e-')
    for key in list(os.environ):
        if key.startswith(('PORTAL_', 'PLATFORM_', 'MODEL_GATEWAY_', 'COMPANY_TRIAL_')):
            del os.environ[key]
    os.environ.update(PORTAL_DEBUG='1', PORTAL_HTTPS='0', DJANGO_SETTINGS_MODULE='config.settings',
        PORTAL_SECRET_KEY='fictional-e2e-test-only-not-for-deployment',
        PORTAL_SQLITE_PATH=str(directory / 'synthetic.sqlite3'),
        PORTAL_BAICHUAN_STORAGE_ROOT=str(directory / 'private'),
        PORTAL_PRODUCT_STORAGE_ROOT=str(directory / 'product-private'),
        PORTAL_BAICHUAN_SCOPE_CONFIRMED='1', PORTAL_BAICHUAN_READ_ROLE_CODES='product',
        PORTAL_BAICHUAN_ASSISTANT_ENABLED='1', PORTAL_BAICHUAN_ASSISTANT_SCOPE='fictional-e2e',
        PORTAL_BAICHUAN_RESEARCH_ENABLED='0', PORTAL_BAICHUAN_FACT_WRITE_ENABLED='0',
        PORTAL_PRODUCT_MODEL_CALLS_ALLOWED='0', PORTAL_PRODUCT_WEB_SEARCH_ENABLED='0')
    sys.path[:0] = [str(ROOT / 'services/business/backend'), str(ROOT / 'services/business')]
    import django
    django.setup()
    from django.conf import settings
    from django.core.management import call_command
    from django.db import connection, connections, transaction
    from django.utils import timezone as djtime
    from openpyxl import load_workbook
    from portal import baichuan_assistant_worker as worker
    from portal import baichuan_exports as exports
    from portal import baichuan_matching as matching
    from portal.baichuan_assistant_models import (BaichuanCertificate, BaichuanConversation,
        BaichuanEvidence, BaichuanExport, BaichuanQualificationRule)
    from portal.baichuan_assistant_service import submit_message
    from portal.baichuan_conditions import compile_conditions
    from portal.baichuan_models import Company
    from portal.models import ModelCallLog, Role, User
    assert Path(settings.DATABASES['default']['NAME']).resolve() == directory / 'synthetic.sqlite3'
    report = {'task': 'AI-T16', 'synthetic': True, 'target_database_accessed': False,
              'scope': 'deterministic resumed query task -> worker -> fixed snapshot -> XLSX -> authorized download',
              'natural_language_model_verified': False, 'network_calls': 0, 'model_calls': 0,
              'sqlite': __import__('sqlite3').sqlite_version, 'python': sys.version.split()[0],
              'openpyxl': __import__('openpyxl').__version__, 'cpu': platform.processor(),
              'logical_cpus': os.cpu_count(), 'platform': platform.platform(), 'concurrency': 1,
              'cold_definition': 'Django connection closed, first sample; OS cache not evicted',
              'p95_method': 'nearest rank over warm samples; limited exploratory sample, not population guarantee',
              'warm_samples': args.warm_samples, 'datasets': [], 'complete': False,
              'wall_limit_seconds': args.wall_seconds, 'started_at_utc': datetime.now(timezone.utc).isoformat()}
    output = Path(args.output).resolve()

    def save():
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    def forbid(*_args, **_kwargs):
        raise AssertionError('No network/model calls allowed in synthetic capacity verification')

    try:
        with ExitStack() as guards:
            for target in ('socket.socket.connect', 'socket.socket.connect_ex', 'socket.create_connection',
                           'socket.getaddrinfo', 'portal.baichuan_assistant_worker.generate_for_use'):
                guards.enter_context(patch(target, side_effect=forbid))
            started = perf_counter()
            call_command('migrate', verbosity=0, stdout=StringIO())
            call_command('seed_portal', stdout=StringIO())
            report['schema_setup_ms'] = round((perf_counter() - started) * 1000, 3)
            actor = User.objects.create_user(username='fictional-capacity-only', must_change_password=False)
            actor.roles.add(Role.objects.get(code='product'))
            now, day = djtime.now(), djtime.localdate()
            rule = BaichuanEvidence.objects.create(scope=settings.BAICHUAN_ASSISTANT_SCOPE,
                source_type='fixture', import_ref='fixture:e2e-rule',
                retrieved_at=now, verified_on=day, access_status='verified', excerpt='FICTIONAL EXACT GRADE RULE',
                locator={'fixture': True}, content_hash=hashlib.sha256(b'fictional-rule').hexdigest())
            BaichuanQualificationRule.objects.create(category='itss', evidence=rule, legal_grades=['1', '2', '3', '4'])
            for count in args.sizes:
                started = perf_counter()
                previous = Company.objects.count()
                companies = [Company(name=f'虚构容量验收企业{i:05}', created_by=actor, updated_by=actor)
                             for i in range(previous, count)]
                Company.objects.bulk_create(companies, batch_size=200)
                sources, certificates = [], []
                for i, company in enumerate(companies, previous):
                    source = BaichuanEvidence(scope=settings.BAICHUAN_ASSISTANT_SCOPE,
                        source_type='fixture', import_ref='fixture:e2e-certificate',
                        retrieved_at=now, verified_on=day, access_status='verified', excerpt='FICTIONAL CERTIFICATE',
                        locator={'fixture': True}, content_hash=hashlib.sha256(str(company.pk).encode()).hexdigest(),
                        subject={'company_id': str(company.pk), 'identity_verified': True, 'holder_type': 'bidder',
                                 'authority': 'issuer', 'scope_verified': True, 'revocation_checked': True,
                                 'revocation_checked_on': day.isoformat()})
                    sources.append(source)
                    certificates.append(BaichuanCertificate(company=company, category='itss', grade='2',
                        holder_type='bidder', holder_company_id=company.pk, assertion_status='PRESENT',
                        verification_status='verified' if i % 2 == 0 else 'unverified', evidence=source,
                        valid_from=day, valid_until=day, revocation_status='clear'))
                BaichuanEvidence.objects.bulk_create(sources, batch_size=200)
                BaichuanCertificate.objects.bulk_create(certificates, batch_size=200)
                dataset = {'companies': count, 'certificates': count,
                           'fixture_setup_ms': round((perf_counter() - started) * 1000, 3), 'samples': []}
                report['datasets'].append(dataset)
                save()
                for attempt in range(args.warm_samples + 1):
                    if attempt == 0:
                        connection.close()
                    timings = {}
                    condition = compile_conditions({'schema_version': 1, 'revision': 1,
                        'must': {'field': 'qualification', 'op': 'grade_exact', 'category': 'itss',
                                 'value': '2', 'holder': 'bidder'},
                        'as_of': {'date': day.isoformat(), 'basis': 'query_date'},
                        'output': {'fields': ['name', 'phone', 'province', 'city'], 'limit': count}})
                    total_start = perf_counter()
                    conversation = BaichuanConversation.objects.create(owner=actor, scope=settings.BAICHUAN_ASSISTANT_SCOPE)
                    enqueue_start = perf_counter()
                    with transaction.atomic():
                        _, job, created = submit_message(actor, conversation.pk, {'text': '虚构容量验收，确定规则恢复任务',
                            'attachment_ids': [], 'parent_revision': 0, 'client_request_id': uuid.uuid4().hex})
                        assert created
                        job.checkpoint = {'conditions': condition, 'synthetic_capacity_fixture': True}
                        job.save(update_fields=['checkpoint', 'updated_at'])
                    timings['query_submit_ms'] = (perf_counter() - enqueue_start) * 1000
                    with patch.object(matching, 'match_companies', measure(matching.match_companies, timings, 'structured_match_ms')), \
                            patch.object(worker, 'save_match_snapshot', measure(worker.save_match_snapshot, timings, 'match_and_snapshot_ms')):
                        start = perf_counter()
                        worker.run_one('capacity-query', job.pk)
                        timings['query_worker_ms'] = (perf_counter() - start) * 1000
                    timings['first_snapshot_ready_ms'] = (perf_counter() - total_start) * 1000
                    job.refresh_from_db()
                    # AI-D19: queries finish against stored data independently
                    # of public research; unavailable research is not partial.
                    assert job.status == 'completed', (job.status, job.error_code)
                    assert job.checkpoint['research_policy'] == 'independent_manifest_only'
                    snapshot = job.snapshots.get()
                    assert snapshot.counts == {'confirmed': (count + 1) // 2, 'candidate': count // 2, 'excluded': 0}, snapshot.counts
                    assert len(snapshot.results) == count
                    assert job.checkpoint['capability_gaps'] == []
                    assert snapshot.research_status == 'not_requested'
                    assert job.counts.get('_budget_model_calls', 0) == 0
                    assert job.counts.get('_budget_sources', 0) == 0
                    start = perf_counter()
                    with transaction.atomic():
                        _, export_job, created = submit_message(actor, conversation.pk,
                            {'text': f'导出当前结果前{count}条，包含候选', 'attachment_ids': [], 'parent_revision': 1,
                             'client_request_id': uuid.uuid4().hex})
                        assert created and export_job.kind == 'export'
                    timings['export_submit_ms'] = (perf_counter() - start) * 1000
                    with ExitStack() as export_instrumentation:
                        export_instrumentation.enter_context(patch.object(exports, 'build_workbook',
                            measure(exports.build_workbook, timings, 'xlsx_build_ms')))
                        export_instrumentation.enter_context(patch.object(exports, '_publish_file',
                            measure(exports._publish_file, timings, 'xlsx_publish_ms')))
                        if args.profile_max_row:
                            from openpyxl.worksheet.worksheet import Worksheet
                            original_max_row = Worksheet.max_row.fget
                            def scan_rows(sheet):
                                timings['worksheet_max_row_calls'] = timings.get('worksheet_max_row_calls', 0) + 1
                                return measure(original_max_row, timings, 'worksheet_max_row_scan_ms')(sheet)
                            export_instrumentation.enter_context(patch.object(Worksheet, 'max_row', property(scan_rows)))
                        start = perf_counter()
                        worker.run_one('capacity-export', export_job.pk)
                        timings['export_worker_ms'] = (perf_counter() - start) * 1000
                    export_job.refresh_from_db()
                    assert export_job.status == 'completed', (export_job.status, export_job.error_code)
                    artifact = BaichuanExport.objects.get(job=export_job)
                    start = perf_counter()
                    authorized, stream = exports.open_export(actor, artifact.pk)
                    timings['authorized_download_ms'] = (perf_counter() - start) * 1000
                    try:
                        start = perf_counter()
                        workbook = load_workbook(stream, read_only=True, data_only=False)
                        assert workbook.sheetnames == ['企业结果', '逐项依据', '待核实信息']
                        assert workbook['企业结果'].max_row == 7 + (count + 1) // 2
                        assert workbook['待核实信息'].max_row == 7 + count // 2
                        assert workbook['逐项依据'].max_row == 7 + count
                        actual_rows = 0
                        for sheet in workbook:
                            for row in sheet.iter_rows(min_row=8):
                                actual_rows += 1
                                assert all(cell.data_type != 'f' for cell in row)
                        assert actual_rows == count * 2
                        workbook.close()
                        timings['xlsx_read_validate_ms'] = (perf_counter() - start) * 1000
                    finally:
                        stream.close()
                    timings['total_ms'] = (perf_counter() - total_start) * 1000
                    timings['snapshot_publish_overhead_ms'] = timings['match_and_snapshot_ms'] - timings['structured_match_ms']
                    sample = {'cache': 'connection_cold' if attempt == 0 else 'warm',
                              **{key: round(value, 3) for key, value in timings.items()},
                              'xlsx_bytes': authorized.size_bytes, 'xlsx_sha256': authorized.file_hash,
                              'counts': snapshot.counts, 'xlsx_rows_validated': actual_rows,
                              'query_terminal': job.status, 'export_terminal': export_job.status}
                    if args.profile_max_row:
                        sample['worksheet_max_row_calls'] = timings.get('worksheet_max_row_calls', 0)
                    dataset['samples'].append(sample)
                    if args.compare_row_algorithm and attempt == 0:
                        # Re-render the exact same immutable snapshot and time,
                        # preserving the old implementation only as a benchmark.
                        def previous_row(sheet, values, _positions):
                            number = sheet.max_row + 1 if sheet.max_row > 1 or sheet.cell(1, 1).value is not None else 1
                            if number > exports.MAX_SHEET_ROWS:
                                raise AssertionError('Comparison fixture unexpectedly exceeds sheet row limit')
                            for column, value in enumerate(values, 1):
                                exports._cell(sheet, number, column, value)
                            return number
                        fixed_time = djtime.now()
                        start = perf_counter()
                        with patch.object(exports, '_row', previous_row):
                            previous_content, previous_counts = exports.build_workbook(artifact, fixed_time)
                        before_ms = (perf_counter() - start) * 1000
                        start = perf_counter()
                        current_content, current_counts = exports.build_workbook(artifact, fixed_time)
                        after_ms = (perf_counter() - start) * 1000
                        from io import BytesIO
                        with ZipFile(BytesIO(previous_content)) as previous_zip, ZipFile(BytesIO(current_content)) as current_zip:
                            entries = ['xl/workbook.xml', 'xl/styles.xml',
                                       *sorted(name for name in previous_zip.namelist() if name.startswith('xl/worksheets/'))]
                            assert all(previous_zip.read(name) == current_zip.read(name) for name in entries)
                        assert previous_counts == current_counts == artifact.job.counts
                        dataset['identical_snapshot_row_algorithm_comparison'] = {
                            'before_ms': round(before_ms, 3), 'after_ms': round(after_ms, 3),
                            'speedup': round(before_ms / after_ms, 3),
                            'same_snapshot_and_generated_at': True, 'worksheet_and_style_xml_identical': True,
                            'xml_entries_compared': entries, 'counts_identical': True}
                    if len(dataset['samples']) > 1:
                        dataset['warm_summary'] = summarize(dataset['samples'][1:])
                        dataset['meets_2s_structured_target'] = dataset['warm_summary']['structured_match_ms']['p95_ms'] <= 2000
                        dataset['meets_5s_first_snapshot_without_ai_target'] = dataset['warm_summary']['first_snapshot_ready_ms']['p95_ms'] <= 5000
                    save()
                    print(json.dumps({'companies': count, 'attempt': attempt, 'total_ms': sample['total_ms'],
                                      'match_ms': sample['structured_match_ms'], 'xlsx_ms': sample['xlsx_build_ms']}), flush=True)
            assert ModelCallLog.objects.count() == 0
            report['complete'] = True
            report['completed_at_utc'] = datetime.now(timezone.utc).isoformat()
            save()
    finally:
        connections.close_all()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--sizes', nargs='+', type=int, default=[1000, 5000])
    parser.add_argument('--warm-samples', type=int, default=3)
    parser.add_argument('--wall-seconds', type=int, default=540)
    parser.add_argument('--output', required=True)
    parser.add_argument('--profile-max-row', action='store_true', help='Instrument worksheet row scans; report separately')
    parser.add_argument('--compare-row-algorithm', action='store_true', help='Compare old/new row algorithms on the same snapshot')
    parser.add_argument('--_directory', help=argparse.SUPPRESS)
    args = parser.parse_args()
    if (not args.sizes or args.sizes != sorted(set(args.sizes)) or min(args.sizes) < 2 or max(args.sizes) > 5000
            or not 0 <= args.warm_samples <= 5 or not 30 <= args.wall_seconds <= 600):
        parser.error('Require distinct ascending sizes 2..5000, 0..5 warm samples, 30..600 seconds')
    if args._directory:
        execute(args)
        return
    output = Path(args.output).resolve()
    if output.exists():
        parser.error('Refusing to overwrite a previous receipt')
    output.parent.mkdir(parents=True, exist_ok=True)
    PARENT.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix='assistant-e2e-', dir=PARENT) as directory:
        command = [sys.executable, '-B', str(Path(__file__).resolve()), *sys.argv[1:], '--_directory', directory]
        try:
            subprocess.run(command, check=True, timeout=args.wall_seconds)
        except subprocess.TimeoutExpired:
            if output.exists():
                report = json.loads(output.read_text(encoding='utf-8'))
                report.update(complete=False, stopped_reason='hard_wall_limit', hard_wall_limit_seconds=args.wall_seconds)
                output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
            raise SystemExit('Capacity check stopped at hard wall limit; partial receipt retained')


if __name__ == '__main__':
    main()

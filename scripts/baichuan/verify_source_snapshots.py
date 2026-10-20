"""Prepare private browser fixtures from restored source declarations, without AI.

Default invocation only describes the prepared verifier. Run --execute AFTER the
owned local source import, with its full --source-sha256. This creates two new
private conversations, leased query snapshots and controlled Excel artifacts.
It never upgrades source declarations to verified facts and never calls models,
research or search. Natural-language understanding is outside this verification.
"""
import argparse
from contextlib import ExitStack
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import uuid
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
LABEL = "源表确定规则验收"


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def load_local():
    spec = importlib.util.spec_from_file_location("source_fixture_local_model", ROOT / "scripts/baichuan/local_model.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def assert_denied(function, *args):
    from portal.baichuan_service import BaichuanError
    try:
        function(*args)
    except BaichuanError as error:
        assert error.status == 404, error.code
    else:
        raise AssertionError("Another owner could read a private object")


def execute(expected_hash):
    if not re.fullmatch(r"[0-9a-f]{64}", expected_hash or ""):
        raise ValueError("Provide the complete lowercase source SHA256")
    local_model = load_local()
    local = local_model.local
    output = local.RUNTIME / "source-browser-fixtures.json"
    # Lock coordinates owned process helpers and prevents an implicit rerun.
    with local.service_lock():
        if output.exists():
            raise ValueError("Private fixtures already exist; refusing duplicate writes")
        state = local.source_state()
        if not state["source_connected"] or state["source_sha256"] != expected_hash:
            raise ValueError("Completed owned source receipt does not match the supplied source hash")
        source = local.RUNTIME / ("source-restore-" + expected_hash[:8]) / "original.xlsx"
        if source.is_symlink() or not source.is_file() or sha256(source) != expected_hash:
            raise ValueError("Owned source backup file does not match the receipt")
        local_model.application()  # Existing owned SQLite guard; no network call.
        from django.conf import settings
        from django.db import connection, transaction
        from django.utils import timezone
        from openpyxl import load_workbook
        from portal.baichuan_assistant_api import snapshot_for
        from portal.baichuan_assistant_models import (BaichuanCertificate, BaichuanConversation,
            BaichuanEvidence, BaichuanFactRevision)
        from portal.baichuan_assistant_service import conversation_for, require_assistant, submit_message
        from portal.baichuan_assistant_worker import _reply
        from portal.baichuan_conditions import compile_conditions
        from portal.baichuan_exports import enqueue_export, export_for, open_export, process_export
        from portal.baichuan_jobs import claim_job, job_for, save_job
        from portal.baichuan_matching import save_match_snapshot
        from portal.baichuan_models import Company, CompanyPerson, CompanyQualification
        from portal.baichuan_service import require_access
        from portal.models import ModelCallLog, User
        database = Path(settings.DATABASES["default"]["NAME"]).resolve()
        assert database == local.RUNTIME / "portal.sqlite3"
        assert settings.BAICHUAN_ASSISTANT_SCOPE == "baichuan-ai-local"
        assert not settings.BAICHUAN_RESEARCH_ENABLED
        count = Company.objects.count()
        assert count > 0 and count == state["imported_company_count"]
        accounts = local.read_json(local.RUNTIME / "accounts.json")["accounts"]
        by_role = {account["role"]: account for account in accounts}
        assert {"product", "general_manager"} <= by_role.keys()
        actors = [(role, User.objects.get(username=by_role[role]["username"]))
                  for role in ("product", "general_manager")]
        assert actors[0][1].pk != actors[1][1].pk
        for role, actor in actors:
            require_assistant(actor)
            require_access(actor, write=True)
            assert actor.roles.filter(code=role).exists()

        def fact_digest():
            digest = hashlib.sha256()
            for model in (Company, CompanyQualification, CompanyPerson,
                          BaichuanCertificate, BaichuanEvidence, BaichuanFactRevision):
                digest.update(model._meta.label.encode())
                for row in model.objects.order_by("pk").values().iterator(chunk_size=200):
                    digest.update(json.dumps(row, sort_keys=True, ensure_ascii=False, default=str).encode())
            return digest.hexdigest()

        baseline = fact_digest()
        calls_before = ModelCallLog.objects.count()
        record = {"validation_label": LABEL, "source_sha256": expected_hash,
                  "scope": settings.BAICHUAN_ASSISTANT_SCOPE,
                  "environment_label": state["environment_label"], "company_count": count,
                  "db_sha256_before": sha256(database), "model_calls": 0,
                  "research_calls": 0, "natural_language_verified": False,
                  "validation_ready": False, "actors": []}
        # Reserve a private recovery record before writes. A failed run cannot
        # silently create a second set of conversations on implicit rerun.
        local.private_json(output, record)
        # Fail closed if future edits accidentally add external calls to this path.
        with ExitStack() as blocked:
            for target in ("socket.socket.connect", "socket.socket.connect_ex", "socket.create_connection",
                           "portal.baichuan_assistant_worker.generate_for_use"):
                blocked.enter_context(patch(target, side_effect=AssertionError("External calls forbidden in source verification")))
            for role, actor in actors:
                worker = "source-fixture-" + uuid.uuid4().hex
                # Reserve and claim before commit: a running local worker can
                # never claim this query and invoke its model planner.
                with transaction.atomic():
                    conversation = BaichuanConversation.objects.create(owner=actor, scope=settings.BAICHUAN_ASSISTANT_SCOPE)
                    _, queued, created = submit_message(actor, conversation.pk,
                        {"text": LABEL + "：仅验证当前源表声明，ITSS 资质待核实；不作合格结论。",
                         "attachment_ids": [], "parent_revision": 0,
                         "client_request_id": "source-fixture:" + uuid.uuid4().hex})
                    assert created
                    query = claim_job(worker, job_id=queued.pk, lease_seconds=300)
                    assert query and query.owner_id == actor.pk
                conditions = compile_conditions({"schema_version": 1, "revision": 1,
                    "must": {"field": "qualification", "category": "itss", "op": "present",
                             "value": "present", "holder": "bidder"},
                    "as_of": {"date": timezone.localdate().isoformat(), "basis": "query_date"},
                    "source_refs": [], "should": [], "unresolved": [],
                    "output": {"fields": ["name", "phone", "province", "city", "website", "credit_code"], "limit": 10}})
                save_job(query.pk, worker, query.fence, stage="matching",
                         checkpoint={"conditions": conditions, "validation_label": LABEL})
                snapshot = save_match_snapshot(actor, query.pk, worker, query.fence, conditions)
                assert snapshot.counts == {"confirmed": 0, "candidate": count, "excluded": 0}
                assert snapshot.research_status == "not_started" and len(snapshot.results) == count
                assert all(row["status"] == "candidate" and row["verdict"] == "unknown"
                           and all(item["verdict"] == "unknown" for item in row["conditions"])
                           for row in snapshot.results)
                summary = (LABEL + f"：当前在库 {count} 家，已确认 0 家，待核实 {count} 家。"
                           "仅按结构化确定条件生成固定快照，未调用模型或联网；源表声明未核实，不能视为满足资质。")
                _reply(query, worker, summary, {"snapshot_id": str(snapshot.pk), "validation_label": LABEL})
                save_job(query.pk, worker, query.fence, status="completed", counts=snapshot.counts,
                         result_ref={"snapshot_id": str(snapshot.pk)},
                         checkpoint={"conditions": conditions, "summary": summary, "validation_label": LABEL})
                row = {"role": role, "actor_id": actor.pk, "conversation_id": str(conversation.pk),
                       "snapshot_id": str(snapshot.pk), "query_job_id": str(query.pk),
                       "counts": snapshot.counts, "condition_revision": 1}
                for include, key in ((False, "default"), (True, "candidates")):
                    with transaction.atomic():
                        export_job = enqueue_export(actor, snapshot.pk,
                            {"client_request_id": "source-export:" + uuid.uuid4().hex,
                             "include_candidates": include, "limit": 10})
                        claimed = claim_job(worker, job_id=export_job.pk, lease_seconds=300)
                        assert claimed
                    exported = process_export(claimed, worker)
                    exported_job = job_for(actor, export_job.pk)
                    assert exported_job.status == "completed"
                    assert exported_job.counts == {"confirmed": 0, "candidate": min(10, count) if include else 0}
                    _, stream = open_export(actor, exported.pk)
                    try:
                        workbook = load_workbook(stream, read_only=True, data_only=False)
                        assert workbook.sheetnames == ["企业结果", "逐项依据", "待核实信息"]
                        assert workbook["企业结果"].max_row == 7
                        expected_rows = min(10, count) if include else 0
                        assert workbook["待核实信息"].max_row == 7 + expected_rows
                        assert workbook["逐项依据"].max_row == 7 + expected_rows
                        workbook.close()
                    finally:
                        stream.close()
                    row["default_export_id" if key == "default" else "export_id"] = str(exported.pk)
                    row["default_export_job_id" if key == "default" else "export_job_id"] = str(export_job.pk)
                record["actors"].append(row)
                output.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
            # Both users are authorized readers, but private results remain
            # inaccessible to the other owner via every relevant public guard.
            for index, row in enumerate(record["actors"]):
                other = actors[1 - index][1]
                for function, identifier in ((conversation_for, row["conversation_id"]),
                    (snapshot_for, row["snapshot_id"]), (job_for, row["query_job_id"]),
                    (export_for, row["export_id"]), (open_export, row["export_id"]),
                    (export_for, row["default_export_id"])):
                    assert_denied(function, other, identifier)
        assert ModelCallLog.objects.count() == calls_before
        assert fact_digest() == baseline, "Source companies, legacy declarations or evidence changed"
        record.update(owner_isolation_verified=True, source_facts_unchanged=True,
                      default_confirmed_exports_empty=True, candidate_exports_nonempty=True, validation_ready=True)
        connection.close()
        output.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"prepared": True, "fixture": str(output), "validation_label": LABEL,
                          "actors": len(actors), "model_calls": 0, "source_facts_unchanged": True}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--source-sha256")
    args = parser.parse_args()
    if not args.execute:
        print(json.dumps({"prepared_only": True, "database_accessed": False,
                          "execute_after_import": "--execute --source-sha256 FULL_SOURCE_SHA256"}))
        return
    execute(args.source_sha256)


if __name__ == "__main__":
    main()

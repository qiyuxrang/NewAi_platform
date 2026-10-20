"""Bounded, isolated AI-D28 worker/snapshot/default-export performance probe.

Only a newly created file SQLite and private artifacts are writable. Model
understanding is excluded from capacity timing; a separate delayed local stub
exercises the real elapsed-budget failure path. No HTTP/model/network is used.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack
from datetime import datetime, timedelta, timezone
import hashlib
from io import StringIO
import json
import math
import os
from pathlib import Path
import platform
import statistics
import subprocess
import sys
import threading
from time import perf_counter, sleep
import traceback
from unittest.mock import patch
import uuid

ROOT = Path(__file__).resolve().parents[2]
NIGHT = ROOT / '.runtime/baichuan-ai-local/night-acceptance-20261018'


def now():
    return datetime.now(timezone.utc).isoformat()


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def save(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf8')


def summarize(samples):
    return {k: dict(samples=len(samples), median_ms=round(statistics.median(r[k] for r in samples), 3),
                    p95_ms=round(sorted(r[k] for r in samples)[math.ceil(.95 * len(samples)) - 1], 3))
            for k in samples[0] if k.endswith('_ms')}


def execute(args):
    directory = Path(args.directory).resolve(strict=True)
    assert directory.parent.parent == NIGHT.resolve() and directory.name == 'isolated'
    for key in list(os.environ):
        if key.startswith(('PORTAL_', 'PLATFORM_', 'MODEL_GATEWAY_', 'COMPANY_TRIAL_')):
            del os.environ[key]
    os.environ.update(DJANGO_SETTINGS_MODULE='config.settings', PORTAL_DEBUG='1', PORTAL_HTTPS='0',
        PORTAL_SECRET_KEY='fictional-trusted-benchmark-only-not-for-deployment-20261018',
        PORTAL_SQLITE_PATH=str(directory / 'synthetic.sqlite3'),
        PORTAL_BAICHUAN_STORAGE_ROOT=str(directory / 'private'), PORTAL_PRODUCT_STORAGE_ROOT=str(directory / 'product'),
        PORTAL_BAICHUAN_SCOPE_CONFIRMED='1', PORTAL_BAICHUAN_READ_ROLE_CODES='product',
        PORTAL_BAICHUAN_ASSISTANT_ENABLED='1', PORTAL_BAICHUAN_ASSISTANT_SCOPE='fictional-trusted-capacity',
        PORTAL_BAICHUAN_TRUST_LEGACY_QUALIFICATIONS='1', PORTAL_BAICHUAN_RESEARCH_ENABLED='0',
        PORTAL_BAICHUAN_FACT_WRITE_ENABLED='0', PORTAL_PRODUCT_MODEL_CALLS_ALLOWED='0', PORTAL_PRODUCT_WEB_SEARCH_ENABLED='0')
    sys.path[:0] = [str(ROOT / 'services/business/backend'), str(ROOT / 'services/business')]
    import django
    django.setup()
    from django.conf import settings
    from django.core.management import call_command
    from django.db import close_old_connections, connection, connections, transaction
    from django.utils import timezone as djtime
    from django.test import override_settings
    from rest_framework.test import APIRequestFactory, force_authenticate
    from openpyxl import load_workbook
    from portal import baichuan_assistant_api as api, baichuan_assistant_worker as worker
    from portal import baichuan_exports as exports, baichuan_matching as matching
    from portal.baichuan_assistant_models import BaichuanCertificate, BaichuanConversation, BaichuanExport, BaichuanJob
    from portal.baichuan_assistant_service import submit_message
    from portal.baichuan_conditions import compile_conditions
    from portal.baichuan_models import Company, CompanyQualification
    from portal.baichuan_rules import install_compatibility_rules
    from portal.baichuan_service import BaichuanError
    from portal.models import ModelCallLog, Role, User
    assert Path(settings.DATABASES['default']['NAME']).resolve() == directory / 'synthetic.sqlite3'
    assert settings.BAICHUAN_TRUST_LEGACY_QUALIFICATIONS and not settings.BAICHUAN_RESEARCH_ENABLED
    report = dict(task='AI-T16', status='STARTED', started_at=now(), synthetic=True,
        current_policy='AI-D28 trusted current legacy exact-grade bidder', real_model_calls=0, network_calls=0,
        forbidden_network_attempts=0, model_stub_calls=0, target_database_accessed=False, frontend_changes=0,
        sizes=[1000, 5000], warm_samples=args.warm_samples, datasets=[], concurrency=[], skipped=0,
        python=sys.version, django=django.get_version(), sqlite=__import__('sqlite3').sqlite_version,
        openpyxl=__import__('openpyxl').__version__, platform=platform.platform(), processor=platform.processor(),
        logical_cpus=os.cpu_count(), sqlite_options=settings.DATABASES['default']['OPTIONS'],
        cold_definition='Close Django DB connection before first sample at each size; OS/page cache not evicted.',
        p95_definition='Nearest rank over explicitly limited warm samples; no population or end-to-end model guarantee.',
        timing_scope='Real submit/checkpoint resume -> worker -> fixed snapshot -> rendered authorized API -> default confirmed XLSX -> authorized open/readback.',
        fixture='index mod7: exact2; othergrade1; UNKNOWN; only new unverified2; expired legacy2; legacy2+revoked unverified2; explicit ABSENT')
    output = directory.parent / 'report.json'

    def persist():
        save(output, report)

    def forbid(*a, **k):
        report['forbidden_network_attempts'] += 1
        raise AssertionError('Network/real model forbidden in synthetic benchmark')

    def uid(kind, index):
        return uuid.uuid5(uuid.NAMESPACE_URL, f'fictional-trusted-capacity/{kind}/{index}')

    def protected():
        return {model.__name__: digest(list(model.objects.order_by('pk').values()))
                for model in (Company, CompanyQualification, BaichuanCertificate)}

    def error_code(fn, expected='not_found'):
        try:
            fn()
        except BaichuanError as e:
            assert e.code == expected, e.code
            return
        raise AssertionError('Expected protected operation denial')

    try:
        with ExitStack() as guards:
            for target in ('socket.socket.connect', 'socket.socket.connect_ex', 'socket.create_connection',
                           'socket.getaddrinfo', 'portal.baichuan_assistant_worker.generate_for_use',
                           'portal.model_gateway.generate_for_use'):
                guards.enter_context(patch(target, side_effect=forbid))
            call_command('migrate', verbosity=0, stdout=StringIO())
            call_command('seed_portal', stdout=StringIO())
            actor = User.objects.create_user(username='fictional-trusted-owner', must_change_password=False)
            outsider = User.objects.create_user(username='fictional-trusted-other-owner', must_change_password=False)
            for user in (actor, outsider):
                user.roles.add(Role.objects.get(code='product'))
            # Fixture setup only: installer retains its real write permission
            # gate; benchmark users return to read-only scope immediately.
            with override_settings(PORTAL_BAICHUAN_WRITE_USER_IDS=[actor.pk]):
                install_compatibility_rules(actor)
            day = djtime.localdate()
            request_factory = APIRequestFactory()

            def condition(count):
                return compile_conditions(dict(schema_version=1, revision=1,
                    must=dict(field='qualification', category='electronic', op='grade_exact', value='2', holder='bidder'),
                    as_of=dict(date=day.isoformat(), basis='query_date'), output=dict(fields=['name', 'phone'], limit=count)))

            def submit(count, checkpoint=True):
                convo = BaichuanConversation.objects.create(owner=actor, scope=settings.BAICHUAN_ASSISTANT_SCOPE)
                start = perf_counter()
                with transaction.atomic():
                    _, job, created = submit_message(actor, convo.pk, dict(text='合成性能样本：投标人电子与智能化精确二级',
                        attachment_ids=[], parent_revision=0, client_request_id=uuid.uuid4().hex))
                    assert created
                    if checkpoint:
                        job.checkpoint = dict(conditions=condition(count), synthetic_capacity_fixture=True)
                        job.save(update_fields=['checkpoint', 'updated_at'])
                return job, (perf_counter() - start) * 1000

            def response_for(user, snapshot):
                request = request_factory.get('/synthetic/results/', dict(page=1, page_size=50))
                force_authenticate(request, user=user)
                result = api.results(request, snapshot_id=snapshot.pk)
                result.render()
                return result

            for count in (1000, 5000):
                previous = Company.objects.count()
                companies = [Company(id=uid('company', i), name=f'纯虚构性能企业{i:05}', phone=f'000TEST{i:05}',
                    created_by=actor, updated_by=actor) for i in range(previous, count)]
                Company.objects.bulk_create(companies, batch_size=200)
                quals, certs = [], []
                for i, company in enumerate(companies, previous):
                    group = i % 7
                    if group != 3:
                        quals.append(CompanyQualification(id=uid('qualification', i), company=company, category='electronic',
                            status='UNKNOWN' if group == 2 else ('ABSENT' if group == 6 else 'PRESENT'),
                            grade=None if group in (2, 6) else ('1' if group == 1 else '2'),
                            valid_until=day-timedelta(days=1) if group == 4 else None))
                    if group in (3, 5):
                        certs.append(BaichuanCertificate(id=uid('certificate', i), company=company, category='electronic', grade='2',
                            holder_type='bidder', holder_company_id=company.pk, assertion_status='PRESENT',
                            verification_status='unverified', revocation_status='revoked' if group == 5 else 'unknown'))
                CompanyQualification.objects.bulk_create(quals, batch_size=200)
                BaichuanCertificate.objects.bulk_create(certs, batch_size=200)
                expected_ids = {str(uid('company', i)) for i in range(count) if i % 7 == 0}
                expected_counts = dict(confirmed=len(expected_ids), excluded=sum(i % 7 == 6 for i in range(count)))
                expected_counts['candidate'] = count - sum(expected_counts.values())
                dataset = dict(companies=count, qualifications=CompanyQualification.objects.count(),
                    unverified_certificates=BaichuanCertificate.objects.count(), expected_counts=expected_counts,
                    expected_confirmed_ids_sha256=digest(sorted(expected_ids)), fixture_sha256=protected(), samples=[])
                report['datasets'].append(dataset); persist()
                for attempt in range(args.warm_samples + 1):
                    if attempt == 0:
                        connection.close()
                    timing = {}
                    start = perf_counter();job, timing['submit_ms'] = submit(count)
                    original_match, original_save = matching.match_companies, worker.save_match_snapshot
                    def timed(fn, key):
                        def wrap(*a, **k):
                            begin = perf_counter()
                            try:return fn(*a, **k)
                            finally:timing[key] = (perf_counter()-begin)*1000
                        return wrap
                    begin = perf_counter()
                    with patch.object(matching, 'match_companies', timed(original_match, 'structured_match_ms')), \
                            patch.object(worker, 'save_match_snapshot', timed(original_save, 'match_snapshot_ms')):
                        assert worker.run_one('synthetic-trusted-query', job.pk)
                    timing['query_worker_ms'] = (perf_counter()-begin)*1000
                    timing['first_snapshot_ready_ms'] = (perf_counter()-start)*1000
                    job.refresh_from_db();assert job.status == 'completed' and job.counts.get('_budget_model_calls', 0) == 0
                    snapshot = job.snapshots.get();assert snapshot.counts == expected_counts and len(snapshot.results) == count
                    assert {r['company_id'] for r in snapshot.results if r['status']=='confirmed'} == expected_ids
                    assert all('用户确认的旧资质库' in r['conditions'][0]['reason'] for r in snapshot.results if r['status']=='confirmed')
                    assert snapshot.research_status=='not_requested' and job.checkpoint['research_policy']=='independent_manifest_only'
                    begin=perf_counter();response=response_for(actor,snapshot);timing['authorized_rendered_results_ms']=(perf_counter()-begin)*1000
                    assert response.status_code==200 and response.data['counts']==expected_counts and len(response.data['results'])==50
                    assert response_for(outsider,snapshot).status_code==404
                    begin=perf_counter();export_job=exports.enqueue_export(actor,snapshot.pk,dict(client_request_id=uuid.uuid4().hex));timing['default_export_submit_ms']=(perf_counter()-begin)*1000
                    assert not export_job.export.selection['include_candidates']
                    begin=perf_counter();assert worker.run_one('synthetic-trusted-export',export_job.pk);timing['export_worker_ms']=(perf_counter()-begin)*1000
                    export_job.refresh_from_db();assert export_job.status=='completed'
                    artifact=BaichuanExport.objects.get(job=export_job)
                    error_code(lambda:exports.open_export(outsider,artifact.pk))
                    begin=perf_counter();authorized,stream=exports.open_export(actor,artifact.pk);timing['authorized_download_open_ms']=(perf_counter()-begin)*1000
                    try:
                        begin=perf_counter();book=load_workbook(stream,read_only=True,data_only=False)
                        assert book.sheetnames==['企业结果','逐项依据','待核实信息']
                        actual=list(book['企业结果'].values)[7:]
                        expected=[[r['company']['name'],r['company']['phone']] for r in snapshot.results if r['status']=='confirmed']
                        assert [[r[2],r[3] or ''] for r in actual]==expected
                        assert book['待核实信息'].max_row==7 and book['逐项依据'].max_row==7+len(expected)
                        assert all(c.data_type!='f' for sheet in book for row in sheet for c in row)
                        book.close();timing['xlsx_readback_validate_ms']=(perf_counter()-begin)*1000
                    finally:stream.close()
                    timing['total_ms']=(perf_counter()-start)*1000
                    sample=dict(cache='connection_cold' if attempt==0 else 'warm',attempt=attempt,
                        **{k:round(v,3) for k,v in timing.items()},counts=snapshot.counts,query_id=str(job.pk),snapshot_id=str(snapshot.pk),
                        default_confirmed_export_rows=len(expected),xlsx_sha256=authorized.file_hash,xlsx_bytes=authorized.size_bytes,
                        unauthorized_snapshot_and_export_denied=True)
                    dataset['samples'].append(sample)
                    if attempt:
                        dataset['warm_summary']=summarize(dataset['samples'][1:])
                    persist();print(json.dumps(dict(companies=count,attempt=attempt,match_ms=sample['structured_match_ms'],total_ms=sample['total_ms'])),flush=True)
                assert protected()==dataset['fixture_sha256']
                dataset['meets_2s_structured_p95']=dataset['warm_summary']['structured_match_ms']['p95_ms']<=2000
                dataset['meets_5s_snapshot_excluding_model_p95']=dataset['warm_summary']['first_snapshot_ready_ms']['p95_ms']<=5000

            baseline=protected()
            jobs=[submit(5000)[0] for _ in range(2)]
            gate=threading.Barrier(2);intervals=[];interval_lock=threading.Lock();original_match=matching.match_companies
            def overlapping_match(*a,**k):
                entered=perf_counter();gate.wait(timeout=20)
                try:return original_match(*a,**k)
                finally:
                    with interval_lock:intervals.append([entered,perf_counter()])
            def run(index, job_id, gate=None):
                close_old_connections()
                try:
                    if gate:gate.wait(timeout=20)
                    begin=perf_counter();claimed=worker.run_one(f'synthetic-parallel-{index}',job_id)
                    return dict(index=index,claimed=claimed,elapsed_ms=round((perf_counter()-begin)*1000,3))
                finally:connections.close_all()
            begin=perf_counter()
            with patch.object(matching,'match_companies',overlapping_match),ThreadPoolExecutor(max_workers=2) as pool:
                futures=[pool.submit(run,i,j.pk) for i,j in enumerate(jobs)];outcomes=[f.result(timeout=60) for f in futures]
            assert len(intervals)==2 and max(x[0] for x in intervals)<min(x[1] for x in intervals)
            for j in jobs:
                j.refresh_from_db();assert j.status=='completed' and j.snapshots.count()==1 and j.attempts==1
                assert j.snapshots.get().counts==report['datasets'][-1]['expected_counts']
            report['concurrency'].append(dict(case='two_real_workers_independent_jobs',workers=2,companies=5000,
                wall_ms=round((perf_counter()-begin)*1000,3),outcomes=outcomes,matching_intervals=intervals,
                actual_matching_overlap=True,one_snapshot_each=True))
            job=submit(5000)[0];gate=threading.Barrier(2)
            with ThreadPoolExecutor(max_workers=2) as pool:
                futures=[pool.submit(run,i,job.pk,gate) for i in range(2)];outcomes=[f.result(timeout=60) for f in futures]
            job.refresh_from_db();assert sum(x['claimed'] for x in outcomes)==1 and job.status=='completed' and job.snapshots.count()==1 and job.attempts==job.fence==1
            report['concurrency'].append(dict(case='same_job_two_real_workers',outcomes=outcomes,unique_snapshot=True,attempts=job.attempts,fence=job.fence))
            persist()

            # Independent OS process holds the real SQLite write reservation.
            job=submit(5000)[0];ready=directory/'lock-ready';connection.close()
            holder_source="""import sqlite3,sys,time,socket,json
from pathlib import Path
def deny(*a,**k):raise AssertionError('network forbidden')
socket.socket=socket.create_connection=socket.getaddrinfo=deny
db=sqlite3.connect(sys.argv[1],timeout=10);db.execute('BEGIN IMMEDIATE')
start=time.monotonic();Path(sys.argv[2]).write_text(str(start));time.sleep(.6);db.commit();db.close()
print(json.dumps({'held_ms':(time.monotonic()-start)*1000,'network_calls':0}))
"""
            (directory/'lock-holder.py').write_text(holder_source,encoding='utf8')
            process=subprocess.Popen([sys.executable,'-B',str(directory/'lock-holder.py'),str(directory/'synthetic.sqlite3'),str(ready)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            try:
                deadline=perf_counter()+10
                while not ready.exists():
                    assert perf_counter()<deadline and process.poll() is None,'Lock holder failed to start'
                    sleep(.005)
                timings={};original_claim=worker.claim_job
                def timed_claim(*a,**k):
                    begin=perf_counter()
                    try:return original_claim(*a,**k)
                    finally:timings['claim_ms']=(perf_counter()-begin)*1000
                with patch.object(worker,'claim_job',timed_claim):assert worker.run_one('synthetic-lock-contention',job.pk)
                stdout,stderr=process.communicate(timeout=10);assert process.returncode==0,stderr.decode(errors='replace')
                lock_result=json.loads(stdout);assert timings['claim_ms']>=300 and lock_result['held_ms']>=590
                job.refresh_from_db();assert job.status=='completed' and job.snapshots.count()==1
                report['concurrency'].append(dict(case='OS_process_SQLite_BEGIN_IMMEDIATE_contention',**timings,holder=lock_result,terminal=job.status,unique_snapshot=True))
            finally:
                if process.poll() is None:process.kill();process.communicate(timeout=5)

            # No real model invocation: delayed in-memory stub triggers natural
            # elapsed-budget exhaustion at the next real fenced checkpoint.
            job=submit(5000,checkpoint=False)[0];job.budget={**job.budget,'timeout_seconds':1,'max_model_calls':1};job.save(update_fields=['budget'])
            def delayed_stub(*a,**k):
                report['model_stub_calls']+=1;sleep(1.2)
                return dict(content=json.dumps(condition(5000)))
            begin=perf_counter()
            with patch.object(worker,'generate_for_use',side_effect=delayed_stub):assert worker.run_one('synthetic-delay-budget',job.pk)
            elapsed=(perf_counter()-begin)*1000;job.refresh_from_db()
            assert job.status=='failed' and job.error_code=='budget_exhausted' and not job.snapshots.exists()
            assert job.counts['_budget_model_calls']==1 and report['model_stub_calls']==1
            report['delayed_stub_budget_failure']=dict(elapsed_ms=round(elapsed,3),budget_seconds=1,stub_delay_seconds=1.2,
                status=job.status,error_code=job.error_code,reserved_model_calls=1,real_model_calls=0,snapshots=0,automatic_retries=job.attempts-1)
            assert protected()==baseline and ModelCallLog.objects.count()==0 and report['forbidden_network_attempts']==0
            report.update(status='PASS',completed_at=now(),fixture_rows_unchanged=True,real_model_log_count=0)
            persist()
    except Exception:
        report.update(status='FAIL',failed_at=now(),traceback=traceback.format_exc());persist();raise
    finally:connections.close_all()


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-directory',required=True)
    parser.add_argument('--warm-samples',type=int,default=5)
    parser.add_argument('--wall-seconds',type=int,default=360)
    parser.add_argument('--directory',help=argparse.SUPPRESS)
    args=parser.parse_args()
    if not 2<=args.warm_samples<=10 or not 60<=args.wall_seconds<=600:parser.error('Require 2..10 warm samples, 60..600 wall seconds')
    if args.directory:return execute(args)
    output=Path(args.output_directory).resolve()
    if output.parent!=NIGHT.resolve() or output.exists():parser.error('Require a new unique direct child directory of private night-acceptance-20261018')
    output.mkdir();(output/'isolated').mkdir()
    reg=json.loads((NIGHT/'trusted-legacy-final-regression/receipt.json').read_text(encoding='utf8'))
    frontend=json.loads((ROOT/'.runtime/baichuan-ai-local/frontend-freeze-20261017.json').read_text(encoding='utf8'))
    inputs={**reg['input_hashes'],**frontend['files']}
    assert len(reg['input_hashes'])==64 and len(frontend['files'])==163
    assert all(sha(ROOT/p)==h for p,h in inputs.items())
    inputs[str(Path(__file__).relative_to(ROOT))]=sha(__file__)
    inputs['scripts/baichuan/benchmark_assistant_e2e.py']=sha(ROOT/'scripts/baichuan/benchmark_assistant_e2e.py')
    save(output/'frozen-inputs.json',inputs)
    command=[sys.executable,'-B','-X','utf8',str(Path(__file__).resolve()),*sys.argv[1:],'--directory',str(output/'isolated')]
    start=perf_counter();code=None;failure=None
    with (output/'execution.log').open('wb') as log:
        try:code=subprocess.run(command,stdout=log,stderr=subprocess.STDOUT,timeout=args.wall_seconds,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0)).returncode
        except subprocess.TimeoutExpired:failure='hard_wall_limit'
    unchanged=all(sha(ROOT/p)==h for p,h in inputs.items())
    report=json.loads((output/'report.json').read_text(encoding='utf8')) if (output/'report.json').exists() else {}
    receipt=dict(task='AI-T16',started_at=report.get('started_at'),completed_at=now(),command=command,returncode=code,
        hard_wall_seconds=args.wall_seconds,elapsed_seconds=round(perf_counter()-start,3),failure=failure,
        status='PASS' if code==0 and report.get('status')=='PASS' and unchanged else 'FAIL',
        current64_and_frontend163_unchanged=unchanged,input_count=len(inputs),input_manifest_sha256=sha(output/'frozen-inputs.json'),
        script_sha256=sha(__file__),report_sha256=sha(output/'report.json') if report else None,log_sha256=sha(output/'execution.log'),
        database_sha256=sha(output/'isolated/synthetic.sqlite3'),real_model_network_calls=0,skipped=report.get('skipped'))
    save(output/'receipt.json',receipt);print(json.dumps(receipt,ensure_ascii=False))
    raise SystemExit(0 if receipt['status']=='PASS' else 1)


if __name__=='__main__':main()

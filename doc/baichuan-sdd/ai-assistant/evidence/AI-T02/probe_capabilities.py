"""Reproducible synthetic capability probes; never loads deployment .env or database.

Run with .runtime/company-test/python/Scripts/python.exe.
Optional --worker-tests runs selected existing tests against a temporary SQLite DB.
All fixtures are fictional, temporary and contain no enterprise data or credentials.
"""
from __future__ import annotations

import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import zipfile

REPO = Path(__file__).resolve().parents[5]
BACKEND = REPO / "services/business/backend"
PARSER = BACKEND / "portal/source_parsers/cli.py"
OCR_PYTHON = REPO / ".runtime/company-test/cloud-ocr/Scripts/python.exe"


def safe_environment():
    # Explicit allowlist: inherited database, cloud, model and proxy secrets absent.
    return {key: value for key, value in os.environ.items() if key.upper() in {
        "SYSTEMROOT", "WINDIR", "TEMP", "TMP", "PATH", "PATHEXT", "COMSPEC",
        "USERPROFILE", "LOCALAPPDATA", "APPDATA", "PROGRAMDATA", "PROGRAMFILES",
        "PROGRAMFILES(X86)", "SYSTEMDRIVE", "PROCESSOR_ARCHITECTURE",
    }}


def configure(directory):
    clean = safe_environment()
    os.environ.clear()
    os.environ.update(clean)
    os.environ.update(PORTAL_SECRET_KEY="synthetic-baichuan-capability-probe-20261017-only",
                      PORTAL_DEBUG="1", PORTAL_HTTPS="0", PYTHONIOENCODING="utf-8",
                      DJANGO_SETTINGS_MODULE="config.settings",
                      PORTAL_SQLITE_PATH=str(directory / "unused-isolated.sqlite3"))
    sys.path.insert(0, str(BACKEND))
    sys.path.insert(0, str(BACKEND.parent))
    import django
    django.setup()
    from django.conf import settings
    assert settings.DATABASES["default"]["ENGINE"] == "django.db.backends.sqlite3"
    assert str(settings.DATABASES["default"]["NAME"]) == str(directory / "unused-isolated.sqlite3")
    return settings


def parser_probe(source, directory, *, ocr=True):
    output = directory / (source.stem + ("-ocr" if ocr else "-no-ocr") + ".json")
    environment = safe_environment()
    environment.update(PRODUCT_PARSER_OCR="1" if ocr else "0", PYTHONIOENCODING="utf-8")
    start = time.monotonic()
    completed = subprocess.run([str(OCR_PYTHON), "-I", str(PARSER), str(source), str(output)],
                               env=environment, capture_output=True, timeout=60,
                               creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
    payload = json.loads(output.read_text(encoding="utf-8"))
    parsed = payload.get("parsed", {})
    return {"exit_code": completed.returncode, "elapsed_seconds": round(time.monotonic() - start, 3),
            "ok": payload.get("ok"), "error": payload.get("code"),
            "status": parsed.get("status"), "method": parsed.get("method"),
            "blocks": [{"kind": b["kind"], "text": b["text"], "location": b["location"],
                        **({"confidence": b["confidence"]} if "confidence" in b else {})}
                       for b in parsed.get("blocks", [])],
            "warning_codes": [w["code"] for w in parsed.get("warnings", [])]}


def fixtures(directory):
    import fitz
    from PIL import Image, ImageDraw, ImageFont
    document = fitz.open()
    page = document.new_page()
    page.insert_text((50, 60), "SYNTHETIC BIDDER: ITSS LEVEL 2 REQUIRED", fontsize=15)
    page = document.new_page()
    page.insert_text((50, 60), "SYNTHETIC MANUFACTURER: ISO 9001 PREFERRED", fontsize=15)
    text_pdf = directory / "synthetic-text.pdf"
    document.save(text_pdf)
    document.close()
    image = Image.new("RGB", (1600, 600), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 52)
    draw.text((50, 100), "虚构测试企业：投标人须具备 ITSS 二级", font=font, fill="black")
    draw.text((50, 220), "制造商 ISO 9001 证书仅为优先项", font=font, fill="black")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")
    document = fitz.open()
    page = document.new_page(width=800, height=300)
    page.insert_image(page.rect, stream=image_bytes.getvalue())
    scanned = directory / "synthetic-scanned.pdf"
    document.save(scanned)
    document.close()
    docx = directory / "synthetic-clauses.docx"
    xml = '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
<w:p><w:r><w:t>虚构招标：投标人必须具备 ITSS 二级。</w:t></w:r></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>制造商</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>ISO 9001 优先项</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
</w:body></w:document>'''
    with zipfile.ZipFile(docx, "w") as archive:
        archive.writestr("word/document.xml", xml.encode())
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
    return text_pdf, scanned, docx


def xlsx_probe(directory):
    import openpyxl
    output = directory / "synthetic-export.xlsx"
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    values = ["00123456789", "001234567890123456", "0000123", "=1+1", "+SUM(A1:A2)", "@test"]
    for name in ["企业结果", "逐项依据", "待核实信息"]:
        sheet = workbook.create_sheet(name)
        for index, value in enumerate(values, 1):
            cell = sheet.cell(index, 1, value)
            cell.data_type = "s"
            cell.number_format = "@"
    workbook.save(output)
    loaded = openpyxl.load_workbook(output, data_only=False)
    observed = [cell for row in loaded.worksheets[0].iter_rows() for cell in row]
    assert [cell.value for cell in observed] == values
    assert all(cell.data_type == "s" and cell.number_format == "@" for cell in observed)
    return {"status": "pass", "sheets": loaded.sheetnames, "bytes": output.stat().st_size,
            "leading_zero_and_formula_strings_preserved_as_text": True,
            "note": "Writer dependency only; business export, authorization and immutable snapshots not implemented."}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    with tempfile.TemporaryDirectory(prefix="baichuan-AI-T02-", dir=REPO / ".runtime") as temp:
        directory = Path(temp)
        settings = configure(directory)
        if "--public-search-tests" in sys.argv:
            from django.core.management import call_command
            settings.DATABASES["default"]["TEST"]["NAME"] = str(directory / "test.sqlite3")
            call_command("test", "portal.tests.test_baichuan_bing_rss", "portal.tests.test_product_web_search",
                "portal.tests.test_baichuan_research_runtime", interactive=False, verbosity=2)
            return
        if "--file-condition-tests" in sys.argv:
            from django.core.management import call_command
            settings.DATABASES["default"]["TEST"]["NAME"] = str(directory / "test.sqlite3")
            call_command("test", "portal.tests.test_baichuan_file_conditions", "portal.tests.test_baichuan_integration_review",
                "portal.tests.test_baichuan_assistant", "portal.tests.test_baichuan_conditions", interactive=False, verbosity=2)
            return
        if "--review-tests" in sys.argv:
            from django.core.management import call_command
            settings.DATABASES["default"]["TEST"]["NAME"] = str(directory / "test.sqlite3")
            call_command("test", "portal.tests.test_baichuan_integration_review", "portal.tests.test_baichuan_assistant", "portal.tests.test_baichuan_jobs", interactive=False, verbosity=2)
            return
        if "--job-export-tests" in sys.argv:
            from django.core.management import call_command
            settings.DATABASES["default"]["TEST"]["NAME"] = str(directory / "test.sqlite3")
            call_command("test", "portal.tests.test_baichuan_jobs", "portal.tests.test_baichuan_exports", interactive=False, verbosity=2)
            return
        if "--export-tests" in sys.argv:
            from django.core.management import call_command
            settings.DATABASES["default"]["TEST"]["NAME"] = str(directory / "test.sqlite3")
            call_command("test", "portal.tests.test_baichuan_exports", interactive=False, verbosity=2)
            return
        if "--jobs-tests" in sys.argv:
            from django.core.management import call_command
            settings.DATABASES["default"]["TEST"]["NAME"] = str(directory / "test.sqlite3")
            call_command("test", "portal.tests.test_baichuan_jobs", interactive=False, verbosity=2)
            return
        if "--worker-tests" in sys.argv:
            from django.core.management import call_command
            settings.DATABASES["default"]["TEST"]["NAME"] = str(directory / "test.sqlite3")
            call_command("test",
                         "portal.tests.test_product_worker_command",
                         "portal.tests.test_product_worker.ProductWorkerTests.test_expired_lease_reclaims_with_new_fence_and_late_write_rejected",
                         "portal.tests.test_product_worker.ProductWorkerTests.test_cancellation_fences_inflight_output",
                         interactive=False, verbosity=2)
            return
        results = {"python": sys.executable, "version": sys.version.split()[0],
                   "isolation": "deployment .env not loaded; DB never opened; fixtures temporary and fictional"}
        if "--web-probe" in sys.argv:
            from portal.product_assistant_tools import fetch_public_page
            import hashlib
            results["public_pages"] = []
            for url in ["https://www.python.org/", "https://example.com/"]:
                start = time.monotonic()
                try:
                    text = fetch_public_page(url)
                    row = {"status": "pass", "url": url, "characters": len(text),
                           "sha256": hashlib.sha256(text.encode()).hexdigest()}
                except Exception as error:
                    row = {"status": "failed", "url": url, "exception_type": type(error).__name__}
                    if str(error) in {"page unavailable, redirected or unsupported format",
                                      "page exceeds reading limit", "empty page",
                                      "only public HTTPS pages are supported"}:
                        row["reason"] = str(error)
                row["elapsed_seconds"] = round(time.monotonic() - start, 3)
                results["public_pages"].append(row)
            # Raw responses and exception messages can contain upstream data; omit both.
            print(json.dumps(results, ensure_ascii=False, indent=2))
            return
        from portal.model_gateway import _gateway_request, GatewayError
        from portal.product_web_search import search_public_query
        from portal.product_service import ProductError
        try:
            _gateway_request({"messages": [{"role": "user", "content": "synthetic probe"}]})
            results["model_gateway"] = {"status": "unexpected_configured"}
        except GatewayError as error:
            results["model_gateway"] = {"status": "unavailable_in_isolation", "code": error.code,
                                        "external_calls": 0}
        try:
            results["search"] = {"results": len(search_public_query("公开技术标准 ITSS"))}
        except ProductError as error:
            results["search"] = {"status": "unavailable_in_isolation", "code": error.code,
                                 "external_calls": 0}
        text_pdf, scanned, docx = fixtures(directory)
        results["text_pdf"] = parser_probe(text_pdf, directory)
        results["docx"] = parser_probe(docx, directory)
        results["scanned_pdf"] = parser_probe(scanned, directory)
        results["scanned_pdf_ocr_disabled"] = parser_probe(scanned, directory, ocr=False)
        malformed = directory / "synthetic-broken.pdf"
        malformed.write_bytes(b"%PDF-1.7\nsynthetic broken")
        results["malformed_pdf"] = parser_probe(malformed, directory)
        assert [b["location"]["page"] for b in results["text_pdf"]["blocks"]] == [1, 2]
        assert results["docx"]["blocks"][0]["text"] == "虚构招标：投标人必须具备 ITSS 二级。"
        assert results["docx"]["blocks"][1]["location"]["table"] == 1
        assert results["scanned_pdf"]["method"] == "ocr"
        assert "ITSS" in results["scanned_pdf"]["blocks"][0]["text"]
        assert "ocr_review_required" in results["scanned_pdf"]["warning_codes"]
        assert "ocr_unavailable" in results["scanned_pdf_ocr_disabled"]["warning_codes"]
        assert not results["scanned_pdf_ocr_disabled"]["blocks"]
        assert results["malformed_pdf"]["error"] == "invalid_document"
        results["xlsx"] = xlsx_probe(directory)
        import importlib.metadata
        results["package_versions"] = {name: importlib.metadata.version(name) for name in
                                       ["Django", "pypdf", "PyMuPDF", "openpyxl", "Pillow"]}
        print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
